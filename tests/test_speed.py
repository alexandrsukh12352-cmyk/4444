"""Ускорение: быстрое чтение xlsx, общий индекс цен на все группы, прайсы прошлых дат по требованию.

Всё это не должно менять ни одной цифры: тесты сверяют результат с прежним способом.
"""
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))

from pricing import excel, prices, report as R  # noqa: E402
from pricing.sources import AcpCross, AcpItem, Offer, Sources, norm_article, scan  # noqa: E402


# ---------------------------------------------------------------- быстрое чтение xlsx


def _tricky_book(path: Path, date1904: bool = False) -> None:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    wb = openpyxl.Workbook()
    if date1904:
        wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "Прайс"
    ws["A1"] = "Прайс-лист"
    ws.merge_cells("A1:F2")
    ws["A3"] = dt.datetime(2026, 10, 6, 13, 5)
    ws["B3"] = dt.date(2025, 1, 1)
    ws["C3"] = dt.time(12, 30)
    ws["D3"] = dt.timedelta(hours=30)
    ws["D3"].number_format = "[h]:mm:ss"
    ws["E3"] = True
    ws["F3"] = False
    ws["A4"], ws["B4"], ws["C4"], ws["D4"] = 1, 1.5, 1e20, -0.0001
    ws["E4"] = "  пробелы  "
    ws["F4"] = "_x000D_ строка"
    ws["A5"] = CellRichText(["обычный ", TextBlock(InlineFont(b=True), "жирный")])
    ws["B5"] = "=1+1"
    ws["C5"] = "#N/A"
    ws["A6"] = 44000
    ws["A6"].number_format = "DD.MM.YYYY"
    ws["H10"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFFF00")
    ws.merge_cells("J12:L14")
    ws["J12"] = "объединено"
    wb.create_sheet("Пустой")
    only = wb.create_sheet("Только объединение")
    only.merge_cells("B2:C3")
    hidden = wb.create_sheet("Скрытый")
    hidden["A1"] = "x"
    hidden.sheet_state = "hidden"
    wb.save(path)


def _grids(path: Path, fast: bool, monkeypatch) -> list:
    monkeypatch.setattr(excel, "FAST_XLSX", fast)
    return [(s.name, [[(type(v), v) for v in row] for row in s.grid]) for s in excel.read_sheets(path)]


def test_fast_xlsx_gives_the_same_grid_as_openpyxl(tmp_path, monkeypatch):
    for date1904 in (False, True):
        path = tmp_path / f"книга{int(date1904)}.xlsx"
        _tricky_book(path, date1904)
        assert _grids(path, True, monkeypatch) == _grids(path, False, monkeypatch)


def test_fast_xlsx_reads_real_exports_the_same(monkeypatch):
    books = sorted((ROOT / "примеры").rglob("*.xlsx"))
    assert books
    for path in books:
        assert _grids(path, True, monkeypatch) == _grids(path, False, monkeypatch), path.name


def test_unusual_book_falls_back_to_openpyxl(tmp_path, monkeypatch):
    from pricing import xlsx_fast

    path = tmp_path / "книга.xlsx"
    _tricky_book(path)
    expected = _grids(path, False, monkeypatch)

    def broken(_path):
        raise xlsx_fast.XlsxFastError("непривычная книга")

    monkeypatch.setattr(xlsx_fast, "read_xlsx", broken)
    assert _grids(path, True, monkeypatch) == expected


# ---------------------------------------------------------------- индекс цен


def _old_index(src, opt, price_files, monitoring):
    """Прежний PriceIndex: один общий словарь, очистка выбросов по каждому ключу."""
    from collections import defaultdict

    offers = defaultdict(list)
    if opt.price_source in ("acp", "both"):
        seen = defaultdict(set)
        for item in src.acp.values():
            for brand, crosses in item.competitors.items():
                b = prices.brand_key(brand)
                for c in crosses:
                    key = (b, norm_article(c.article))
                    bucket = offers[key]
                    for o in c.offers:
                        if not o.price or opt.in_stock_only and o.qty is not None and o.qty <= 0:
                            continue
                        if (o.supplier, o.price) in seen[key]:
                            continue
                        seen[key].add((o.supplier, o.price))
                        bucket.append(R.PriceOffer(o.supplier, o.price, o.qty, "АЦПКД"))
    low = 0
    if opt.price_source in ("prices", "both"):
        part, _ = R._file_offer_part(price_files, monitoring, opt.monitoring_only, opt.in_stock_only)
        for key, vals in part.items():
            offers[key].extend(vals)
    result = {}
    for key, vals in offers.items():
        kept = R.without_anomalies(vals)
        low += len(vals) - len(kept)
        if kept:
            result[key] = kept
    return result, low


def _price_world(tmp_path):
    root = tmp_path / "Прайсы"
    rows = ["Бренд;Артикул;Цена;Наличие"]
    for i in range(60):
        rows.append(f"Brembo;B{i};{1000 + i * 10};{i % 3}")
        rows.append(f"Kashiyama;K{i % 20};{500 + i};1")
    rows += ["Brembo;B1;2;5", "Brembo;B2;50;5", "Brembo;B3;1010;0"]  # заглушка, выброс, дубль цены
    (root / "Armtek").mkdir(parents=True)
    (root / "Armtek" / "a.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    (root / "TISS").mkdir(parents=True)
    (root / "TISS" / "t.csv").write_text("Бренд;Артикул;Цена\nBrembo;B1;990\nKashiyama;K3;3\nBrembo;B7;1070\n", encoding="utf-8")
    src = scan([root], [root])
    files = prices.load_price_files(src.price_sheets, {}, {}, ["Armtek", "TISS"])
    acp = {}
    for i in range(0, 30, 3):
        item = AcpItem(str(i), f"OUR{i}", "", "ADVICS", "", "", "", "", "", "", "", None, None, None, None, None, None)
        item.competitors = {
            "BREMBO": [AcpCross(f"B{i}", [Offer("Rossko", 1000 + i * 10, 2), Offer("Berg", 5, 1)])],
            "KASHIYAMA": [AcpCross(f"K{i}", [Offer("Rossko", 0, 1)]), AcpCross("NOPE1", [])],
        }
        acp[str(i)] = item
    src = Sources({}, [])
    src.acp = acp
    return src, files


def test_shared_price_part_gives_the_same_prices_as_one_dictionary(tmp_path):
    src, files = _price_world(tmp_path)
    monitoring = {"BREMBO": {"armtek"}}
    for price_source in ("both", "prices", "acp"):
        for in_stock in (False, True):
            for only in (True, False):
                opt = R.Options(price_source=price_source, in_stock_only=in_stock, monitoring_only=only)
                index = R.PriceIndex(src, opt, files, monitoring)
                expected, low = _old_index(src, opt, files, monitoring)
                assert index.offers == expected
                assert index.skipped_low == low
                for brand, article in expected:
                    assert index.get(brand, article) == expected[(brand, article)]
                assert index.get("BREMBO", "NOPE") == [] and index.get("KASHIYAMA", "NOPE1") == []


def test_price_part_is_cached_by_the_files_not_by_the_list(tmp_path):
    src, files = _price_world(tmp_path)
    opt = R.Options()
    monitoring = {}
    first = R._file_offers_clean(list(files), monitoring, True, False)
    again = R._file_offers_clean(list(files), monitoring, True, False)  # новый список — те же файлы
    assert again is first
    other = R._file_offers_clean(list(files[:1]), monitoring, True, False)
    assert other is not first
    assert R.index_for(src, opt, list(files), monitoring) is R.index_for(src, opt, list(files), monitoring)


# ---------------------------------------------------------------- рабочая папка: прайсы общие, прошлые даты по требованию


def _workspace(tmp_path, monkeypatch, options=None):
    from pricing import dynamics, filecache, server, settings as settings_mod

    cache = tmp_path / "кэш"
    for mod in (filecache, prices, dynamics):
        monkeypatch.setattr(mod, "CACHE_DIR", cache)
    filecache._memory.clear()
    user = tmp_path / "настройки.json"
    user.write_text(json.dumps({"options": options or {}}, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(settings_mod, "USER_FILE", user)
    monkeypatch.setattr(server, "ROOT", tmp_path)
    server.STATE.workspace_cache = None
    server.STATE.price_shared = []
    incoming = tmp_path / "Входящие"
    incoming.mkdir(parents=True)
    (incoming / "пст.csv").write_text(
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа;Группа заказа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски;Диски\n"
        "2;BBB;Колодка;ADVICS;Тормозные колодки;Дисковые тормозные колодки;Колодки\n",
        encoding="utf-8",
    )
    (incoming / "сбс.csv").write_text("Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;AAA;5;1000\n2;BBB;1;100\n", encoding="utf-8")
    (incoming / "продажи.csv").write_text("Код;Артикул;Количество;Сумма продажи\n1;AAA;3;500\n2;BBB;2;200\n", encoding="utf-8")
    crosses = tmp_path / "Данные" / "Кроссы"
    crosses.mkdir(parents=True)
    (crosses / "кроссы.csv").write_text(
        "Артикул;Бренд;Артикул кросс;Бренд кросс;Использовать для сравнения цен\n"
        "AAA;ADVICS;B1;BREMBO;Да\nBBB;ADVICS;K1;AKEBONO;Да\n",
        encoding="utf-8",
    )
    day = tmp_path / "Данные" / "Прайсы"
    for date, price in (("29.09.2026", 900), ("06.10.2026", 1000)):
        folder = day / date / "TISS"
        folder.mkdir(parents=True)
        (folder / "tiss.csv").write_text(f"Бренд;Артикул;Цена\nBrembo;B1;{price}\nAkebono;K1;{price // 2}\n", encoding="utf-8")
    (day / "валюта.csv").write_text(";29.09.2026;06.10.2026\nUSD;90;95\n", encoding="utf-8")
    return server, user


def test_old_dates_are_read_only_when_dynamics_needs_them(tmp_path, monkeypatch):
    server, _ = _workspace(tmp_path, monkeypatch)
    scan_ = server.api_scan({"folder": str(tmp_path)})
    assert [f["date"] for f in scan_["price_files"]] == ["06.10.2026"]
    ws = server.STATE.workspace(tmp_path)
    assert [p.parent.parent.name for p in ws.shared.dormant_prices] == ["29.09.2026"]

    disks = server.api_analyze({"folder": str(tmp_path), "options": {"brand": "ADVICS", "group": "Тормозные диски"}})
    assert disks["price_date"] == "06.10.2026"
    row = disks["preview"][0]
    assert row["prices"]["BREMBO"] == 1000.0

    dyn = server.api_prices_dynamics({"folder": str(tmp_path)})
    assert dyn["date_from"] == "29.09.2026" and dyn["date_to"] == "06.10.2026"
    assert sorted(f["date"] for f in dyn["price_files"]) == ["06.10.2026", "29.09.2026"]
    assert not ws.shared.dormant_prices
    # Отчёт по-прежнему берёт только последнюю дату.
    again = server.api_analyze({"folder": str(tmp_path), "options": {"brand": "ADVICS", "group": "Тормозные диски"}})
    assert again["preview"][0]["prices"]["BREMBO"] == 1000.0


def test_all_files_of_the_last_date_off_falls_back_to_the_previous_date(tmp_path, monkeypatch):
    server, user = _workspace(tmp_path, monkeypatch)
    first = server.api_scan({"folder": str(tmp_path)})
    (latest,) = first["price_files"]
    data = json.loads(user.read_text(encoding="utf-8"))
    data["price_files"] = {latest["id"]: {"enabled": False}}
    user.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    disks = server.api_analyze({"folder": str(tmp_path), "options": {"brand": "ADVICS", "group": "Тормозные диски"}})
    assert disks["price_date"] == "29.09.2026"
    assert disks["preview"][0]["prices"]["BREMBO"] == 900.0


def test_groups_share_one_set_of_parsed_prices(tmp_path, monkeypatch):
    """У каждой группы своя копия источника, а разобранные прайсы — одни на всех."""
    import copy

    server, _ = _workspace(tmp_path, monkeypatch)
    server.api_scan({"folder": str(tmp_path)})
    settings = server.Settings.load()
    disks = server.STATE.for_group(tmp_path, "ADVICS", "Тормозные диски")
    other = copy.copy(disks)
    other.__dict__.pop("_price_cache", None)
    other.price_sheets = list(disks.price_sheets)
    assert server.STATE.price_files(other, settings) is server.STATE.price_files(disks, settings)
    a = server._report_prices(disks, settings)
    b = server._report_prices(other, settings)
    assert [id(pf) for pf in a] == [id(pf) for pf in b]


def test_columns_set_on_an_old_date_still_teach_the_new_date(tmp_path, monkeypatch):
    """Ручные столбцы поставщика, указанные на прошлой дате, применяются к новой — без чтения всех дат."""
    server, user = _workspace(tmp_path, monkeypatch)
    day = tmp_path / "Данные" / "Прайсы"
    for date in ("29.09.2026", "06.10.2026"):
        (day / date / "TISS" / "tiss.csv").unlink()
    (day / "29.09.2026" / "TISS" / "a.csv").write_text(
        "Прайс дня\nКонтрагент;Код;Сумма\nBrembo;B1;100\nAkebono;K1;200\n", encoding="utf-8")
    (day / "06.10.2026" / "TISS" / "b.csv").write_text(
        "Другой текст\nКонтрагент;Код;Сумма\nBrembo;B1;300\nAkebono;K1;400\n", encoding="utf-8")
    data = json.loads(user.read_text(encoding="utf-8"))
    data["price_files"] = {"29.09.2026 / TISS / a.csv": {
        "header_row": 1, "cols": {"brand": 0, "article": 1, "price": 2}, "supplier": "TISS"}}
    user.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    scan_ = server.api_scan({"folder": str(tmp_path)})
    (b,) = scan_["price_files"]
    assert b["date"] == "06.10.2026" and b["status"] == "saved" and b["offers"] == 2
    ws = server.STATE.workspace(tmp_path)
    assert ws.shared.dormant_prices  # прошлая дата целиком не читалась
