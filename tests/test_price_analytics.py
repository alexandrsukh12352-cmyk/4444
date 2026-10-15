"""Аналитика прайсов по датам и городам, выгрузка цен для макроса."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))

import openpyxl  # noqa: E402

from pricing import price_analytics  # noqa: E402

from test_speed import _workspace  # noqa: E402


def _three_dates(tmp_path, monkeypatch):
    server, user = _workspace(tmp_path, monkeypatch)
    monkeypatch.setattr(price_analytics, "CACHE_DIR", tmp_path / "кэш")
    price_analytics._PANELS.clear()
    price_analytics._ROWS_MEM.clear()
    day = tmp_path / "Данные" / "Прайсы" / "13.10.2026" / "TISS"
    day.mkdir(parents=True)
    # Екатеринбург поднял Brembo на 10%, Москва оставила; Akebono подешевел в обоих городах.
    (day / "TissEKB.csv").write_text("Бренд;Артикул;Цена;Кол-во\nBrembo;B1;1100;4\nAkebono;K1;450;2\nNoname;Z1;5;1\n", encoding="utf-8")
    (day / "TissMSK.csv").write_text("Бренд;Артикул;Цена;Кол-во\nBrembo;B1;1000;1\nAkebono;K1;400;3\n", encoding="utf-8")
    msk = tmp_path / "Данные" / "Прайсы" / "06.10.2026" / "TISS" / "TissMSK.csv"
    msk.write_text("Бренд;Артикул;Цена;Кол-во\nBrembo;B1;1000;1\nAkebono;K1;500;3\n", encoding="utf-8")
    for date in ("29.09.2026", "06.10.2026"):
        ekb = tmp_path / "Данные" / "Прайсы" / date / "TISS" / "tiss.csv"
        ekb.rename(ekb.with_name("TissEKB.csv"))
    return server


def test_changes_by_city_brand_and_group(tmp_path, monkeypatch):
    server = _three_dates(tmp_path, monkeypatch)
    server.api_scan({"folder": str(tmp_path)})
    a = server.api_prices_analytics({"folder": str(tmp_path)})
    assert a["dates"] == ["29.09.2026", "06.10.2026", "13.10.2026"]
    assert (a["date_from"], a["date_to"]) == ("06.10.2026", "13.10.2026")
    assert a["presets"] == {"prev": "06.10.2026", "week": "06.10.2026", "first": "29.09.2026"}
    brembo = next(b for b in a["brands"] if b["key"] == "BREMBO")
    assert (brembo["n"], brembo["up"], brembo["same"]) == (2, 1, 1)
    cell = next(w for w in a["where"] if w["key"] == "BREMBO" and w["city"] == "Екатеринбург")
    assert abs(cell["median"] - 0.10) < 1e-9
    disks = next(p for p in a["places"] if p["key"] == "BREMBO" and p["group"] == "Тормозные диски" and not p["subgroup"])
    assert disks["n"] == 2 and disks["category"] == "ДЕТАЛИ ТОРМОЗНОЙ СИСТЕМЫ" and disks["up"] == 1
    # Noname нет в позиционировании и в кроссах — его нет нигде.
    assert all(b["key"] != "NONAME" for b in a["brands"])
    # Индекс по всем датам: 29.09 → 06.10 Brembo +11,1% (900 → 1000), дальше медиана двух городов +5%.
    line = a["index"]["lines"]["BREMBO"]
    assert line[0] == 100 and abs(line[1] - 111.111) < 0.01 and abs(line[2] - 116.667) < 0.01
    assert a["rates"]["codes"] == ["USD"] and a["rates"]["values"]["USD"] == [90, 95, 95]
    first = server.api_prices_analytics({"folder": str(tmp_path), "date_from": "29.09.2026"})
    assert first["date_from"] == "29.09.2026" and first["summary"]["n"] == 2


def test_export_for_macro_has_one_row_per_price_and_date(tmp_path, monkeypatch):
    server = _three_dates(tmp_path, monkeypatch)
    server.api_scan({"folder": str(tmp_path)})
    r = server.api_prices_export({"folder": str(tmp_path)})
    wb = openpyxl.load_workbook(r["path"])
    assert wb.sheetnames == ["Цены", "Даты и файлы", "Курсы", "Проверка"]
    ws = wb["Цены"]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    # 29.09: 2 позиции; 06.10: 2 файла × 2; 13.10: 2 файла × 2 — Noname не из кроссов.
    assert len(rows) == r["rows"] == 10
    ekb = [x for x in rows if x["Город"] == "Екатеринбург" and x["Бренд"] == "Brembo"]
    last = max(ekb, key=lambda x: x["Дата"])
    assert last["Цена"] == 1100 and last["Цена на прошлую дату"] == 1000 and abs(last["Изменение"] - 0.1) < 1e-9
    assert last["Дата"].strftime("%d.%m.%Y") == "13.10.2026" and last["Прошлая дата"].strftime("%d.%m.%Y") == "06.10.2026"
    assert last["Наш артикул"] == "AAA" and last["Группа"] == "Тормозные диски" and last["Кросс «Да»"] == "да"
    assert last["Строка в файле"] == 2 and last["Наличие"] == 4 and last["Первая строка цены"] == 1
    check = wb["Проверка"]
    first = [c.value for c in check[2]]
    assert first[-1].startswith(("Brembo | B1", "Akebono | K1"))
    files = list(wb["Даты и файлы"].iter_rows(min_row=2, values_only=True))
    assert len(files) == 5
