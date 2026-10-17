"""Формулы книги, рекомендации, марки авто и все сочетания вида отчёта."""

import dataclasses
import re
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))
sys.path.insert(0, str(ROOT / "tests"))

from pricing import report as R  # noqa: E402
from pricing.cars import CarParser  # noqa: E402
from pricing.recommend import RecOptions, recommend, weight_classes  # noqa: E402
from pricing.settings import Settings  # noqa: E402
from pricing.sources import scan  # noqa: E402
from xlcalc import XLErr, cell  # noqa: E402

SAMPLES = ROOT / "примеры" / "выгрузки"
CROSSES = ROOT / "примеры" / "Кроссы свод (тестовая выгрузка).csv"
GROUP = dict(brand="ADVICS", group="Тормозные диски")


@pytest.fixture(scope="module")
def src():
    return scan([SAMPLES, CROSSES])


@pytest.fixture(scope="module")
def settings():
    return Settings.load(None)


def _head_row(ws):
    return next(r for r in range(1, 8) if any(c.value == "Артикул" for c in ws[r]))


def _header(ws):
    return {c.value: c.column for c in ws[_head_row(ws)] if c.value}


def _row(ws, code):
    for row in ws.iter_rows(min_row=_head_row(ws)):
        if row[0].value == code:
            return row[0].row
    raise AssertionError(code)


# ---------------------------------------------------------------- марки и рекомендации


def test_colors_stay_visible_in_excel_2016():
    """Excel 2016 читает альфа-канал: 00RRGGBB — прозрачный цвет, и оформление пропадает."""
    import zipfile
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule

    from pricing.workbook import CAT_A_FILL, GREEN, INK

    wb = Workbook()
    ws = wb.active
    ws["A1"].font = INK
    ws["A1"].fill = CAT_A_FILL
    ws.conditional_formatting.add("B1:B2", FormulaRule(formula=["B1<0"], fill=GREEN[0], font=GREEN[1]))
    buf = BytesIO()
    wb.save(buf)
    styles = zipfile.ZipFile(buf).read("xl/styles.xml").decode()
    assert "FF1F2937" in styles and "FFFFFBE0" in styles
    assert "FFC6EFCE" in styles and "FF006100" in styles
    assert "00C6EFCE" not in styles and "00006100" not in styles
    # как в файле, который Excel 2016 сохраняет сам: цвет условного формата в bgColor, без patternType
    assert '<patternFill><bgColor rgb="FFC6EFCE" /></patternFill>' in styles

    from pricing.workbook import _excel2016_stylesheet

    built = Path("/tmp/excel2016-flags.xlsx")
    wb.save(built)
    _excel2016_stylesheet(built)
    flagged = zipfile.ZipFile(built).read("xl/styles.xml").decode()
    assert 'fontId="1"' in flagged and 'applyFont="1"' in flagged
    assert 'applyFill="1"' in flagged


def test_sheet_xml_repairs_what_excel_deletes():
    from pricing.workbook import _repair_sheet_xml

    broken = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetViews><sheetView workbookViewId="0">'
        '<pane xSplit="4" ySplit="3" topLeftCell="E4" activePane="bottomRight" state="frozen" />'
        '<selection pane="topRight" />'
        '<selection pane="bottomLeft" />'
        '<selection pane="bottomRight" activeCell="A1" sqref="A1" />'
        '</sheetView></sheetViews>'
        '<sheetFormatPr baseColWidth="8" defaultRowHeight="15" />'
        '<cols><col hidden="1" outlineLevel="1" width="20" customWidth="1" min="5" max="5" />'
        '<col width="11" customWidth="1" min="1" max="1" /></cols>'
        '<sheetData><row r="2"><c r="A2"><f>SUM(A4:A9)</f><v /></c></row>'
        '<row r="3"><c r="A3"><f>SUM(A4:A9)</f><v></v></c></row>'
        '<row r="4"><c r="Z4"><f>IF(OR(AND($B$1="Скрыть",1)),"","да")</f></c></row>'
        '<row r="5"><c r="Z5" t="str"><f>IF(OR(AND($B$1="Скрыть",0)),"","да")</f><v>да</v></c></row>'
        '</sheetData>'
        '<autoFilter ref="A3:E9"><filterColumn colId="4" hiddenButton="1" showButton="1">'
        '<filters><filter val="да" /></filters></filterColumn></autoFilter>'
        '<dataValidations count="1"><dataValidation type="list" sqref="B1">'
        '<formula1>"Показать,Скрыть"</formula1></dataValidation></dataValidations>'
        '</worksheet>'
    )
    fixed = _repair_sheet_xml(broken)
    assert "<dataValidations" not in fixed
    assert "Показать,Скрыть" not in fixed
    assert "<v />" not in fixed and "<v></v>" not in fixed
    assert 'activeCell="A1"' not in fixed
    assert '<selection pane="bottomRight" activeCell="E4" sqref="E4"/>' in fixed
    assert '<selection pane="topRight" activeCell="E1" sqref="E1"/>' in fixed
    assert '<selection pane="bottomLeft" activeCell="A4" sqref="A4"/>' in fixed
    assert 'outlineLevelCol="1"' in fixed
    assert '<row r="5" hidden=' not in fixed
    assert fixed.startswith("<?xml")
    col = re.search(r'<col\b[^>]*outlineLevel="1"[^>]*/>', fixed).group(0)
    assert col.index("min=") < col.index("outlineLevel=")
    assert col.startswith("<col min=")


def test_car_parser_segments():
    p = CarParser()
    land = p.parse("(R) TOYOTA LAND CRUISER J200 (07-), LEXUS LX570 (07-)")
    assert land.segment == "Премиум"
    assert "Toyota Land Cruiser" in land.label and "Lexus LX" in land.label
    assert "Mazda Mazda" not in p.parse("MAZDA 3").label
    assert p.parse("HYUNDAI GETZ").segment == "Бюджетный"
    assert p.parse("LEXUS RX350").segment == "Премиум"
    prado = p.parse("TOYOTA LAND CRUISER PRADO 150")
    assert prado.segment == "Средний" and "Prado" in prado.label
    assert p.parse("SUZUKI SOLIO").fleet == "Японский (праворульный)"
    assert p.parse("LADA GRANTA").segment == "Бюджетный"
    assert p.parse("").segment == "Не определён"


def test_price_cut_is_limited_by_payback():
    opt = RecOptions(enabled=True, tolerance=3, down=10, min_vr=0, rounding=1, max_lift=20)
    # ВР 50%: снижение окупается ростом продаж на 20%, если оно не глубже 20 × 0,5 / 1,2 ≈ 8,3%.
    cut = recommend(100, 50, 50, "Средний", "tail", opt, "Brembo")
    assert cut.price == 92 and cut.kind == "down" and "окупаемостью" in cut.reason
    # ВР 10%: окупаемое снижение меньше коридора — цену не трогаем, объясняем почему.
    thin = recommend(100, 90, 50, "Средний", "tail", opt, "Brembo")
    assert thin.price == 100 and thin.kind == "keep" and "окупаемостью" in thin.reason
    # Повышение правило не трогает.
    assert recommend(100, 50, 200, "Средний", "tail", opt, "Brembo").price == 105


def test_recommend_respects_limits_and_margin():
    opt = RecOptions(enabled=True, tolerance=3, down=10, up={"Премиум": 10, "Бюджетный": 2, "Средний": 5}, min_vr=32, rounding=1,
                     max_lift=0)
    up = recommend(100, 50, 200, "Премиум", "tail", opt, "Brembo")
    assert up.price == 110 and up.kind == "up"
    assert recommend(100, 50, 200, "Бюджетный", "tail", opt, "Lada").price == 102
    down = recommend(100, 50, 50, "Средний", "tail", opt, "Brembo")
    assert down.price == 90 and down.kind == "down"
    # ниже пола ВР не опускаем: 80 / (1 − 0,32) ≈ 117,6, а лимит снижения — только до 90, поэтому цена остаётся
    held = recommend(100, 80, 50, "Средний", "tail", opt, "Brembo")
    assert held.price == 100 and held.kind == "keep"
    inside = recommend(100, 50, 102, "Средний", "tail", opt, "Brembo")
    assert inside.kind == "keep" and inside.price == 100
    # топ продаж: лимит +5% сжимается вдвое → +2,5% → 103 ₽
    top = recommend(100, 40, 200, "Средний", "top", opt, "Brembo")
    assert top.price == 103 and "топ" in top.reason
    # ВР ниже 32% и цели нет — поднимаем до пола, но не выше лимита сегмента
    raised = recommend(100, 80, None, "Премиум", "tail", opt, "")
    assert raised.kind == "vr" and raised.price == 110
    # оверсток и распродажа: вверх нельзя, даже если цель выше и ВР ниже пола
    stuck = recommend(100, 80, 200, "Премиум", "tail", opt, "Brembo", raise_cap=0)
    assert stuck.price == 100
    europe = recommend(100, 50, 200, "Премиум", "tail", opt, "Brembo", raise_cap=0.02)
    assert europe.price == 102
    # топ категории A дороже цели — не снижаем, даже если разрыв большой
    core = recommend(100, 40, 50, "Средний", "top", opt, "Brembo", abc="AA")
    assert core.price == 100 and core.kind == "keep"
    assert recommend(100, 40, 50, "Средний", "top", opt, "Brembo", abc="А").price == 100
    # прочий топ: общий лимит −10% × 0,5 дал бы −5%, держим не больше −2%
    other = recommend(100, 40, 50, "Средний", "top", opt, "Brembo", abc="B")
    assert other.price == 98 and other.kind == "down"


def test_weight_classes_by_accumulated_sales():
    rows = [type("R", (), {"code": c, "sales_qty": q, "abc": ""}) for c, q in (("a", 70), ("b", 16), ("c", 14), ("d", 0))]
    got = weight_classes(rows, RecOptions())
    assert got == {"a": "top", "b": "mid", "c": "tail", "d": "tail"}


def test_weight_classes_follow_abc_from_1c():
    rows = [type("R", (), {"code": c, "sales_qty": q, "abc": abc})
            for c, q, abc in (("a", 1, "AA"), ("b", 90, "А"), ("c", 5, "B"), ("d", 4, "C"), ("e", 0, ""))]
    assert weight_classes(rows) == {"a": "top", "b": "mid", "c": "tail", "d": "tail", "e": "tail"}


# ---------------------------------------------------------------- книга


def _build(tmp_path, src, settings, **kw):
    opt = R.Options(**GROUP, **kw)
    prep = R.prepare(src, settings, opt)
    result = R.build_report(src, settings, opt, tmp_path)
    wb = openpyxl.load_workbook(result.path)
    return opt, prep, result, wb


def _block(ws, title):
    for cell in ws["A"]:
        if cell.value == title:
            return cell.row
    raise AssertionError(title)


def test_full_report_formulas_match_python(tmp_path, src, settings):
    opt, prep, result, wb = _build(tmp_path, src, settings, rec={"enabled": True, "target": "main"})
    assert result.kind == "full"
    assert wb.sheetnames[:3] == ["Сводка", "Товары", "Цены конкурентов"] and "Дашборд" not in wb.sheetnames and "Параметры" not in wb.sheetnames
    assert {"ПСТ", "СБС", "Продажи"} <= set(wb.sheetnames)
    assert wb["ПСТ"].sheet_state == "hidden" and wb["Товары"].sheet_state == "visible"
    assert any(c.value == "Код товара" for row in wb["ПСТ"].iter_rows(max_row=4) for c in row)
    assert result.path.name == "ADVICS Тормозные диски.xlsx"
    brief = result.brief.read_text(encoding="utf-8-sig")
    assert result.brief.name == "Разбор.txt" and result.brief.parent == tmp_path
    assert result.path.parent.name == "Целевые группы"
    assert "Advics Тормозные диски" in brief and "===" not in brief

    ws = wb["Товары"]
    head = _header(ws)
    assert "Курс" not in head and "ГЗ" not in head and "Скидка, %" not in head
    assert head["Продажи, шт"] and head["Продажи, ₽"]
    n_fmt = ws.cell(_row(ws, 130909), head["Кат. А, ₽"]).number_format
    assert n_fmt.startswith("#,##0")
    assert ws.cell(_row(ws, 130909), head["Остаток, шт"]).number_format.startswith("#,##0")
    n = _row(ws, 130909)
    row = next(r for r in prep.kept if r.code == "130909")
    col = lambda title: openpyxl.utils.get_column_letter(head[title])  # noqa: E731

    def val(title):
        return cell(wb, ws.title, f"{col(title)}{n}")

    assert val("Закупка, ₽") == pytest.approx(row.purchase_rub)
    assert val("СБС план, ₽") == pytest.approx(row.sbs_plan)
    assert val("Кат. А, ₽") == pytest.approx(row.cat_a)
    assert val("ВР") == pytest.approx(1 - row.sbs / row.cat_a)
    assert "СБС план" in str(ws.cell(n, head["ВР"]).value) or head["СБС план, ₽"] and f'{openpyxl.utils.get_column_letter(head["СБС план, ₽"])}{n}' in str(ws.cell(n, head["ВР"]).value)
    assert val("Сумма продажи, ₽") == pytest.approx(row.sales_qty * row.cat_a)
    assert isinstance(ws.cell(n, head["Brembo, ₽"]).value, float)
    assert val("Brembo, ₽") == pytest.approx(row.prices["BREMBO"])
    assert "INDEX" in str(ws.cell(n, head["МО, ₽"]).value)
    assert "ПСТ" in str(ws.cell(n, head["МО, ₽"]).value)
    assert {"Цель, ₽", "По кому цель", "Откл. от цели", "Причина", "Подробно"} <= set(head)
    assert "В пути, шт" not in head and "Потенциал, шт" not in head
    assert val("Новая цена, ₽") == pytest.approx(row.rec.price)
    assert val("Изм.") == pytest.approx(row.rec.change)
    assert val("ВР после") == pytest.approx(1 - row.sbs / row.rec.price)
    assert "Основание" not in head and "Brembo" not in head
    assert wb["Товары"].cell(4, 1).fill.fill_type == "solid"
    for other in wb.worksheets:
        if other.title in ("Товары", "Списки", "Сводка", "Цены конкурентов", "Графики"):
            continue
        assert other.tables == {}
        assert list(other.conditional_formatting) == []
        assert other.sheet_view.showGridLines
        assert not other["A1"].font.bold
        assert other["A1"].fill.fill_type is None
    assert any(str(c.value).startswith("=IF") for c in ws[2] if isinstance(c.value, str))


def test_subgroup_sheets_and_modes(tmp_path, src, settings):
    half = len(src.pst) // 2
    pst = [dataclasses.replace(p, subgroup="Спортивные тормозные диски") if i < half else p for i, p in enumerate(src.pst)]
    split = dataclasses.replace(src, pst=pst)
    for kind in ("full",):
        for mode in ("sheets", "single"):
            for enabled in (True, False):
                opt = R.Options(**GROUP, report_kind=kind, subgroup_mode=mode, rec={"enabled": enabled, "target": "avg" if enabled else "main"})
                result = R.build_report(split, settings, opt, tmp_path / f"{kind}-{mode}-{enabled}")
                names = result.sheets
                assert "Дашборд" not in names
                assert "Спортивные тормозные диски" not in names
                assert "Товары" in names
                assert "Параметры" not in names
                assert "ПСТ" in names and "СБС" in names and "Продажи" in names
                assert (any(n.startswith("Цены ") for n in names) if mode == "sheets" else "Цены конкурентов" in names) and "Справочники" not in names and "Данные" not in names
                if enabled:
                    wb = openpyxl.load_workbook(result.path)
                    assert "Новая цена, ₽" in [c.value for c in wb["Товары"][4]]
                    assert wb["ПСТ"].sheet_state == "hidden"
                    assert wb["Товары"].sheet_state == "visible"


def test_competitor_price_is_a_number(tmp_path, src, settings):
    _, prep, _, wb = _build(tmp_path, src, settings)
    ws = wb["Товары"]
    n = _row(ws, 130909)
    head = _header(ws)
    row = next(r for r in prep.kept if r.code == "130909")
    assert ws.cell(n, head["Brembo, ₽"]).value == pytest.approx(row.prices["BREMBO"])
    assert ws.cell(n, head["Brembo, ₽"]).number_format.startswith("#,##0")
    assert "Кроссы" not in str(ws.cell(n, head["Brembo, ₽"]).value)


def test_summary_and_prices_sheets_link_to_the_report(tmp_path, src, settings):
    opt, prep, result, wb = _build(tmp_path, src, settings, rec={"enabled": True, "target": "main"})
    assert wb.active.title == "Сводка"
    itogi = wb["Сводка"]
    assert itogi.sheet_state == "visible" and wb["Цены конкурентов"].sheet_state == "visible"
    texts = [c.value for row in itogi.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("=") and "'Товары'!" in t for t in texts)
    assert "По подгруппам" in texts and "Итого" in texts
    assert not any("#REF" in t for t in texts)

    report = wb["Товары"]
    head = _header(report)
    assert {"Цель, ₽", "По кому цель", "Откл. от цели", "Причина", "Подробно"} <= set(head)
    row = next(r for r in prep.kept if r.target)
    n = _row(report, int(row.code) if row.code.isdigit() else row.code)
    assert report.cell(n, head["Цель, ₽"]).value == pytest.approx(round(row.target, 2))
    assert report.cell(n, head["Подробно"]).value == row.rec.reason

    prices = wb["Цены конкурентов"]
    phead = {c.value: c.column for c in prices[3] if c.value}
    assert list(phead)[:6] == ["Артикул", "Наименование", "Группа", "Подгруппа", "Разновидность", "ABC"]
    assert "Кат. А, ₽" in phead and "Brembo, ₽" in phead and "Откл." in phead
    m = next(c.row for c in prices["A"] if c.value == row.article)
    assert "'Товары'!" in prices.cell(m, phead["Кат. А, ₽"]).value


def test_cross_sheet_lists_selected_articles_and_supplier_prices(tmp_path, src, settings):
    _, prep, _, wb = _build(tmp_path, src, settings, price_pick="min")
    assert wb.sheetnames[3] == "Кроссы" and wb["Графики"].sheet_state == "hidden"
    cross = wb["Кроссы"]
    assert cross.sheet_state == "hidden"
    assert list(cross.merged_cells.ranges) == []
    assert cross.sheet_view.showGridLines
    row = next(r for r in prep.kept if r.code == "130909")
    brand_col = next(c.column for c in cross[1] if c.value == "Brembo")
    assert cross.cell(2, brand_col).value == "Кросс-артикул"
    headers = [c.value for c in cross[2]]
    assert "Мин" in headers
    articles = ", ".join(p.article for p in row.picks["BREMBO"] if p.article)
    data = next(r for r in cross.iter_rows(min_row=3) if r[0].value == row.article)
    assert data[brand_col - 1].value == articles
    min_col = next(i for i, value in enumerate(headers) if value == "Мин" and i >= brand_col - 1)
    assert str(data[min_col].value).startswith("=IF(COUNT(")
    assert cell(wb, "Кроссы", data[min_col].coordinate) == pytest.approx(row.prices["BREMBO"])
    report = wb["Товары"]
    assert isinstance(report.cell(_row(report, 130909), _header(report)["Brembo, ₽"]).value, float)


def test_sheet_filters_do_not_write_the_excel_row(tmp_path, src, settings):
    pst = [dataclasses.replace(p, overstock="да") if p.code == "130909" else p for p in src.pst]
    split = dataclasses.replace(src, pst=pst)
    opt = R.Options(**GROUP, filters={"overstock": "exclude"},
                    sheet_filters=["overstock", "sale", "competitor_price"])
    prep = R.prepare(split, settings, opt)
    assert all(r.code != "130909" for r in prep.kept)
    result = R.build_report(split, settings, opt, tmp_path)
    wb = openpyxl.load_workbook(result.path)
    ws = wb["Товары"]
    assert not any("Показ" in str(c.value or "") for r in range(1, 5) for c in ws[r])
    assert "Списки" not in wb.sheetnames
    assert ws.auto_filter.ref.startswith("A4:")
    assert list(ws.auto_filter.filterColumn) == []
    with pytest.raises(AssertionError):
        _row(ws, 130909)
    kept_code = next(r.code for r in prep.kept)
    shown = int(kept_code) if str(kept_code).isdigit() else kept_code
    assert _row(ws, shown)
    import zipfile
    packed = zipfile.ZipFile(result.path)
    sheets = [packed.read(name).decode() for name in packed.namelist() if name.startswith("xl/worksheets/sheet")]
    report_xml = next(xml for xml in sheets if "Кат. А" in xml and "Разновидность" in xml)
    assert "Показать,Скрыть" not in report_xml
    assert '="Скрыть"' not in report_xml
    assert 'activeCell="A1"' not in report_xml
    assert "outlineLevelCol=" in report_xml
    assert "<v />" not in report_xml and "<v></v>" not in report_xml
    assert "workbookProtection" not in packed.read("xl/workbook.xml").decode()
    pane = re.search(r'<pane\b[^>]*topLeftCell="([A-Z]+\d+)"', report_xml)
    assert pane
    top = pane.group(1)
    assert f'<selection pane="bottomRight" activeCell="{top}" sqref="{top}"' in report_xml
    styles = packed.read("xl/styles.xml").decode()
    first_font = re.search(r"<font>.*?</font>", styles).group(0)
    assert first_font.index("<sz") < first_font.index("<name")
    for col in re.findall(r'<col\b[^>]*outlineLevel="1"[^>]*/>', report_xml):
        assert col.index("min=") < col.index("outlineLevel=")
        assert col.index("max=") < col.index("outlineLevel=")


def test_no_sales_never_raises_price():
    opt = RecOptions(enabled=True, tolerance=3, down=10, min_vr=30, max_lift=20)
    # мы дешевле цели, продаж нет — дело не в цене
    cheap = recommend(100, 40, 120, "Средний", "tail", opt, "Brembo", sold=False, stock=True)
    assert cheap.price == 100 and cheap.kind == "keep" and "нет спроса" in cheap.reason
    # продаж нет, остатка нет — не трогаем
    gone = recommend(100, 40, 80, "Средний", "tail", opt, "Brembo", sold=False, stock=False)
    assert gone.price == 100 and gone.kind == "keep"
    # продаж нет, есть остаток, мы дороже цели — снижаем к цели, окупаемость не мешает
    dear = recommend(100, 40, 90, "Средний", "tail", opt, "Brembo", sold=False, stock=True)
    assert dear.price == 90 and dear.kind == "down" and "нет продаж" in dear.reason
    # нет цены конкурента и ВР ниже порога — без продаж не поднимаем
    blind = recommend(100, 80, None, "Средний", "tail", opt, sold=False, stock=True)
    assert blind.price == 100 and blind.kind == "keep"
    # новинка без продаж — как обычно: к цели вверх
    new = recommend(100, 40, 120, "Средний", "tail", opt, "Brembo", sold=False, stock=True, new=True)
    assert new.price > 100


def test_min_margin_does_not_lift_price_above_the_market():
    opt = RecOptions(enabled=True, tolerance=3, down=0, min_vr=40, max_lift=0)
    # ВР 20%, мы дороже цели: поднимать к мин. ВР — уходить от рынка дальше
    dear = recommend(100, 80, 95, "Средний", "top", opt, "Brembo", abc="AA")
    assert dear.price == 100 and "рынок ниже" in dear.reason
    # ВР 20%, цель выше нас на 3%: поднимаем, но не выше цели
    near = recommend(100, 80, 103, "Средний", "tail", RecOptions(enabled=True, tolerance=5, min_vr=40), "Brembo")
    assert near.price == 103 and near.kind == "vr"
    # цену недавно меняли — не трогаем
    frozen = recommend(100, 40, 130, "Средний", "tail", opt, "Brembo", frozen="цену меняли 01.10 — ждём до 30.12")
    assert frozen.price == 100 and frozen.kind == "keep" and "ждём" in frozen.reason


def test_recent_sales_file_is_told_apart_from_the_year():
    from pricing.excel import Sheet
    from pricing.sources import classify, is_recent_sales

    grid = [["Код", "Артикул", "Сумма продажи", "Количество"], ["1", "A1", 100, 2]]
    year = classify(Sheet(Path("Продажи 2025.xlsx"), "Лист1", grid))
    recent = classify(Sheet(Path("Продажи за 3 месяца.xlsx"), "Лист1", grid))
    assert year.kind == recent.kind == "sales"
    assert not is_recent_sales(year) and is_recent_sales(recent)
    head = [["Начало периода: 01.07.2026"], ["Конец периода: 30.09.2026"], *grid]
    assert is_recent_sales(classify(Sheet(Path("Продажи.xlsx"), "Лист1", head)))


def test_recent_sales_drive_the_no_sales_rule(tmp_path, settings):
    import shutil

    folder = tmp_path / "in"
    shutil.copytree(SAMPLES, folder)
    # та же выгрузка «Продажи», отдельным файлом с «3 месяца» в имени
    book = openpyxl.load_workbook(folder / "Advics диски.xlsx")
    sheet = next(ws for ws in book.worksheets if any("Сумма продажи" in str(c.value or "") for row in ws.iter_rows(max_row=12) for c in row))
    out = openpyxl.Workbook()
    for row in sheet.iter_rows(values_only=True):
        out.active.append(list(row))
    out.save(folder / "Продажи за 3 месяца.xlsx")
    src3 = scan([folder, CROSSES])
    assert src3.located("sales3") and src3.located("sales")
    assert src3.sales3 == src3.sales
    opt = R.Options(**GROUP, rec={"enabled": True})
    prep = R.prepare(src3, settings, opt)
    sold = [r for r in prep.kept if (r.sales or 0) > 0]
    assert sold and all(r.sales3 == r.sales for r in sold)
    quiet = [r for r in prep.kept if r.sales3 == 0 and r.rec and r.rec.kind in ("up", "vr") and not R.is_new(r)]
    assert quiet == []
    result = R.build_report(src3, settings, opt, tmp_path / "out")
    head = _header(openpyxl.load_workbook(result.path)["Товары"])
    assert "Продажи 3 мес, шт" in head or "Продажи 3 мес, ₽" in head


def test_short_reasons_for_the_report():
    from pricing.workbook import short_reason

    assert short_reason("к цели Brembo, ограничено +5%") == "к Brembo, шаг ограничен"
    assert short_reason("нет цены конкурента") == "нет цены конкурента"
    assert short_reason("топ продаж, категория A — не снижаем (Brembo)") == "топ — не снижаем"
    assert short_reason("в пределах позиционирования (Brembo)") == "у цели (Brembo)"
    assert short_reason("нет продаж — цену не поднимаем") == "нет продаж"
    assert short_reason("цену меняли 15.09.2026: +5,0%, по рекомендации — не трогаем до 14.12.2026, ждём") == "меняли 15.09, ждём до 14.12"
    assert short_reason("к цели Brembo; не ниже ВР 30%") == "к Brembo, мин. ВР"


def test_reason_columns_can_be_left_out_and_prices_split_by_subgroup(tmp_path, src, settings, monkeypatch):
    monkeypatch.setattr(R, "_subgroup_in_positioning", lambda *a, **k: True)
    half = len(src.pst) // 2
    pst = [dataclasses.replace(p, subgroup="Спортивные тормозные диски") if i < half else p for i, p in enumerate(src.pst)]
    split = dataclasses.replace(src, pst=pst)
    for cols, expect in (("all", {"Причина", "Подробно"}), ("short", {"Причина"}), ("none", set())):
        opt = R.Options(**GROUP, subgroup_mode="single", reason_cols=cols, rec={"enabled": True})
        wb = openpyxl.load_workbook(R.build_report(split, settings, opt, tmp_path / cols).path)
        head = set(_header(wb["Товары"]))
        assert {"Причина", "Подробно"} & head == expect, cols
        assert "Новая цена, ₽" in head
    opt = R.Options(**GROUP, subgroup_mode="sheets", rec={"enabled": True})
    names = openpyxl.load_workbook(R.build_report(split, settings, opt, tmp_path / "split").path).sheetnames
    prices = [n for n in names if n.startswith("Цены ")]
    assert len(prices) == 2 and "Цены конкурентов" not in names, names
    assert R.Options.from_dict({"reason_cols": "что-то"}).reason_cols == "all"


def test_short_reasons_for_sales_rules():
    from pricing.workbook import short_reason

    assert short_reason("продажи падают (−30% за 4 нед.) — не поднимаем, дело не в цене (Brembo)") == "падают −30% — не поднимаем"
    assert short_reason("продажи растут (+40% за 4 нед.) и так — не снижаем (Brembo)") == "растут +40% — не снижаем"
    assert short_reason("остатка на 40 нед. — не поднимаем (Brembo)") == "запас 40 нед. — не поднимаем"
    assert short_reason("остатка на 60 нед. (обычно до 50) — не поднимаем (Brembo)") == "запас 60 нед. — не поднимаем"
    assert short_reason("к цели Brembo, ограничено −6,0%; продажи падают (−25% за 4 нед.)") == "к Brembo, шаг ограничен, падают −25%"
