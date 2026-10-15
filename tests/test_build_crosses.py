import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))
sys.path.insert(0, str(ROOT / "tests"))

import build_crosses as bc  # noqa: E402
import make_sample_price  # noqa: E402

SETTINGS = bc.load_settings(ROOT / "настройки.ini")


def cross(art, yes):
    return bc.Cross(art, bc.norm_article(art), "X", yes, "")


def test_select_takes_only_yes_when_enough():
    picked = bc.select_crosses([cross("A1", True), cross("A2", False)], min_yes=1, max_no=0)
    assert [(c.article, s) for c, s in picked] == [("A1", bc.STATUS_YES)]


def test_select_falls_back_to_no_when_yes_missing():
    picked = bc.select_crosses([cross("A1", False), cross("A2", False)], min_yes=1, max_no=0)
    assert [s for _, s in picked] == [bc.STATUS_NO_FILL, bc.STATUS_NO_FILL]


def test_select_threshold_and_limit():
    available = [cross("Y1", True), cross("N1", False), cross("N2", False), cross("N3", False)]
    picked = bc.select_crosses(available, min_yes=2, max_no=2)
    assert [c.article for c, _ in picked] == ["Y1", "N1", "N2"]


def test_norm_article_handles_cyrillic_lookalikes_and_separators():
    assert bc.norm_article("ВR-0449 s") == bc.norm_article("BR0449S") == "BR0449S"


def test_positioning_resolves_by_subgroup_and_merges_case():
    pos = bc.builtin_positioning(SETTINGS)
    level, row, comps, _ = pos.resolve("avantech", "", "Дисковые тормозные колодки")
    assert level == "подгруппа"
    assert [c.key for c in comps] == ["KASHIYAMA", "LYNXAUTO", "G-BRAKE"]
    assert comps[0].percent == -15
    # «G-brake» и «G-BRAKE» в позиционировании — один бренд с одним написанием.
    _, _, discs, _ = pos.resolve("AVANTECH", "", "Стандартные тормозные диски")
    assert {c.name for c in discs if c.key == "G-BRAKE"} == {c.name for c in comps if c.key == "G-BRAKE"}


def test_positioning_reasons():
    pos = bc.builtin_positioning(SETTINGS)
    assert pos.resolve("NOBRAND", "", "x")[3] == "бренда нет в позиционировании"
    assert pos.resolve("AVANTECH", "", "Ремни ГРМ")[3] == "в позиционировании не указаны конкуренты"
    assert pos.resolve("AVANTECH", "", "Нет такой")[3].startswith("подгруппа не найдена")


def test_columns_bound_by_name_not_position(tmp_path):
    """Столбцы переставлены, есть лишние, заголовок не в первой строке, csv в cp1251 через «;»."""
    crosses = tmp_path / "кроссы.csv"
    crosses.write_bytes(
        (
            "Выгрузка от 01.10\n"
            "Лишний;Бренд кросс;Использовать для сравнения цен;Артикул кросс;Бренд;Артикул\n"
            "x;LYNXAUTO;Да;BD-1;AVANTECH;BR0001S\n"
            "x;LYNXAUTO;Нет;BD-2;AVANTECH;BR0001S\n"
            "x;PATRON;Нет;PBD1;AVANTECH;BR0001S\n"
            "x;PATRON;Нет;PBD2;AVANTECH;BR0001S\n"
            "x;MILES;Да;M1;AVANTECH;BR0001S\n"
        ).encode("cp1251")
    )
    price = tmp_path / "прайс.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Подгруппа", "Цена", "Торговая марка", "Артикул"])
    wb.active.append(["Стандартные тормозные диски", 100, "Avantech", "br-0001s"])
    wb.save(price)

    out, stats = bc.run([tmp_path], tmp_path / "out", SETTINGS)
    rows = list(openpyxl.load_workbook(out)["Для заноса"].iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    got = {(r[header.index("Бренд конкурент")], r[header.index("Артикул кросс")], r[header.index("Использовать для сравнения цен")]) for r in data}
    # LYNXAUTO: есть «Да» -> только он; PATRON: «Да» нет -> оба «Нет»; G-brake: нет кроссов; MILES не в позиционировании.
    assert got == {("LYNXAUTO", "BD-1", "Да"), ("PATRON", "PBD1", "Нет"), ("PATRON", "PBD2", "Нет")}
    assert stats["pair_не найден"] == 1
    assert data[0][header.index("Цена")] == 100


def test_end_to_end_on_sample(tmp_path):
    sample_price = make_sample_price.main(tmp_path / "in" / "прайс.xlsx")
    crosses = make_sample_price.CROSSES
    out, stats = bc.run([crosses, sample_price], tmp_path / "out", SETTINGS)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Сводка", "Для заноса", "По артикулам", "Не найдено", "Без позиционирования"]
    assert stats["cross_yes"] > 0 and stats["cross_no"] > 0
    assert stats["no_pos"] == 2

    # Файл результата, оставленный рядом со входными, не должен приниматься за кроссы.
    _, again = bc.run([crosses, sample_price, out], tmp_path / "out2", SETTINGS)
    assert again["cross_yes"] == stats["cross_yes"]
    assert any("файл результата" in w for w in again["warnings"])


def test_missing_crosses_file_is_user_error(tmp_path):
    (tmp_path / "пусто.csv").write_text("a;b\n1;2\n", encoding="utf-8")
    with pytest.raises(bc.UserError):
        bc.run([tmp_path], tmp_path / "out", SETTINGS)
