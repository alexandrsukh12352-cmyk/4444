import dataclasses
import json
import sys
import threading
import urllib.request
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))

from pricing import report as R  # noqa: E402
from pricing.settings import Settings  # noqa: E402
from pricing.sources import CrossRow, scan  # noqa: E402

SAMPLES = ROOT / "примеры" / "выгрузки"
CROSSES = ROOT / "примеры" / "Кроссы свод (тестовая выгрузка).csv"
GROUP = dict(brand="ADVICS", group="Тормозные диски")


@pytest.fixture(scope="module")
def src():
    return scan([SAMPLES, CROSSES])


@pytest.fixture(scope="module")
def settings():
    return Settings.load(None)


def rows_for(src, settings, **kw):
    opt = R.Options(**GROUP, **kw)
    rows = R.build_rows(src, settings, opt)
    kept, dropped = R.apply_filters(rows, opt)
    index = R.PriceIndex(src, opt)
    R.attach_crosses(src, kept, opt, index, R.suggestions_map(R.analyze_crosses(src, kept, opt, index)))
    return opt, rows, kept, dropped


def test_positioning_is_per_subgroup():
    import build_crosses as bc

    pos = bc.builtin_positioning(bc.load_settings(None))
    disks = {c.key: R._percent(c) for c in pos.resolve("ADVICS", "Тормозные диски", "Стандартные тормозные диски")[2]}
    pads = {c.key: R._percent(c) for c in pos.resolve("ADVICS", "Тормозные колодки", "Дисковые тормозные колодки")[2]}
    assert disks["BREMBO"] == 10 and pads["BREMBO"] == 20


def test_settings_positioning_overrides_builtin(src):
    s = Settings.load(None)
    s.data["positioning"] = [[
        "ADVICS", "Тормозные диски", "Стандартные тормозные диски", "Brembo", "-3", "Kashiyama", "5", "", "",
    ]]
    rows = R.build_rows(src, s, R.Options(**GROUP))
    r = next(x for x in rows if x.code == "130909")
    assert r.percents["BREMBO"] == -3 and r.percents["KASHIYAMA"] == 5


def test_clearance_is_withdrawal_or_assortment_r_and_new_is_not_abc(src, settings):
    withdrawn = dataclasses.replace(src, pst=[
        dataclasses.replace(p, discontinued=True, assortment="") if p.code == "130909" else p for p in src.pst
    ])
    opt = R.Options(**GROUP, filters={"sale": "exclude"})
    kept, _ = R.apply_filters(R.build_rows(withdrawn, settings, opt), opt)
    assert all(r.code != "130909" for r in kept)
    marked = dataclasses.replace(src, pst=[
        dataclasses.replace(p, discontinued=False, assortment="R", abc="A") if p.code == "130909" else p for p in src.pst
    ])
    kept, _ = R.apply_filters(R.build_rows(marked, settings, opt), opt)
    assert all(r.code != "130909" for r in kept)
    abc_d = dataclasses.replace(src, pst=[
        dataclasses.replace(p, discontinued=False, assortment="", abc="D") if p.code == "130909" else p for p in src.pst
    ])
    kept, _ = R.apply_filters(R.build_rows(abc_d, settings, opt), opt)
    assert all(r.code != "130909" for r in kept)
    third = dataclasses.replace(src, pst=[
        dataclasses.replace(p, order_group="Товары сторонних поставщиков") if p.code == "130909" else p for p in src.pst
    ])
    rows = R.build_rows(third, settings, R.Options(**GROUP))
    assert all(r.code != "130909" for r in rows)


def test_rebuilding_rows_does_not_keep_a_previous_edit(src, settings):
    opt = R.Options(**GROUP)
    first = R.build_rows(src, settings, opt)
    first[0].cat_a = -12345
    second = R.build_rows(src, settings, opt)
    assert second[0] is not first[0]
    assert second[0].cat_a != -12345
    assert [r.code for r in first] == [r.code for r in second]


def test_stock_slice_folder_keeps_target_groups_inside(tmp_path):
    from pricing.workbook import _report_path

    opt = R.Options.from_dict({"brand": "ADVICS", "group": "Тормозные диски", "stock_slice": "sale",
                               "filters": {"sale": "exclude", "overstock": "exclude"}})
    assert opt.filters["sale"] == "only" and "overstock" not in opt.filters
    path = _report_path(tmp_path, opt, 4)
    assert path.parts[-4:] == ("Распродажа и оверсток", "Распродажа", "Целевые группы", "ADVICS Тормозные диски.xlsx")
    over = R.Options.from_dict({"brand": "DOUBLE FORCE", "group": "Пружины подвески", "stock_slice": "overstock"})
    over_path = _report_path(tmp_path, over, 2)
    assert over_path.parts[-3:] == ("Оверсток", "Мелкие группы", "DOUBLE FORCE Пружины подвески.xlsx")


def test_result_bucket_is_target_small_or_no_price():
    assert R.result_bucket("ADVICS", "Тормозные диски", 12) == "target"
    assert R.result_bucket("AVANTECH", "Тормозные колодки", 3) == "target"
    assert R.result_bucket("DOUBLE FORCE", "Фильтры автомобильные", 8) == "target"
    assert R.result_folder_name("target") == "Целевые группы"
    assert R.result_bucket("ADVICS", "Тормозные диски", 0) == "noprice"
    assert R.result_bucket("ADVICS", "Тормозные суппорты", 4) == "small"
    assert R.result_bucket("DOUBLE FORCE", "Пружины подвески", 4) == "small"
    assert R.result_bucket("ADVICS", "Прочие товары", 4) == "small"
    assert R.result_folder_name("small") == "Мелкие группы"
    assert R.result_folder_name("noprice") == "Нет цен"


def test_only_positioning_groups_stay_in_the_list():
    from pricing.server import _only_positioning

    groups = [
        {"brand": "ADVICS", "group": "Тормозные диски"},
        {"brand": "ADVICS", "group": "Нет такой группы"},
        {"brand": "Другой", "group": "Тормозные диски"},
    ]
    allowed = {(R.norm_text("ADVICS"), R.norm_text("Тормозные диски"))}
    got = _only_positioning(groups, allowed)
    assert [(g["brand"], g["group"]) for g in got] == [("ADVICS", "Тормозные диски")]


def test_workspace_reads_incoming_until_group_folders_have_files():
    from pricing.sources import attach_shared, scan_workspace

    ws = scan_workspace(ROOT)
    assert any(p.sources.pst for p in ws.packs)
    pack = next(p for p in ws.packs if p.sources.pst)
    ready = attach_shared(pack.sources, ws.shared)
    assert ready.pst and (ready.crosses or pack.sources.crosses)


def test_top_seller_is_not_cut_toward_a_later_competitor():
    from types import SimpleNamespace

    comp = lambda key, name, order: SimpleNamespace(key=key, name=name, order=order)  # noqa: E731
    row = SimpleNamespace(
        code="1", sales_qty=80, abc="AA", cat_a=200, sbs=80, sbs_plan=None,
        overstock="", assortment="", car=None,
        competitors=[comp("BREMBO", "Brembo", 1), comp("GBRAKE", "G-brake", 3)],
        picks={"BREMBO": [], "GBRAKE": []},
        prices={"BREMBO": None, "GBRAKE": 100},
        percents={"BREMBO": 10, "GBRAKE": 35},
    )
    opt = R.Options(rec={"enabled": True, "target": "main", "down": 10, "tolerance": 3, "min_vr": 32})
    R.apply_recommendations([row], opt)
    assert row.weight == "top"
    assert row.target is None and row.rec.price == 200

    priced = SimpleNamespace(**row.__dict__)
    priced.prices = {"BREMBO": 100, "GBRAKE": 40}
    priced.abc = "B"
    priced.cat_a = 200
    R.apply_recommendations([priced], opt)
    assert priced.target_by == "Brembo"
    assert priced.target == pytest.approx(110)
    assert priced.rec.price == pytest.approx(196)


def test_competitor_without_a_price_does_not_get_a_column():
    from types import SimpleNamespace

    from pricing.report import competitor_order

    comp = lambda key, name, order: SimpleNamespace(key=key, name=name, order=order)  # noqa: E731
    row = SimpleNamespace(
        competitors=[comp("BREMBO", "Brembo", 1), comp("GBRAKE", "G-brake", 3)],
        picks={"BREMBO": [SimpleNamespace(article="X")], "GBRAKE": [SimpleNamespace(article="Y")]},
        prices={"BREMBO": 100.0, "GBRAKE": None},
    )
    assert [c.name for c in competitor_order([row])] == ["Brembo"]
    # Одна цена внизу отчёта, без продаж — столбец всё равно есть.
    sold = SimpleNamespace(
        competitors=[comp("BREMBO", "Brembo", 1), comp("KASHIYAMA", "Kashiyama", 2)],
        picks={"BREMBO": [SimpleNamespace(article="X")], "KASHIYAMA": [SimpleNamespace(article="Y")]},
        prices={"BREMBO": 100.0, "KASHIYAMA": None},
        sales_qty=10, sales=0,
    )
    idle = SimpleNamespace(
        competitors=sold.competitors,
        picks=sold.picks,
        prices={"BREMBO": None, "KASHIYAMA": 50.0},
        sales_qty=0, sales=0,
    )
    assert [c.name for c in competitor_order([sold, idle])] == ["Brembo", "Kashiyama"]


def test_missing_rates_cover_every_group():
    from types import SimpleNamespace

    from pricing.report import missing_rates_from_pst
    from pricing.settings import Settings

    settings = Settings.load(None)
    items = [
        SimpleNamespace(code="1", purchase_currency="KRW", purchase_price=100),
        SimpleNamespace(code="1", purchase_currency="KRW", purchase_price=100),
        SimpleNamespace(code="2", purchase_currency="CNY", purchase_price=5),
        SimpleNamespace(code="3", purchase_currency="руб.", purchase_price=10),
    ]
    got = {row["currency"]: row["items"] for row in missing_rates_from_pst(items, settings)}
    assert got.get("KRW") == 1 and got.get("CNY") == 1
    assert "руб." not in got


def test_recommendation_can_stick_to_the_first_competitor():
    from pricing.recommend import target_for

    order = ["BREMBO", "KASHIYAMA"]
    percents = {"BREMBO": 10, "KASHIYAMA": 20}
    no_brembo = {"BREMBO": None, "KASHIYAMA": 100}
    assert target_for(no_brembo, percents, order, "first") == (None, "")
    goal, who = target_for(no_brembo, percents, order, "main")
    assert who == "KASHIYAMA" and goal == pytest.approx(120)
    goal, who = target_for({"BREMBO": 100, "KASHIYAMA": 50}, percents, order, "first")
    assert who == "BREMBO" and goal == pytest.approx(110)


def test_placeholder_and_outlier_prices_are_not_used():
    from pricing.report import PriceOffer, without_anomalies

    junk = [
        PriceOffer("Форум-Авто", 2, 1, "АЦПКД"),
        PriceOffer("Форум-Авто", 15, 1, "АЦПКД"),
        PriceOffer("TISS", 3594.5, 4, "прайс"),
    ]
    assert [o.price for o in without_anomalies(junk)] == [3594.5]
    assert without_anomalies([PriceOffer("Форум-Авто", 2, None, "АЦПКД")]) == []
    assert without_anomalies([PriceOffer("KOITO", 47, 10, "прайс")])[0].price == 47
    both = without_anomalies([PriceOffer("A", 2000, None, "прайс"), PriceOffer("B", 3500, None, "прайс")])
    assert [o.price for o in both] == [2000, 3500]


def test_article_signs_are_stripped_before_match():
    from pricing.sources import norm_article

    assert norm_article("09.A417.11") == norm_article("09-A417-11") == norm_article("09A41711") == "09A41711"
    assert norm_article("GR-20751") == norm_article("GR20751") == "GR20751"
    assert norm_article("G20011_ПоврежУпак") == "G20011"
    assert norm_article("ВR-0449 s") == "BR0449S"


def test_price_article_matches_cross_without_signs(src, settings):
    from pricing.excel import Sheet
    from pricing.prices import Offer, PriceFile

    sheet = Sheet(Path("Signs.xlsx"), "Лист1", [])
    signed = PriceFile(sheet, "TISS", "", None, {}, "", "auto", offers=[
        Offer("BREMBO", "09.A417.11", 10.0, 5, 1, "09A41711"),
        Offer("G-BRAKE", "GR-20751", 20.0, 1, 2, "GR20751"),
    ])
    opt = R.Options(**GROUP, price_source="prices", monitoring_only=False)
    # 10 ₽ против диска за тысячи — отсев «в разы» убрал бы цену; здесь проверяем только сопоставление артикула.
    settings.options["screen"] = {"enabled": False}
    try:
        prep = R.prepare(src, settings, opt, [signed])
    finally:
        settings.options.pop("screen", None)
    row = next(x for x in prep.kept if x.code == "130909")
    assert row.prices["BREMBO"] == 10.0


def test_gz_file_is_recognised(src):
    assert src.located("gz") is not None
    assert src.gz["ПРОДУКЦИЯ ADVICS"] == 1.37


def test_values_match_template(src, settings):
    _, rows, _, _ = rows_for(src, settings)
    r = next(x for x in rows if x.code == "130909")
    # Шаблон ADVICS_тормозные_диски.xlsx, строка 3: J, R, S
    assert r.purchase_rub == pytest.approx(1572.819392)
    assert r.sbs_plan == pytest.approx(2154.76256704)
    assert r.cat_a == pytest.approx(4130.83)
    assert (r.discount, r.gz, r.abc, r.stock) == (15, 1.37, "AA", 4941)


def test_source_per_field(src, settings):
    _, rows, _, _ = rows_for(src, settings, sources={"stock": "acp", "sbs": "acp"})
    r = next(x for x in rows if x.code == "130909")
    assert (r.stock, r.sbs) == (63, 2207)


def test_acp_as_base_takes_pst_columns(src, settings):
    _, rows, _, _ = rows_for(src, settings, base="acp")
    assert len(rows) == len(src.acp)
    r = next(x for x in rows if x.code == "130909")
    assert (r.order_group, r.purchase_currency, r.cat_a) == ("Продукция ADVICS", "JPY", 4130.83)


def test_new_is_assortment_n_not_the_abc_category(src, settings):
    _, rows, kept, dropped = rows_for(src, settings, filters={"new": "exclude"})
    hidden = [r for r, reason in dropped if reason == "NEW"]
    assert hidden and all(r.assortment == "N" for r in hidden)
    assert all(r.assortment != "N" for r in kept)
    assert any(r.abc == "NEW" for r in kept)


def test_filters(src, settings):
    _, rows, kept, dropped = rows_for(src, settings, filters={"stock_transit": "exclude"})
    assert any(reason == "нет остатка и нет в пути" for _, reason in dropped)
    assert all((r.stock or 0) > 0 or (r.in_transit or 0) > 0 for r in kept)
    _, rows, kept, _ = rows_for(src, settings, filters=dict(R.CLEAN_PRESET))
    assert len(kept) < len(rows)
    assert not any(r.overstock or r.discontinued or (r.assortment or "").upper() == "R" or (r.abc or "").upper() in {"D", "Д"} or not r.mo or (not r.sbs and not r.sbs_plan) for r in kept)
    _, _, only, _ = rows_for(src, settings, filters={"overstock": "only"})
    assert only and all(r.overstock for r in only)
    _, _, aa, _ = rows_for(src, settings, abc=["AA"])
    assert aa and {r.abc for r in aa} == {"AA"}


def test_cross_modes():
    yes, no = CrossRow("Y1", "Y1", "B", True), CrossRow("N1", "N1", "B", False)

    class Index:
        def get(self, _brand, article):
            return [R.PriceOffer("S", 100.0, 1, "АЦПКД")] if article == "N1" else []

    pick = lambda mode, svod: [p.article for p in R.select(mode, svod, [], Index(), "B")]  # noqa: E731
    assert pick("yes", [yes, no]) == ["Y1"]
    assert pick("no", [yes, no]) == ["N1"]
    assert pick("yes_fill", [yes, no]) == ["Y1", "N1"]  # у «Да» нет цены — добираем «Нет»
    assert pick("yes_fill", [no]) == ["N1"]
    assert pick("yes_no", [yes, no]) == ["Y1", "N1"]
    assert pick("off", [yes, no]) == []


def test_cross_analytics_suggests_no_when_yes_is_scarce(src, settings):
    opt, _, kept, _ = rows_for(src, settings)
    stats = {c["key"]: c for c in R.analyze_crosses(src, kept, opt, R.PriceIndex(src, opt))}
    brembo, kashiyama = stats["BREMBO"]["segments"][0], stats["KASHIYAMA"]["segments"][0]
    assert brembo["seg"] == "*" and brembo["suggested"] == "yes"
    assert kashiyama["yes"] == 0 and kashiyama["suggested"] == "yes_fill"
    assert kashiyama["priced"]["yes_fill"] == 27


def test_cross_split_by_abc_and_modes_per_segment(src, settings):
    opt = R.Options(**GROUP, cross_split="abc", cross_modes={"BREMBO|*": "no", "BREMBO|abc=AA": "yes"})
    index = R.PriceIndex(src, opt)
    rows, _ = R.apply_filters(R.build_rows(src, settings, opt), opt)
    analysis = {c["key"]: c for c in R.analyze_crosses(src, rows, opt, index)}
    segs = {s["abc"]: s for s in analysis["BREMBO"]["segments"]}
    assert {"AA", "A", "B", "C"} <= set(segs)
    assert sum(s["items"] for s in segs.values()) == len(rows)
    assert segs["AA"]["mode"] == "yes" and segs["AA"]["explicit"]
    assert segs["C"]["mode"] == "no"  # «для всех» конкурента
    R.attach_crosses(src, rows, opt, index, R.suggestions_map(list(analysis.values())))
    aa = [r for r in rows if r.abc == "AA" and r.picks.get("BREMBO")]
    c = [r for r in rows if r.abc == "C" and r.picks.get("BREMBO")]
    assert aa and all(p.yes for r in aa for p in r.picks["BREMBO"])
    assert c and all(p.yes is False for r in c for p in r.picks["BREMBO"])


def test_minimal_price_and_source_filter(src, settings):
    from pricing.excel import Sheet
    from pricing.prices import Offer, PriceFile

    r = next(x for x in rows_for(src, settings)[2] if x.code == "130909")
    acp_price = r.prices["BREMBO"]
    sheet = Sheet(Path("Cheap.xlsx"), "Лист1", [])
    cheap = PriceFile(sheet, "Cheap", "", None, {}, "", "auto", offers=[Offer("BREMBO", "09.A417.11", 10.0, 5, 1)])
    # 10 ₽ рядом с настоящей ценой АЦПКД — заглушка и в минимум не идёт. Одна такая цена, без сравнения, остаётся.
    settings.options["screen"] = {"enabled": False}  # здесь проверяем минимум и источники, не отсев «в разы»
    try:
        for source, expected in [("acp", acp_price), ("prices", 10.0), ("both", acp_price)]:
            opt = R.Options(**GROUP, price_source=source, monitoring_only=False)
            prep = R.prepare(src, settings, opt, [cheap])
            row = next(x for x in prep.kept if x.code == "130909")
            assert row.prices["BREMBO"] == expected, source
    finally:
        settings.options.pop("screen", None)
    # С отсевом одинокая цена 10 ₽ против диска за тысячи в сравнение не идёт.
    prep = R.prepare(src, settings, R.Options(**GROUP, price_source="prices", monitoring_only=False), [cheap])
    row = next(x for x in prep.kept if x.code == "130909")
    assert row.prices["BREMBO"] is None and "дороже" in row.excluded["BREMBO"]["reason"]
    assert R.prepare(src, settings, R.Options(**GROUP, price_source="both", monitoring_only=False), [cheap]).index.skipped_low >= 1
    monitored = R.prepare(src, settings, R.Options(**GROUP, price_source="prices", monitoring_only=True), [cheap])
    assert monitored.index.skipped_by_monitoring == 1


def test_report_layout(tmp_path, src, settings):
    opt = R.Options(**GROUP, cross_modes={"BREMBO|*": "yes"})
    result = R.build_report(src, settings, opt, tmp_path)
    wb = openpyxl.load_workbook(result.path)
    assert wb.sheetnames[0] == "Отчет"
    assert "Дашборд" not in wb.sheetnames
    assert "Параметры" not in wb.sheetnames
    assert {"ПСТ", "СБС", "Продажи"} <= set(wb.sheetnames)
    assert wb["ПСТ"].sheet_state == "hidden" and wb["Отчет"].sheet_state == "visible"
    ws = wb["Отчет"]
    head = [c.value for c in ws[3]]
    assert head[:4] == ["Код товара", "Наименование русс.", "Артикул", "Разновидность"]
    variety_at = head.index("Разновидность") + 1
    freeze_at = openpyxl.utils.column_index_from_string("".join(ch for ch in ws.freeze_panes if ch.isalpha()))
    assert variety_at < freeze_at
    assert not ws.column_dimensions[openpyxl.utils.get_column_letter(variety_at)].outlineLevel
    for i, title in enumerate(head, 1):
        level = ws.column_dimensions[openpyxl.utils.get_column_letter(i)].outlineLevel or 0
        if level:
            assert i >= freeze_at, title
    assert "Brembo" not in head and "Основание" not in head
    assert "Курс" not in head and "ГЗ" not in head and "Скидка, %" not in head
    assert head.index("Разновидность") < head.index("Продажи, шт") < head.index("Кат. А")
    assert head.index("Сумма прибыли") < head.index("Brembo цена") < head.index("Откл Brembo")
    rows = [{h: c.value for h, c in zip(head, r)} for r in ws.iter_rows(min_row=3)]
    row = next(r for r in rows if r["Код товара"] == 130909)
    assert isinstance(row["Brembo цена"], float)
    assert str(row["Кат. А"]).startswith("=IF(")
    assert "ПСТ" in str(row["МО"])
    assert str(row["Откл Brembo"]).startswith("=IF(")
    assert any(c.value == "Код товара" for line in wb["ПСТ"].iter_rows(max_row=4) for c in line)


def test_subgroup_sheets(tmp_path, src, settings):
    half = len(src.pst) // 2
    pst = [dataclasses.replace(p, subgroup="Спортивные тормозные диски") if i < half else p for i, p in enumerate(src.pst)]
    split = dataclasses.replace(src, pst=pst)
    built = R.build_report(split, settings, R.Options(**GROUP, subgroup_mode="sheets"), tmp_path / "a")
    assert "Спортивные тормозные диски" not in built.sheets
    assert built.rows == len(src.pst) - half
    assert "Отчет" in built.sheets and "Дашборд" not in built.sheets
    single = R.build_report(split, settings, R.Options(**GROUP, subgroup_mode="single"), tmp_path / "b").sheets
    assert "Отчет" in single and "Спортивные тормозные диски" not in single


def test_settings_roundtrip(tmp_path):
    file = tmp_path / "настройки.json"
    s = Settings.load(file)
    assert s.gz["ПРОДУКЦИЯ ADVICS"] == 1.37
    s.data["gz"] = [["Продукция ADVICS", 1.5]]
    s.data["currency_rates"] = {"JPY": 0.6}
    s.save(file)
    again = Settings.load(file)
    assert again.gz == {"ПРОДУКЦИЯ ADVICS": 1.5} and again.rate("JPY") == 0.6
    assert again.discounts  # скидки не трогали — остались зашитые


def test_server_endpoints(tmp_path, monkeypatch):
    from pricing import server, settings as settings_mod

    monkeypatch.setattr(settings_mod, "USER_FILE", tmp_path / "настройки.json")
    httpd = server.serve("127.0.0.1", 47900)
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{httpd.server_address[1]}"
    try:
        assert b"<title>" in urllib.request.urlopen(base + "/").read()

        def post(path, body):
            req = urllib.request.Request(base + path, json.dumps(body).encode(), {"Content-Type": "application/json"})
            return json.loads(urllib.request.urlopen(req).read())

        scanned = post("/api/scan", {"folder": str(SAMPLES)})
        assert scanned["groups"]["pst"][0]["count"] == 547
        analysis = post("/api/analyze", {"folder": str(SAMPLES), "options": GROUP})
        assert analysis["kept"] == 547 and [c["name"] for c in analysis["competitors"]] == ["Brembo", "Kashiyama", "G-brake"]
        split = post("/api/analyze", {"folder": str(SAMPLES), "options": {**GROUP, "cross_split": "abc"}})
        assert len(split["crosses"][0]["segments"]) > 1
        init = post("/api/init", {})
        assert set(init["price_sources"]) == {"acp", "prices", "both"} and "subgroup_abc" in init["cross_splits"]
        assert init["report_kinds"]["quick"] == "Быстрый" and init["segments"]
        assert analysis["dashboard"]["items"] == 547 and analysis["dashboard"]["vr"]
        assert analysis["preview"][0]["segment"]
    finally:
        httpd.shutdown()
        httpd.server_close()
