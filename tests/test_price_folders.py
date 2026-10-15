"""Прайсы лежат в Данные/Прайсы/<Поставщик>/. Папка задаёт поставщика, лист 02 — ожидаемые бренды."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))

from pricing import prices, report as R, server  # noqa: E402
from pricing.sources import scan  # noqa: E402

CSV = "Бренд;Артикул;Цена\nBrembo;A1;1500\nKashiyama;Z9;900\nMANN;M1;400\n"


def _write(folder: Path, name: str, text: str = CSV) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_text(text, encoding="utf-8")


def test_supplier_comes_from_the_folder_and_city_from_the_file_name(tmp_path):
    root = tmp_path / "Данные" / "Прайсы"
    _write(root / "Armtek", "oct.csv")
    _write(root / "Rossko", "TissNSK.csv")
    _write(root / "Форум-Авто", "price.csv")
    _write(root, "TissEKB.csv")
    _write(root / "Armtek", "same.csv")
    _write(root / "Berg", "same.csv")
    src = scan([root], [root])
    positioning = R._builtin_positioning()
    known = R.known_suppliers(positioning)
    loaded = {prices.price_id(s): prices.load_one(s, {}, {}, known) for s in src.price_sheets}

    arm = next(pf for pf in loaded.values() if pf.folder == "Armtek" and pf.sheet.file.name == "oct.csv")
    assert arm.supplier == "Armtek" and arm.supplier_from == "folder" and arm.city == ""
    assert arm.status == "auto" and arm.offers and not arm.mismatch
    assert "Armtek /" in arm.id

    rossko = next(pf for pf in loaded.values() if pf.folder == "Rossko")
    assert (rossko.supplier, rossko.city, rossko.supplier_from) == ("Rossko", "Новосибирск", "folder")
    assert rossko.mismatch and rossko.name_supplier == "TISS"

    forum = next(pf for pf in loaded.values() if pf.folder == "Форум-Авто")
    assert forum.supplier == "Форум-Авто" and forum.supplier_from == "folder"

    loose = next(pf for pf in loaded.values() if pf.sheet.file.name == "TissEKB.csv")
    assert (loose.supplier, loose.city, loose.supplier_from, loose.folder) == ("TISS", "Екатеринбург", "file", "")
    assert loose.id == loose.sheet.label

    same = [pf for pf in loaded.values() if pf.sheet.file.name == "same.csv"]
    assert len(same) == 2 and len({pf.id for pf in same}) == 2

    described = server.describe_price(arm, positioning)
    assert "BREMBO" in described["expected_brands"] and "KASHIYAMA" in described["expected_brands"]
    assert described["missing_brands"] == []
    assert {b["brand"] for b in described["competitor_brands"]} >= {"BREMBO", "KASHIYAMA"}
    assert described["monitored"]

    book = R.monitoring_suppliers(positioning)
    assert any(row["supplier"] == "Armtek" and "BREMBO" in row["brands"] for row in book)
    assert any(row["supplier"] == "TISS" and "BREMBO" in row["brands"] for row in book)


def test_layout_text_shows_blank_columns_and_groups_the_same_shape(tmp_path):
    root = tmp_path / "Прайсы"
    shifted = "\n\n;;\n;;Бренд;Артикул;Цена\n;;Brembo;A1;100\n"
    _write(root / "Armtek", "a.csv", shifted)
    _write(root / "Armtek", "b.csv", shifted)
    src = scan([root], [root])
    known = R.known_suppliers(R._builtin_positioning())
    loaded = prices.load_price_files(src.price_sheets, {}, {}, known)
    assert len(loaded) == 2
    assert all(pf.header and pf.header.row == 3 for pf in loaded)
    text = prices.layout_text(loaded, "test")
    assert "Форма 1. Файлов: 2." in text
    assert "слева 2 пустых" in text
    assert "шапка 4" in text
    assert "· | · | Бренд | Артикул | Цена" in text
    assert "-> артикул" in text and "-> цена" in text
    assert text.count("\nФорма ") == 1


def test_layout_file_is_written_for_the_prices_tab(tmp_path, monkeypatch):
    from pricing import settings as settings_mod

    monkeypatch.setattr(settings_mod, "USER_FILE", tmp_path / "настройки.json")
    root = tmp_path / "Прайсы"
    _write(root / "Berg", "price.csv", "Бренд;Артикул;Цена\nMILES;Q1;100\n")
    monkeypatch.setattr(server, "resolve_folder", lambda raw, settings: root)
    monkeypatch.setattr(server, "_workspace_root", lambda folder: False)
    monkeypatch.setattr(server, "_report_dir", lambda settings, folder=None: tmp_path / "Результат")
    got = server.api_prices_layout({"folder": str(root)})
    text = (tmp_path / "Результат" / "Устройство прайсов.txt").read_text(encoding="utf-8-sig")
    assert got["name"] == "Устройство прайсов.txt"
    assert got["files"] == 1 and "MILES" in text and "шапка 1" in text


def test_price_column_is_numbers_and_name_is_not(tmp_path):
    root = tmp_path / "Прайсы"
    _write(root, "mix.csv", "Наименование;Артикул;Опт;Комментарий\nКолодки передние;A1;1500;есть\nФильтр масляный;B2;800;нет\n")
    src = scan([root], [root])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, [])
    assert pf.column_titles()["price"] == "Опт"
    assert pf.column_titles()["name"] == "Наименование"
    assert pf.checks["price"]["numeric"] == 1
    assert pf.checks["name"]["numeric"] == 0
    assert not pf.doubt


def test_text_column_called_price_is_not_taken(tmp_path):
    root = tmp_path / "Прайсы"
    _write(root, "bad.csv", "Артикул;Цена\nA1;колодки передние\nB2;фильтр масляный\n")
    src = scan([root], [root])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, [])
    assert pf.status == "none"
    assert "price" not in pf.cols


def test_manual_columns_apply_to_a_file_with_the_same_header(tmp_path):
    root = tmp_path / "Прайсы" / "Armtek"
    text = "Бренд;Артикул;Цена\nBrembo;A1;1500\n"
    _write(root, "a.csv", text)
    _write(root, "b.csv", text)
    src = scan([root.parent], [root.parent])
    known = R.known_suppliers(R._builtin_positioning())
    first = next(s for s in src.price_sheets if s.file.name == "a.csv")
    second = next(s for s in src.price_sheets if s.file.name == "b.csv")
    sample = prices.load_one(first, {}, {}, known)
    own = {prices.price_id(first): {"header_row": sample.header.row, "cols": dict(sample.cols)}}
    loaded = {pf.sheet.file.name: pf for pf in prices.load_price_files(src.price_sheets, {}, own, known)}
    assert loaded["a.csv"].status == "manual"
    assert loaded["b.csv"].status == "saved"
    assert loaded["b.csv"].cols["price"] == sample.cols["price"]
    assert loaded["b.csv"].supplier == "Armtek"


def test_cross_export_matches_our_cross_and_leaves_a_gap(tmp_path):
    import openpyxl

    from pricing.crossdump import write_cross_prices
    from pricing.sources import CrossRow, norm_article

    root = tmp_path / "Прайсы"
    _write(root, "t.csv", "Бренд;Артикул;Цена\nBrembo;09.A417.11;1500\n")
    src = scan([root], [root])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, ["TISS"])
    crosses = {
        ("ADVICS", "111"): {
            "BREMBO": [
                CrossRow("09.A417.11", norm_article("09.A417.11"), "Brembo", True),
                CrossRow("NOPE", "NOPE", "Brembo", True),
            ],
        },
    }
    path = tmp_path / "Кроссы и цены.xlsx"
    info = write_cross_prices(path, crosses, [pf], R._builtin_positioning())
    assert info["matched"] == 1 and info["missing"] == 1
    book = openpyxl.load_workbook(path, read_only=True)
    rows = list(book["Цены"].iter_rows(values_only=True))
    assert rows[1][2] == "Brembo" and rows[1][3] == "09.A417.11" and rows[1][7] == 1500
    gaps = list(book["Без цены"].iter_rows(values_only=True))
    assert gaps[1][3] == "NOPE"
    book.close()


def test_report_uses_only_the_latest_price_date():
    early = prices.PriceFile.__new__(prices.PriceFile)
    late = prices.PriceFile.__new__(prices.PriceFile)
    off = prices.PriceFile.__new__(prices.PriceFile)
    early.shot_date, early.enabled = "12.10.2026", True
    late.shot_date, late.enabled = "13.10.2026", True
    off.shot_date, off.enabled = "01.09.2026", False
    assert prices.latest_price_files([early, late, off]) == [late]


def test_header_on_the_second_row_wins_over_the_first(tmp_path):
    root = tmp_path / "Прайсы"
    _write(root, "row.csv", "Бренд;Артикул;Цена\nБренд;Артикул;Цена\nBrembo;A1;1500\nKashiyama;K1;800\n")
    src = scan([root], [root])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, ["Brembo"])
    assert pf.header.row == 1
    assert pf.offers and {o.brand for o in pf.offers} == {"Brembo", "Kashiyama"}


def test_brand_column_is_found_by_known_competitor_names(tmp_path):
    root = tmp_path / "Прайсы"
    _write(root, "brand.csv", "Контрагент;Артикул;Цена\nBrembo;A1;1500\nKashiyama;K1;800\nBrembo;A2;900\n")
    src = scan([root], [root])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, [], ["Brembo", "Kashiyama"])
    assert pf.column_titles()["brand"] == "Контрагент"
    assert {o.brand for o in pf.offers} == {"Brembo", "Kashiyama"}


def test_folder_outside_prices_does_not_become_the_supplier(tmp_path):
    folder = tmp_path / "входящие" / "Armtek"
    _write(folder, "TissEKB.csv")
    src = scan([folder])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, R.known_suppliers(R._builtin_positioning()))
    assert pf.supplier == "TISS" and pf.city == "Екатеринбург" and pf.supplier_from == "file"
    assert pf.folder == ""


def test_saved_price_settings_follow_the_folder_id(tmp_path, monkeypatch):
    from pricing import settings as settings_mod

    monkeypatch.setattr(settings_mod, "USER_FILE", tmp_path / "настройки.json")
    root = tmp_path / "Прайсы"
    _write(root / "Berg", "price.csv", "Бренд;Артикул;Цена\nMILES;Q1;100\n")
    monkeypatch.setattr(server, "resolve_folder", lambda raw, settings: root)
    monkeypatch.setattr(server, "_workspace_root", lambda folder: False)
    (f,) = server.api_scan({"folder": str(root)})["price_files"]
    assert f["supplier"] == "Berg" and f["supplier_from"] == "folder" and f["id"].startswith("Berg /")
    saved = server.api_prices_save({"folder": str(root), "id": f["id"], "city": "Омск"})
    assert saved["files"][0]["city"] == "Омск" and saved["files"][0]["supplier"] == "Berg"
    again = server.api_scan({"folder": str(root)})["price_files"][0]
    assert again["city"] == "Омск"
    monitoring = server.api_scan({"folder": str(root)})["monitoring_suppliers"]
    assert any(row["supplier"] == "Berg" and "MILES" in row["brands"] for row in monitoring)


def test_supplier_columns_apply_to_the_next_day_without_the_same_titles(tmp_path):
    root = tmp_path / "Прайсы"
    _write(root / "12.10.2026" / "TISS", "a.csv", "Прайс дня\nКонтрагент;Код;Сумма\nBrembo;A1;100\nKashiyama;K1;200\n")
    _write(root / "13.10.2026" / "TISS", "b.csv", "Другой текст\nКонтрагент;Код;Сумма\nBrembo;A2;300\nKashiyama;K2;400\nBrembo;A3;500\n")
    src = scan([root], [root])
    first = next(s for s in src.price_sheets if s.file.name == "a.csv")
    own = {prices.price_id(first): {"header_row": 1, "cols": {"brand": 0, "article": 1, "price": 2}, "supplier": "TISS"}}
    loaded = {pf.sheet.file.name: pf for pf in prices.load_price_files(src.price_sheets, {}, own, ["TISS"], ["Brembo", "Kashiyama"])}
    assert loaded["a.csv"].status == "manual" and loaded["a.csv"].offers
    assert loaded["b.csv"].status == "saved"
    assert {o.brand for o in loaded["b.csv"].offers} >= {"Brembo", "Kashiyama"}
    assert "Шапка не узнана" in loaded["b.csv"].why or loaded["b.csv"].why == ""


def test_parsed_prices_are_not_reread_when_the_file_is_unchanged(tmp_path, monkeypatch):
    root = tmp_path / "Прайсы"
    _write(root, "a.csv", "Бренд;Артикул;Цена\nBrembo;A1;100\nKashiyama;K1;200\n")
    src = scan([root], [root])
    prices.load_price_files(src.price_sheets, {}, {}, [])
    prices._OFFER_MEM.clear()
    calls = {"n": 0}
    original = prices.read_offers

    def wrapped(pf):
        calls["n"] += 1
        return original(pf)

    monkeypatch.setattr(prices, "read_offers", wrapped)
    again = prices.load_price_files(src.price_sheets, {}, {}, [])
    assert calls["n"] == 0
    assert again[0].offers


def test_unrecognized_file_explains_the_first_rows(tmp_path):
    root = tmp_path / "Прайсы" / "Новый"
    _write(root, "a.csv", "просто текст без столбцов\nещё строка\n")
    src = scan([root.parent], [root.parent])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, [])
    assert pf.status == "none"
    assert "Шапка не узнана" in pf.problem
    assert "строка 1" in pf.why
