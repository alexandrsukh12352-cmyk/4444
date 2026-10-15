import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))

from pricing import check  # noqa: E402
from pricing.excel import Sheet, code_key, find_header, num  # noqa: E402
from pricing.sources import discount_for, scan  # noqa: E402

SAMPLES = ROOT / "примеры" / "выгрузки"
CROSSES = ROOT / "примеры" / "Кроссы свод (тестовая выгрузка).csv"


@pytest.fixture(scope="module")
def src():
    return scan([SAMPLES, CROSSES])


def test_unchanged_file_is_not_read_from_excel_again(tmp_path, monkeypatch):
    import os

    from pricing import filecache
    from pricing.sources import scan

    monkeypatch.setattr(filecache, "CACHE_DIR", tmp_path / "кэш")
    filecache._memory.clear()
    folder = tmp_path / "in"
    folder.mkdir()
    target = folder / "пст.csv"
    target.write_text(PST_CSV, encoding="utf-8")
    first = scan([folder])
    assert [p.code for p in first.pst] == ["1", "2", "3"]
    assert filecache.misses >= 1
    filecache.reset_stats()
    filecache._memory.clear()
    second = scan([folder])
    assert [p.code for p in second.pst] == ["1", "2", "3"]
    assert filecache.misses == 0
    target.write_text(PST_CSV.replace("09.A417.11", "НОВЫЙ"), encoding="utf-8")
    os.utime(target, (target.stat().st_mtime + 5, target.stat().st_mtime + 5))
    filecache._memory.clear()
    third = scan([folder])
    assert third.pst[0].article == "НОВЫЙ"
    assert filecache.misses == 1

    from pricing.sources import scan_workspace

    saved = scan_workspace(tmp_path)
    signature = (("x", 1, 1),)
    filecache.save_workspace(signature, saved)
    loaded = filecache.load_workspace(signature)
    assert loaded is not None and loaded.company is not None
    assert filecache.load_workspace((("x", 2, 1),)) is None


def test_second_open_reuses_the_price_gap_cache(tmp_path, monkeypatch):
    incoming = tmp_path / "Входящие"
    incoming.mkdir(parents=True)
    (incoming / "пст.csv").write_text(
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски\n"
        "4;DDD;Суппорт;ADVICS;Тормозные суппорты;Направляющие суппорта\n",
        encoding="utf-8",
    )
    server, _user = _workspace(monkeypatch, tmp_path)
    monkeypatch.setattr(server.filecache, "CACHE_DIR", tmp_path / "кэш")
    server.filecache._memory.clear()
    calls = {"n": 0}
    real = server.report.group_price_gap

    def spy(*args, **kwargs):
        calls["n"] += 1
        return real(*args, **kwargs)

    monkeypatch.setattr(server.report, "group_price_gap", spy)
    server.api_scan({"folder": str(tmp_path)})
    first = calls["n"]
    assert first >= 1
    server.STATE.workspace_cache = None
    server.api_scan({"folder": str(tmp_path)})
    assert calls["n"] == first


def test_numbers_from_1c():
    assert num("4 571,29") == 4571.29
    assert num("1\xa0855\xa0007") == 1855007
    assert num("???") is None
    assert code_key("130 909") == code_key(130909.0) == "130909"


def test_header_with_title_rows_and_merged_levels():
    grid = [
        ["Отчёт ПСТ за период", None, None],
        ["Параметры: 01.01.2025 - 31.12.2025", None, None],
        ["Код товара", "Закупочная цена", "Закупочная цена"],
        ["Код товара", "Вал.", "Цена"],
        [130909, "JPY", 2848],
    ]
    h = find_header(Sheet(Path("x.xlsx"), "ПСТ", grid), [[r"^код товара$"], [r"^закупочная цена / цена$"]])
    assert (h.row, h.depth) == (2, 2)
    assert h.raw == ["Код товара", "Закупочная цена / Вал.", "Закупочная цена / Цена"]
    assert h.data == [[130909, "JPY", 2848]]


def test_every_document_is_recognised_by_content(src):
    assert set(src.found) >= {"pst", "sales", "sbs", "acp", "crosses", "discounts", "positioning"}
    assert not src.skipped


def test_pst_new_structure(src):
    item = next(p for p in src.pst if p.code == "130909")
    assert (item.article, item.brand, item.abc) == ("A6F164B", "ADVICS", "AA")
    assert item.mo_rub == 4859.8 and item.purchase_currency == "JPY" and item.purchase_price == 2848


def test_sales_sbs_and_discount(src):
    assert src.sales["130909"] == pytest.approx(62387933.25)
    stock, sbs = src.sbs["130909"]
    assert stock == 4941 and sbs == pytest.approx(10912627.97 / 4941)
    assert "ИТОГО" not in {k.upper() for k in src.sales}
    assert discount_for(src, "Стандартные тормозные диски", "Продукция ADVICS", "ADVICS") == 15


def test_acp_multi_row_product(src):
    item = src.acp["130909"]
    assert (item.cat_a_price, item.avg_cost, item.stock) == (4130.83, 2207, 63)
    brembo = {c.article: c for c in item.competitors["BREMBO"]}
    assert set(brembo) == {"09a41710", "09a41711"}
    assert brembo["09a41711"].min_price == 4048.33
    assert len(item.competitors["G-BRAKE"]) == 3


def test_check_report(tmp_path, src):
    out = tmp_path / "check.xlsx"
    stats = check.build_report(src, out)
    assert stats["items"] == 549
    assert openpyxl.load_workbook(out).sheetnames == ["Файлы", "Сопоставление", "Кроссы", "ПСТ и АЦПКД"]


TRANSIT_CSV = """Товар.Артикул;СНГ;;;Владивосток;;;Итого;;
;Заказано, шт.;Сумма в вал. пост.;Сумма, руб, тек. курс;Заказано, шт.;Сумма в вал. пост.;Сумма, руб, тек. курс;Заказано, шт.;Сумма в вал. пост.;Сумма, руб, тек. курс
0-03-004;;;;24;;;32;;
09-A417-11;;;;10;;;10;;
0-02-1228;;;;;;;8;;
Итого;1;2;3;4;5;6;7;8;9
"""

PST_CSV = """Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа
1;09.A417.11;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски
2;0-03-004;Диск 2;ADVICS;Тормозные диски;Стандартные тормозные диски
3;НЕТ-В-ПУТИ;Диск 3;ADVICS;Тормозные диски;Стандартные тормозные диски
"""


def test_transit_total_ignores_cities_and_matches_without_signs(tmp_path):
    from pricing.sources import norm_article

    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "путь.csv").write_text(TRANSIT_CSV, encoding="utf-8")
    (folder / "пст.csv").write_text(PST_CSV, encoding="utf-8")
    src = scan([folder])
    assert "transit" in src.found and not any(s.file.name == "путь.csv" for s in src.price_sheets)
    assert src.transit[norm_article("0-03-004")] == 32
    assert src.transit[norm_article("09.A417.11")] == 10
    assert src.transit[norm_article("0-02-1228")] == 8
    assert norm_article("Итого") not in src.transit

    from pricing import report as R
    from pricing.settings import Settings

    prep = R.prepare(src, Settings.load(None), R.Options(brand="ADVICS", group="Тормозные диски"))
    by_art = {r.article: r.in_transit for r in prep.rows}
    assert by_art["09.A417.11"] == 10
    assert by_art["0-03-004"] == 32
    assert by_art["НЕТ-В-ПУТИ"] is None

    result = R.build_report(src, Settings.load(None), R.Options(brand="ADVICS", group="Тормозные диски", report_kind="quick"), tmp_path)
    wb = openpyxl.load_workbook(result.path)
    head = [c.value for c in wb["Отчет"][3]]
    assert head[head.index("Остаток") + 1] == "В пути, шт"
    assert head[head.index("В пути, шт") + 1] == "Потенциал, шт"
    rows = {r[0].value: r for r in wb["Отчет"].iter_rows(min_row=3)}
    transit_col = head.index("В пути, шт")
    potential_col = head.index("Потенциал, шт")
    assert rows[2][transit_col].value == 32
    assert rows[3][transit_col].value is None
    assert str(rows[2][potential_col].value).startswith("=IF(COUNT(")
    sys.path.insert(0, str(ROOT / "tests"))
    from xlcalc import cell
    letter = openpyxl.utils.get_column_letter(potential_col + 1)
    assert cell(wb, "Отчет", f"{letter}{rows[2][0].row}") == 32
    assert "Параметры" not in wb.sheetnames
    assert wb["ПСТ"].sheet_state == "hidden"
    assert wb["ПСТ"]["A1"].fill.fill_type is None
    assert wb["Отчет"].cell(3, 1).fill.fill_type == "solid"


def test_incoming_base_replaces_the_same_document_in_the_group(tmp_path):
    from pricing.sources import apply_company, scan_workspace

    group = tmp_path / "Данные" / "Группы" / "Диски"
    incoming = tmp_path / "Входящие"
    group.mkdir(parents=True)
    incoming.mkdir()
    (group / "свой.csv").write_text(PST_CSV, encoding="utf-8")
    (group / "сбс.csv").write_text(
        "Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;09.A417.11;5;1000\n",
        encoding="utf-8",
    )
    (incoming / "общий.csv").write_text(
        PST_CSV + "9;ОБЩИЙ;Общий;ADVICS;Тормозные диски;Стандартные\n",
        encoding="utf-8",
    )
    ws = scan_workspace(tmp_path)
    assert [p.article for p in ws.company.pst] == ["09.A417.11", "0-03-004", "НЕТ-В-ПУТИ", "ОБЩИЙ"]
    ready = apply_company(ws.packs[0].sources, ws.company)
    assert [p.article for p in ready.pst] == ["09.A417.11", "0-03-004", "НЕТ-В-ПУТИ", "ОБЩИЙ"]
    assert ready.sbs["1"][0] == 5
    assert any("ПСТ" in w and "не используется" in w for w in ready.warnings)
    assert not any("СБС" in w and "не используется" in w for w in ready.warnings)


def test_one_transit_file_is_shared_with_every_group(tmp_path):
    from pricing.sources import attach_shared, norm_article, scan_workspace

    groups = tmp_path / "Данные" / "Группы"
    (groups / "Одни").mkdir(parents=True)
    (groups / "Другие").mkdir(parents=True)
    (groups / "Одни" / "пст.csv").write_text(PST_CSV, encoding="utf-8")
    (groups / "Другие" / "пст.csv").write_text(PST_CSV.replace("ADVICS", "OTHER"), encoding="utf-8")
    (groups / "Одни" / "путь.csv").write_text(TRANSIT_CSV, encoding="utf-8")
    ws = scan_workspace(tmp_path)
    other = next(p for p in ws.packs if p.name == "Другие")
    ready = attach_shared(other.sources, ws.shared)
    assert ready.transit[norm_article("0-03-004")] == 32
    assert ready.located("transit")


def test_finished_report_is_not_parsed_again(tmp_path):
    from pricing.report import Options, build_report
    from pricing.settings import Settings

    folder = tmp_path / "in"
    folder.mkdir()
    (folder / "пст.csv").write_text(PST_CSV, encoding="utf-8")
    src = scan([folder])
    build_report(src, Settings.load(None), Options(brand="ADVICS", group="Тормозные диски", report_kind="quick"), folder)
    report = next(folder.rglob("*.xlsx"))
    again = scan([folder, report])
    assert len(again.pst) == 3
    assert any("готовый отчёт" in s for s in again.skipped)


def test_build_all_keeps_only_groups_with_every_document(tmp_path, monkeypatch):
    import os

    from pricing import server, settings as settings_mod

    incoming = tmp_path / "Входящие"
    incoming.mkdir(parents=True)
    (incoming / "пст.csv").write_text(
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски\n"
        "2;BBB;Колодка;ADVICS;Тормозные колодки;Дисковые тормозные колодки\n"
        "3;CCC;Фильтр;MANN;Фильтры;Масляные\n",
        encoding="utf-8",
    )
    (incoming / "сбс.csv").write_text(
        "Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;AAA;5;1000\n2;BBB;1;100\n",
        encoding="utf-8",
    )
    (incoming / "продажи.csv").write_text(
        "Код;Артикул;Сумма продажи\n1;AAA;500\n2;BBB;200\n",
        encoding="utf-8",
    )
    (incoming / "диски.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "1;100;80;ADVICS;Тормозные диски;Стандартные тормозные диски;AAA\n",
        encoding="utf-8",
    )
    pads = incoming / "колодки.csv"
    pads.write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "2;90;70;ADVICS;Тормозные колодки;Дисковые тормозные колодки;BBB\n",
        encoding="utf-8",
    )
    os.utime(pads, (2_000_000_000, 2_000_000_000))
    user = tmp_path / "настройки.json"
    user.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(settings_mod, "USER_FILE", user)
    monkeypatch.setattr(server, "ROOT", tmp_path)

    scanned = server.api_scan({"folder": str(tmp_path)})
    got = {(g["brand"], g["group"]) for pack in scanned["batch"] for g in pack["groups"]}
    assert got == {("ADVICS", "Тормозные диски"), ("ADVICS", "Тормозные колодки")}

    made = server.api_build_all({"folder": str(tmp_path), "open": False, "options": {"report_kind": "quick", "base": "pst"}})
    assert {(m["brand"], m["group"]) for m in made["made"]} == got
    pads_src = server.STATE.for_group(tmp_path, "ADVICS", "Тормозные колодки")
    assert pads_src.acp["2"].mo_price == 90
    assert all("несколько файлов" not in w for w in pads_src.warnings)


def test_group_acp_beats_the_incoming_file_and_the_newer_slice_wins(tmp_path, monkeypatch):
    import os

    from pricing import server, settings as settings_mod

    incoming = tmp_path / "Входящие"
    group = tmp_path / "Данные" / "Группы" / "Диски"
    incoming.mkdir(parents=True)
    group.mkdir(parents=True)
    pst = (
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски\n"
    )
    (incoming / "пст.csv").write_text(pst, encoding="utf-8")
    (incoming / "сбс.csv").write_text(
        "Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;AAA;5;1000\n",
        encoding="utf-8",
    )
    (incoming / "продажи.csv").write_text("Код;Артикул;Сумма продажи\n1;AAA;500\n", encoding="utf-8")
    older = incoming / "старый.csv"
    newer = incoming / "новый.csv"
    acp = "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n1;{price};80;ADVICS;Тормозные диски;Стандартные тормозные диски;AAA\n"
    older.write_text(acp.format(price=10), encoding="utf-8")
    newer.write_text(acp.format(price=20), encoding="utf-8")
    os.utime(older, (1_000_000_000, 1_000_000_000))
    os.utime(newer, (2_000_000_000, 2_000_000_000))
    user = tmp_path / "настройки.json"
    user.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(settings_mod, "USER_FILE", user)
    monkeypatch.setattr(server, "ROOT", tmp_path)

    fresh = server.STATE.for_group(tmp_path, "ADVICS", "Тормозные диски")
    assert fresh.acp["1"].mo_price == 20
    assert any("несколько файлов" in w for w in fresh.warnings)

    (group / "свой.csv").write_text(acp.format(price=5), encoding="utf-8")
    server.STATE.workspace_cache = None
    local = server.STATE.for_group(tmp_path, "ADVICS", "Тормозные диски")
    assert local.acp["1"].mo_price == 5


def _workspace(monkeypatch, tmp_path, options=None):
    import json

    from pricing import server, settings as settings_mod

    user = tmp_path / "настройки.json"
    user.write_text(json.dumps({"options": options or {}}, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(settings_mod, "USER_FILE", user)
    monkeypatch.setattr(server, "ROOT", tmp_path)
    server.STATE.workspace_cache = None
    return server, user


def test_queue_marks_missing_documents_and_a_groups_own_acp(tmp_path, monkeypatch):
    incoming = tmp_path / "Входящие"
    group = tmp_path / "Данные" / "Группы" / "Диски"
    incoming.mkdir(parents=True)
    group.mkdir(parents=True)
    (incoming / "пст.csv").write_text(
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски\n"
        "5;SPORT;Спорт;ADVICS;Тормозные диски;Спортивные тормозные диски\n"
        "2;BBB;Колодка;ADVICS;Тормозные колодки;Дисковые тормозные колодки\n"
        "4;DDD;Суппорт;ADVICS;Тормозные суппорты;Направляющие суппорта\n"
        "3;CCC;Фильтр;MANN;Фильтры;Масляные\n",
        encoding="utf-8",
    )
    (incoming / "сбс.csv").write_text(
        "Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;AAA;5;1000\n2;BBB;1;100\n",
        encoding="utf-8",
    )
    (incoming / "продажи.csv").write_text("Код;Артикул;Сумма продажи\n1;AAA;500\n2;BBB;200\n", encoding="utf-8")
    (incoming / "колодки.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "2;90;70;ADVICS;Тормозные колодки;Дисковые тормозные колодки;BBB\n",
        encoding="utf-8",
    )
    (group / "свой.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "1;5;4;ADVICS;Тормозные диски;Стандартные тормозные диски;AAA\n",
        encoding="utf-8",
    )
    server, _user = _workspace(monkeypatch, tmp_path, {"rec_groups": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": {"tolerance": 11}}})
    scanned = server.api_scan({"folder": str(tmp_path)})
    rows = {(g["brand"], g["group"]): g for g in scanned["queue"]}
    assert set(rows) == {
        ("ADVICS", "Тормозные диски"),
        ("ADVICS", "Тормозные колодки"),
        ("ADVICS", "Тормозные суппорты"),
    }
    disks = rows[("ADVICS", "Тормозные диски")]
    assert disks["ready"] is True and disks["own"] is True and disks["count"] == 1 and disks["rec_own"] is True
    assert disks["reason"] == ""
    pads = rows[("ADVICS", "Тормозные колодки")]
    assert pads["ready"] is True and pads["own"] is False and pads["acp"] is True
    caliper = rows[("ADVICS", "Тормозные суппорты")]
    assert caliper["ready"] is False
    assert caliper["reason"] == "нет СБС, Продаж и АЦПКД"
    ready = {(g["brand"], g["group"]) for pack in scanned["batch"] for g in pack["groups"]}
    assert ready == {("ADVICS", "Тормозные диски"), ("ADVICS", "Тормозные колодки")}


def test_build_all_uses_each_groups_thresholds(tmp_path, monkeypatch):
    import json

    incoming = tmp_path / "Входящие"
    incoming.mkdir(parents=True)
    (incoming / "пст.csv").write_text(
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски\n"
        "2;BBB;Колодка;ADVICS;Тормозные колодки;Дисковые тормозные колодки\n",
        encoding="utf-8",
    )
    (incoming / "сбс.csv").write_text(
        "Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;AAA;5;1000\n2;BBB;1;100\n",
        encoding="utf-8",
    )
    (incoming / "продажи.csv").write_text("Код;Артикул;Сумма продажи\n1;AAA;500\n2;BBB;200\n", encoding="utf-8")
    (incoming / "диски.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "1;100;80;ADVICS;Тормозные диски;Стандартные тормозные диски;AAA\n",
        encoding="utf-8",
    )
    (incoming / "колодки.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "2;90;70;ADVICS;Тормозные колодки;Дисковые тормозные колодки;BBB\n",
        encoding="utf-8",
    )
    server, user = _workspace(monkeypatch, tmp_path, {"rec": {"enabled": False, "tolerance": 3, "min_vr": 32, "target": "main"}})
    captured = []
    real = server.report.build_report

    def spy(src, settings, opt, out_dir, files):
        captured.append((opt.brand, opt.group, opt.recs.tolerance, opt.recs.min_vr, opt.recs.enabled, opt.recs.target))
        return real(src, settings, opt, out_dir, files)

    monkeypatch.setattr(server.report, "build_report", spy)
    server.api_build_all({
        "folder": str(tmp_path),
        "open": False,
        "options": {
            "report_kind": "quick",
            "base": "pst",
            "rec": {"enabled": True, "target": "first", "tolerance": 99},
            "rec_groups": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": {"tolerance": 11, "min_vr": 40}},
        },
    })
    by = {(b, g): (tol, vr, en, tg) for b, g, tol, vr, en, tg in captured}
    assert by[("ADVICS", "Тормозные диски")] == (11, 40, True, "first")
    assert by[("ADVICS", "Тормозные колодки")] == (3, 32, True, "first")
    saved = json.loads(user.read_text(encoding="utf-8"))
    assert saved["options"]["rec"]["tolerance"] == 3
    assert saved["options"]["rec"]["min_vr"] == 32
    assert saved["options"]["rec"]["enabled"] is True
    assert saved["options"]["rec"]["target"] == "first"
    assert saved["options"]["rec_groups"]["ADVICS\tТОРМОЗНЫЕ ДИСКИ"]["tolerance"] == 11


def test_rec_sets_stay_separate_for_all_groups_and_for_one(tmp_path, monkeypatch):
    server, _user = _workspace(monkeypatch, tmp_path, {
        "rec_scope": "group",
        "rec_mine": {"tolerance": 4},
        "rec_source": "mine",
        "rec_groups": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": {"tolerance": 8}},
        "rec_group_source": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": "mine"},
    })
    settings = server.Settings.load()
    yours = server.rec_for_group(settings, "ADVICS", "Тормозные диски", {
        "rec_scope": "group",
        "rec_group_source": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": "yours"},
        "rec": {"enabled": True, "target": "first"},
    })
    assert yours["tolerance"] == 3 and yours["min_vr"] == 34 and yours["enabled"] is True and yours["target"] == "first"
    mine = server.rec_for_group(settings, "ADVICS", "Тормозные диски", {
        "rec_scope": "group",
        "rec_groups": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": {"tolerance": 8}},
        "rec_group_source": {"ADVICS\tТОРМОЗНЫЕ ДИСКИ": "mine"},
    })
    assert mine["tolerance"] == 8 and mine["down"] == 6
    one = server.rec_for_group(settings, "ADVICS", "Тормозные колодки", {
        "rec_scope": "all", "rec_source": "mine", "rec_mine": {"tolerance": 4},
    })
    assert one["tolerance"] == 4
    one_yours = server.rec_for_group(settings, "ADVICS", "Тормозные колодки", {
        "rec_scope": "all", "rec_source": "yours",
    })
    assert one_yours["tolerance"] == 3


def test_build_all_writes_a_brief_for_every_positioned_group(tmp_path, monkeypatch):
    incoming = tmp_path / "Входящие"
    incoming.mkdir(parents=True)
    (incoming / "пст.csv").write_text(
        "Код товара;Артикул;Наименование;Торговая марка;Группа;Подгруппа\n"
        "1;AAA;Диск;ADVICS;Тормозные диски;Стандартные тормозные диски\n"
        "2;BBB;Колодка;ADVICS;Тормозные колодки;Дисковые тормозные колодки\n"
        "4;DDD;Суппорт;ADVICS;Тормозные суппорты;Направляющие суппорта\n",
        encoding="utf-8",
    )
    (incoming / "сбс.csv").write_text(
        "Код;Артикул;Кол-во, кон. остаток;Себестоимость\n1;AAA;5;1000\n2;BBB;1;100\n",
        encoding="utf-8",
    )
    (incoming / "продажи.csv").write_text("Код;Артикул;Сумма продажи\n1;AAA;500\n2;BBB;200\n", encoding="utf-8")
    (incoming / "диски.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "1;100;80;ADVICS;Тормозные диски;Стандартные тормозные диски;AAA\n",
        encoding="utf-8",
    )
    (incoming / "колодки.csv").write_text(
        "Код;Цена МО;Цена со скидкой кат. А;Бренд;Группа;Подгруппа;Артикул\n"
        "2;90;70;ADVICS;Тормозные колодки;Дисковые тормозные колодки;BBB\n",
        encoding="utf-8",
    )
    server, _user = _workspace(monkeypatch, tmp_path)
    real = server.report.build_report

    def spy(src, settings, opt, out_dir, files):
        if opt.group == "Тормозные колодки":
            raise RuntimeError("сломалась книга")
        return real(src, settings, opt, out_dir, files)

    monkeypatch.setattr(server.report, "build_report", spy)
    result = server.api_build_all({"folder": str(tmp_path), "open": False, "options": {"report_kind": "quick", "base": "pst"}})
    made = {(m["brand"], m["group"]) for m in result["made"]}
    skipped = {(s["brand"], s["group"]) for s in result["skipped"]}
    assert made == {("ADVICS", "Тормозные диски")}
    assert ("ADVICS", "Тормозные колодки") in skipped
    assert ("ADVICS", "Тормозные суппорты") in skipped
    text = (tmp_path / "Результат" / "Разбор.txt").read_text(encoding="utf-8-sig")
    for title in ("Advics Тормозные диски", "Advics Тормозные колодки", "Advics Тормозные суппорты"):
        assert text.count(title) == 1
        body = text.split(title, 1)[1]
        nxt = body.find("\nAdvics ")
        block = body if nxt < 0 else body[:nxt]
        assert "Вывод:" in block


def test_uploaded_transit_export_uses_the_total_column():
    from pricing.sources import norm_article

    path = Path("/home/ubuntu/.cursor/projects/workspace/uploads")
    files = list(path.glob("*.csv")) if path.is_dir() else []
    if not files:
        pytest.skip("нет выгрузки товара в пути")
    src = scan(files)
    assert src.transit[norm_article("00010/10140")] == 960
    assert src.transit[norm_article("0-03-004")] == 32
    assert src.transit[norm_article("0-02-1228")] == 8
    assert src.transit[norm_article("0101-006")] == 2
