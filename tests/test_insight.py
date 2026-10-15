"""Аналитика прайсов и распознавание: отсев несопоставимого, ОПТ, синонимы брендов, «файл поехал»,
воронка, дыры в кроссах, история, готовый XML листов-выгрузок."""
import io
import json
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "Программа"))

from pricing import insight, price_check, prices  # noqa: E402
from pricing.sources import scan  # noqa: E402

from test_speed import _workspace  # noqa: E402


# ---------------------------------------------------------------- отсев несопоставимого


def _row(article, name, cat_a, prices_by):
    offers = {k: SimpleNamespace(supplier="TISS", price=p, qty=1) for k, p in prices_by.items()}
    return SimpleNamespace(article=article, name=name, cat_a=cat_a, prices=dict(prices_by), best=offers, excluded={})


def _settings(**screen):
    return SimpleNamespace(options={"screen": screen} if screen else {})


def test_kits_and_prices_apart_by_times_are_not_compared():
    kit = _row("D6F141B+D6F141B+SN889", "Комплект колодки + диски", 10917, {"KASHIYAMA": 1226})
    far = _row("A1", "Диск", 3300, {"BREMBO": 200250, "GBRAKE": 3000})
    cheap = _row("A2", "Диск", 9000, {"BREMBO": 2000})
    fine = _row("A3", "Диск", 4000, {"BREMBO": 4600})
    reasons = insight.screen_pairs([kit, far, cheap, fine], _settings())
    assert kit.prices["KASHIYAMA"] is None and "комплект" in kit.excluded["KASHIYAMA"]["reason"]
    assert far.prices["BREMBO"] is None and far.prices["GBRAKE"] == 3000
    assert cheap.prices["BREMBO"] is None and "дороже" in cheap.excluded["BREMBO"]["reason"]
    assert fine.prices["BREMBO"] == 4600 and not fine.excluded
    assert reasons == {"комплект против одной детали": 1, "мы дешевле в разы": 1, "мы дороже в разы": 1}


def test_screening_can_be_switched_off_and_the_ratio_changed():
    kit = _row("A+B", "Комплект", 1000, {"X": 100})
    assert not insight.screen_pairs([kit], _settings(enabled=False)) and kit.prices["X"] == 100
    near = _row("A1", "Диск", 3300, {"BREMBO": 1500})
    insight.screen_pairs([near], _settings(ratio=2))
    assert near.prices["BREMBO"] is None
    near = _row("A1", "Диск", 3300, {"BREMBO": 1500})
    insight.screen_pairs([near], _settings(ratio=3))
    assert near.prices["BREMBO"] == 1500


# ---------------------------------------------------------------- распознавание прайсов


def _price(tmp_path, name, text):
    base = tmp_path / name.split(".")[0] / "Прайсы"
    root = base / "TISS"
    root.mkdir(parents=True, exist_ok=True)
    (root / name).write_text(text, encoding="utf-8")
    src = scan([base], [base])
    (pf,) = prices.load_price_files(src.price_sheets, {}, {}, ["TISS"])
    return pf


def test_wholesale_price_wins_over_retail(tmp_path):
    pf = _price(tmp_path, "a.csv", "Бренд;Артикул;Цена розн.;ОПТ\nBrembo;A1;1500;1000\nBrembo;A2;2500;2000\n")
    assert pf.column_titles()["price"] == "ОПТ"
    assert [o.price for o in pf.offers] == [1000, 2000]
    pf = _price(tmp_path, "b.csv", "Бренд;Артикул;Цена;Цена опт\nBrembo;A1;1500;1000\nBrembo;A2;2500;2000\n")
    assert pf.column_titles()["price"] == "Цена опт"


def test_brand_aliases_and_hints(tmp_path):
    pf = _price(tmp_path, "a.csv", "Бренд;Артикул;Цена\nKashiyama Japan;K1;100\nKashiyama Japan;K2;200\nBrembo;B1;300\n")
    hints = price_check.brand_hints([pf], ["Kashiyama", "Brembo"])
    assert [(h["brand"], h["suggest"]) for h in hints] == [("Kashiyama Japan", "Kashiyama")]
    try:
        prices.set_brand_aliases([["Kashiyama Japan", "KASHIYAMA"]])
        assert prices.brand_key("Kashiyama Japan") == "KASHIYAMA" and prices.alias_token()
        assert price_check.brand_hints([pf], ["Kashiyama", "Brembo"]) == []
    finally:
        prices.set_brand_aliases([])
    assert prices.brand_key("Kashiyama Japan") == "KASHIYAMAJAPAN" and prices.alias_token() == ""


def test_a_file_that_moved_is_noted_against_the_previous_date(tmp_path, monkeypatch):
    monkeypatch.setattr(prices, "CACHE_DIR", tmp_path / "кэш")
    root = tmp_path / "Прайсы"
    rows = "\n".join(f"Brembo;B{i};{1000 + i}" for i in range(40))
    for date, body in (("29.09.2026", rows), ("06.10.2026", "\n".join(f"Akebono;A{i};{500 + i}" for i in range(40)))):
        (root / date / "TISS").mkdir(parents=True)
        (root / date / "TISS" / "t.csv").write_text("Бренд;Артикул;Цена\n" + body + "\n", encoding="utf-8")
    src = scan([root], [root])
    files = prices.load_price_files(src.price_sheets, {}, {}, ["TISS"])
    notes = price_check.quality_notes(files, ["Brembo", "Akebono"])
    new = next(pf for pf in files if pf.shot_date == "06.10.2026")
    assert any("пропал бренд BREMBO" in n for n in notes[new.id])
    old = next(pf for pf in files if pf.shot_date == "29.09.2026")
    assert old.id not in notes
    # Сводка прошлой даты сохранилась — сравнение работает и без чтения её прайса.
    again = price_check.quality_notes([new], ["Brembo", "Akebono"])
    assert again[new.id] == notes[new.id]


# ---------------------------------------------------------------- воронка, дыры, история


def test_funnel_names_the_missing_price_list_and_gaps_book_is_written(tmp_path, monkeypatch):
    server, _ = _workspace(tmp_path, monkeypatch)
    monkeypatch.setattr(insight, "_history_path", lambda: tmp_path / "история.json")
    server.api_scan({"folder": str(tmp_path)})
    pads = {"brand": "ADVICS", "group": "Тормозные колодки"}
    a = server.api_analyze({"folder": str(tmp_path), "options": pads})
    akebono = next(f for f in a["funnel"] if f["key"] == "AKEBONO")
    # Akebono мониторят у Berg и Ivers Auto, а цена есть только в прайсе TISS.
    assert (akebono["yes"], akebono["in_prices"], akebono["monitored"], akebono["final"]) == (1, 1, 0, 0)
    assert akebono["missing_suppliers"] == ["Berg", "Ivers Auto"]
    assert any(t.startswith("Нет прайса Berg, Ivers Auto — у 1 товара") for t in akebono["actions"])
    brembo = next(f for f in a["funnel"] if f["key"] == "BREMBO")
    assert brembo["yes"] == 0 and any("Нет кроссов «Да» на Brembo — 1 товар" in t for t in brembo["actions"])

    g = server.api_crosses_gaps({"folder": str(tmp_path), "options": pads, "open": False})
    import openpyxl

    wb = openpyxl.load_workbook(g["path"])
    assert wb.sheetnames == ["Нет кросса Да", "Да без цены", "Несопоставимо"]
    assert g["no_yes"] == 1 and g["yes_no_price"] == 1
    reason = wb["Да без цены"].cell(3, 9).value
    assert reason.startswith("цена есть только у поставщиков вне позиционирования")

    disks = {"brand": "ADVICS", "group": "Тормозные диски"}
    d = server.api_analyze({"folder": str(tmp_path), "options": disks})
    brembo = next(f for f in d["funnel"] if f["key"] == "BREMBO")
    assert brembo["final"] == 1 and brembo["sold_share"] == 1
    assert brembo["reliability"]["leaders"] == [{"supplier": "TISS", "rows": 1}]
    server.api_build({"folder": str(tmp_path), "options": disks, "open": False})
    hist = server.api_analyze({"folder": str(tmp_path), "options": disks})["history"]
    assert len(hist) == 1 and hist[0]["priced"] == 1


def test_history_keeps_one_record_a_day(tmp_path):
    path = tmp_path / "история.json"
    summary = {"items": 10, "priced": 4, "with_target": 4, "in_band": 2, "vr": 0.4,
               "slices": [{"kind": "all", "median": 0.05, "weighted": 0.04}]}
    rows = [SimpleNamespace(sales_qty=5, target=100), SimpleNamespace(sales_qty=5, target=None)]
    insight.record_history("ADVICS", "Тормозные диски", rows, summary, path=path)
    insight.record_history("ADVICS", "Тормозные диски", rows, {**summary, "priced": 5}, path=path)
    insight.record_history("ADVICS", "Тормозные диски", rows, summary, "sale", path=path)
    got = insight.history_for("advics", "тормозные диски", path=path)
    assert len(got) == 1 and got[0]["priced"] == 5 and got[0]["sold_share"] == 0.5
    assert len(json.loads(path.read_text(encoding="utf-8"))) == 2


# ---------------------------------------------------------------- листы-выгрузки готовым XML


def test_export_sheets_from_template_are_byte_for_byte_the_same(tmp_path, monkeypatch):
    from pricing import workbook

    server, _ = _workspace(tmp_path, monkeypatch)
    server.api_scan({"folder": str(tmp_path)})
    opts = {"brand": "ADVICS", "group": "Тормозные диски"}

    def members(path):
        z = zipfile.ZipFile(io.BytesIO(Path(path).read_bytes()))
        return {i.filename: z.read(i.filename) for i in z.infolist() if not i.filename.startswith("docProps")}

    def build():
        return members(server.api_build({"folder": str(tmp_path), "options": opts, "open": False})["path"])

    saved = workbook.raw_template
    workbook.raw_template = lambda grid: None
    try:
        old = build()
    finally:
        workbook.raw_template = saved
    first = build()
    again = build()
    assert first == old and again == old


def test_direct_xml_of_export_sheets_equals_openpyxl():
    import datetime as dt

    from pricing import workbook
    from pricing.excel import read_sheets

    grids = [s.grid for p in sorted((ROOT / "примеры").rglob("*")) if p.suffix in (".xlsx", ".xls", ".csv")
             for s in read_sheets(p)]
    grids.append([["Код", None, "a & <b> \"q\"", " lead", "x\r\ny", "   "], [130909, 1.5, True, -0.0, 1e20, 10**17],
                  [], [None, "", "tail ", 3, 2.25, False, "ё ✓"], [12345678901234567890, 1e-7, -5]])
    grids += [[], [[None, ""]]]
    checked = 0
    for grid in grids:
        fast = workbook._raw_template_fast(grid)
        if fast is None:  # даты или формулы — такой лист пишет сам openpyxl
            continue
        slow = workbook._raw_template_openpyxl(grid)
        assert (fast.sheet_data, fast.dimension, fast.probes) == (slow.sheet_data, slow.dimension, slow.probes)
        checked += 1
    assert checked >= 10
    for odd in ([[dt.datetime(2026, 1, 1)]], [["=1+1"]], [["#N/A"]], [[float("nan")]]):
        assert workbook._raw_template_fast(odd) is None
