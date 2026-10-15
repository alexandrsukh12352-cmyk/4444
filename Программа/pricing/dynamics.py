"""Динамика цен конкурентов по папкам-датам и файл «валюта».

Прайсы лежат так: Данные/Прайсы/12.10.2026/Armtek/файл.xlsx.
Рядом, прямо в «Прайсы», файл «валюта»: в шапке даты, в строках валюта и курс.

Повторный запуск не читает Excel и не собирает индексы заново, если дата и размер
файлов те же. Сетки листов по-прежнему в своём кэше. Здесь хранится только
сжатый индекс цен и уже посчитанное сравнение двух дат.
"""

from __future__ import annotations

import hashlib
import pickle
import re
from collections import defaultdict
from pathlib import Path

from .excel import clean, num, read_sheets
from .filecache import CACHE_DIR
from .prices import PriceFile, alias_token, brand_key, price_date, supplier_key
from .report import MIN_PRICE
from .sources import Sources, _is_currency_file, norm_article, norm_text

DYN_VERSION = 3
_DATE_IN_CELL = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$|^(\d{4})-(\d{2})-(\d{2})$")
_MEM_INDEX: dict[str, dict] = {}
_MEM_RESULT: dict[str, dict] = {}


def _slot(kind: str, token: str) -> Path:
    raw = f"{DYN_VERSION}|{kind}|{token}".encode()
    return CACHE_DIR / f"dyn-{hashlib.sha1(raw).hexdigest()}.pkl"


def _load(kind: str, token: str):
    hit = _MEM_INDEX.get(kind + token) if kind == "index" else _MEM_RESULT.get(token)
    if hit is not None:
        return hit
    path = _slot(kind, token)
    if not path.is_file():
        return None
    try:
        with path.open("rb") as fh:
            blob = pickle.load(fh)
    except (OSError, pickle.PickleError, EOFError, ValueError):
        return None
    if not isinstance(blob, dict) or blob.get("version") != DYN_VERSION or blob.get("token") != token:
        return None
    data = blob.get("data")
    if kind == "index":
        _MEM_INDEX[kind + token] = data
    else:
        _MEM_RESULT[token] = data
    return data


def _save(kind: str, token: str, data) -> None:
    if kind == "index":
        _MEM_INDEX[kind + token] = data
    else:
        _MEM_RESULT[token] = data
    try:
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        tmp = _slot(kind, token).with_suffix(".tmp")
        with tmp.open("wb") as fh:
            pickle.dump({"version": DYN_VERSION, "token": token, "data": data}, fh, protocol=pickle.HIGHEST_PROTOCOL)
        tmp.replace(_slot(kind, token))
    except (OSError, pickle.PickleError, TypeError, ValueError):
        pass


def _file_token(path: Path) -> str:
    st = path.stat()
    return f"{path.resolve()}|{st.st_mtime_ns}|{st.st_size}"


def prices_dir(folder: Path) -> Path | None:
    if folder.name.casefold() in {"прайсы", "prices"} and folder.is_dir():
        return folder
    for candidate in (folder / "Данные" / "Прайсы", folder / "Прайсы"):
        if candidate.is_dir():
            return candidate
    return None


def find_currency_file(folder: Path) -> Path | None:
    root = prices_dir(folder)
    if root is None:
        return None
    found = [p for p in root.rglob("*") if p.is_file() and _is_currency_file(p)]
    if not found:
        return None
    found.sort(key=lambda p: (len(p.parts), str(p)))
    return found[0]


def _cell_date(value) -> str:
    text = clean(value)
    m = _DATE_IN_CELL.match(text)
    if not m:
        return ""
    if m.group(1):
        return text
    return f"{m.group(6)}.{m.group(5)}.{m.group(4)}"


def _rate(value) -> float | None:
    if isinstance(value, str):
        value = value.replace(" ", "").replace(",", ".")
    return num(value)


def parse_currency_grid(grid: list) -> dict:
    """Шапка — даты, первый столбец — код валюты. «89,15» тоже число."""
    header = None
    for row in grid[:15]:
        dates = [_cell_date(v) for v in row[1:]]
        if sum(bool(d) for d in dates) >= 1:
            header = dates
            break
    if not header:
        return {"dates": [], "rates": {}}
    rates: dict[str, dict[str, float]] = {}
    for row in grid:
        code = clean(row[0]).upper() if row else ""
        if not code or _cell_date(code) or code in {"ВАЛЮТА", "CURRENCY"}:
            continue
        if not re.fullmatch(r"[A-Z]{3}", code):
            continue
        bucket = {}
        for i, date in enumerate(header, start=1):
            if not date or i >= len(row):
                continue
            value = _rate(row[i])
            if value:
                bucket[date] = value
        if bucket:
            rates[code] = bucket
    dates = []
    for date in header:
        if date and date not in dates:
            dates.append(date)
    return {"dates": dates, "rates": rates}


def read_currency(folder: Path) -> dict:
    path = find_currency_file(folder)
    empty = {"path": "", "dates": [], "rates": {}, "cached": False}
    if path is None:
        return empty
    token = _file_token(path)
    cached = _load("index", "cur|" + token)
    if isinstance(cached, dict) and "rates" in cached:
        cached = dict(cached)
        cached["path"] = str(path)
        cached["cached"] = True
        return cached
    try:
        sheets = read_sheets(path)
    except (OSError, ValueError):
        return empty
    parsed = {"dates": [], "rates": {}}
    for sheet in sheets:
        parsed = parse_currency_grid(sheet.grid)
        if parsed["rates"]:
            break
    data = {"path": str(path), "dates": parsed["dates"], "rates": parsed["rates"], "cached": False}
    _save("index", "cur|" + token, {"dates": parsed["dates"], "rates": parsed["rates"]})
    return data


def _offer_token(pf: PriceFile) -> str:
    cols = ",".join(f"{k}:{pf.cols[k]}" for k in sorted(pf.cols))
    token = _file_token(pf.sheet.file) + f"|{pf.header.row if pf.header else ''}|{cols}|{pf.supplier}|{pf.shot_date}"
    aliases = alias_token()
    return f"{token}|{aliases}" if aliases else token


def offer_index(pf: PriceFile) -> dict[tuple[str, str], float]:
    """Минимальная цена артикула конкурента в этом файле. Из кэша, если файл не менялся."""
    if not pf.offers or not pf.enabled:
        return {}
    token = _offer_token(pf)
    cached = _load("index", "px|" + token)
    if isinstance(cached, dict):
        return cached
    index: dict[tuple[str, str], float] = {}
    for offer in pf.offers:
        if not offer.price or offer.price < MIN_PRICE:
            continue
        key = (brand_key(offer.brand), offer.article_key or norm_article(offer.article))
        if not key[0] or not key[1]:
            continue
        prev = index.get(key)
        if prev is None or offer.price < prev:
            index[key] = offer.price
    _save("index", "px|" + token, index)
    return index


def _date_key(text: str) -> tuple:
    parts = clean(text).split(".")
    if len(parts) == 3 and all(p.isdigit() for p in parts):
        dd, mm, yy = parts
        return (int(yy), int(mm), int(dd), text)
    return (9999, 99, 99, text)


def _conclusions(rows: list[dict], rates: list[dict]) -> list[str]:
    brands = [row for row in rows if row.get("brand") and not row.get("category") and not row.get("group") and not row.get("subgroup")]
    brands.sort(key=lambda row: abs(row.get("change") or 0), reverse=True)
    usable = [rate for rate in rates if rate.get("change") is not None]
    lines = []
    for row in brands:
        change = row.get("change") or 0
        if change > 0.005:
            direction = "подорожал"
        elif change < -0.005:
            direction = "подешевел"
        else:
            direction = "почти не изменился"
        line = f"{row['brand']} {direction} на {abs(change):.1%} по {row['n']} артикулам."
        if usable:
            closest = min(usable, key=lambda rate: abs(change - rate["change"]))
            gap = change - closest["change"]
            line += f" Ближе всего курс {closest['code']} ({closest['change']:+.1%}), цена {'выше' if gap > 0 else 'ниже'} него на {abs(gap):.1%}."
        lines.append(line)
    if not lines:
        lines.append("Нет общих артикулов брендов из позиционирования в выбранных датах.")
    return lines


def _median(values: list[float]) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    mid = len(ordered) // 2
    if len(ordered) % 2:
        return ordered[mid]
    return (ordered[mid - 1] + ordered[mid]) / 2


def _competitor_in_place(positioning, our_brand: str, place: tuple[str, str, str], comp_key: str) -> bool:
    """Конкурент попадает в срез, только если он есть в позиционировании этой группы или подгруппы."""
    _category, group, subgroup = place
    if not (group or subgroup):
        return False
    brand = norm_text(our_brand)
    rows = []
    if subgroup:
        rows = (getattr(positioning, "by_sub", None) or {}).get((brand, norm_text(subgroup)), [])
    if not rows and group:
        rows = (getattr(positioning, "by_group", None) or {}).get((brand, norm_text(group)), [])
    if not rows:
        rows = [
            row for row in getattr(positioning, "rows", None) or []
            if norm_text(getattr(row, "brand", "")) == brand
            and (not group or norm_text(getattr(row, "group", "")) == norm_text(group))
            and (not subgroup or norm_text(getattr(row, "subgroup", "")) == norm_text(subgroup))
        ]
    for row in rows:
        for comp in getattr(row, "competitors", None) or []:
            if brand_key(getattr(comp, "name", "")) == comp_key or brand_key(getattr(comp, "key", "")) == comp_key:
                return True
    return False


def _catalog(src: Sources, positioning) -> dict[tuple[str, str], list[tuple[str, str, str]]]:
    """Кросс конкурента -> наши категория, группа, подгруппа. Только кроссы «Да»."""
    pst = {}
    for item in src.pst or []:
        pst[(norm_text(item.brand), norm_article(item.article))] = (
            clean(item.category), clean(item.group), clean(item.subgroup),
        )
    pos = {}
    for row in getattr(positioning, "rows", None) or []:
        pos.setdefault((norm_text(row.brand), norm_text(row.group), norm_text(row.subgroup)), clean(row.category))
    places: dict[tuple[str, str], set[tuple[str, str, str]]] = defaultdict(set)
    for (brand, article), groups in (src.crosses or {}).items():
        place = pst.get((brand, article))
        if place:
            category, group, subgroup = place
            if not category:
                category = pos.get((brand, norm_text(group), norm_text(subgroup)), "")
            place = (category, group, subgroup)
        for rows in groups.values():
            for row in rows:
                if not row.yes:
                    continue
                key = (brand_key(row.cross_brand), row.cross_key)
                if place and not _competitor_in_place(positioning, brand, place, key[0]):
                    continue
                places[key].add(place or ("", "", ""))
    return {key: sorted(value) for key, value in places.items()}


def _snap(files: list[PriceFile], date: str) -> dict[tuple[str, str, str], float]:
    """(поставщик, бренд, артикул) -> минимальная цена на эту дату."""
    out: dict[tuple[str, str, str], float] = {}
    for pf in files:
        if (pf.shot_date or "без даты") != date or not pf.enabled:
            continue
        sup = supplier_key(pf.supplier) or "—"
        for (brand, article), price in offer_index(pf).items():
            key = (sup, brand, article)
            prev = out.get(key)
            if prev is None or price < prev:
                out[key] = price
    return out


def competitor_brand_names(positioning) -> list[str]:
    """Бренды, с которыми мы сравниваем цены: мониторинг поставщиков и конкуренты в строках позиционирования."""
    names: list[str] = []
    seen: set[str] = set()
    for brand in getattr(positioning, "monitoring", {}) or {}:
        key = brand_key(brand)
        if key and key not in seen:
            seen.add(key)
            names.append(str(brand))
    for row in getattr(positioning, "rows", None) or []:
        for comp in getattr(row, "competitors", None) or []:
            label = clean(getattr(comp, "name", ""))
            key = brand_key(label)
            if key and key not in seen:
                seen.add(key)
                names.append(label)
    return names


def _brand_names(files: list[PriceFile]) -> dict[str, str]:
    seen: dict[str, str] = {}
    for pf in files:
        for offer in pf.offers:
            key = brand_key(offer.brand)
            if key and key not in seen:
                seen[key] = clean(offer.brand) or key
    return seen


def build(src: Sources, files: list[PriceFile], positioning, folder: Path, date_from: str | None, date_to: str | None) -> dict:
    currency = read_currency(folder)
    dates = sorted({pf.shot_date or "без даты" for pf in files if pf.enabled and pf.offers}, key=_date_key)
    priced = [d for d in dates if d != "без даты"] or dates
    if date_from not in priced or date_to not in priced or _date_key(date_from or "") >= _date_key(date_to or ""):
        date_from = priced[-2] if len(priced) >= 2 else (priced[0] if priced else "")
        date_to = priced[-1] if priced else ""
    pieces = [_offer_token(pf) for pf in files if pf.enabled]
    cur = find_currency_file(folder)
    pieces.append(_file_token(cur) if cur else "no-cur")
    pieces.append(f"{len(src.crosses or {})}|{len(src.pst or [])}")
    token = hashlib.sha1(("\n".join(sorted(pieces)) + f"|{date_from}|{date_to}|{DYN_VERSION}").encode()).hexdigest()
    cached = _load("result", token)
    if isinstance(cached, dict):
        cached = dict(cached)
        cached["cached"] = True
        return cached
    names = {brand_key(n): n for n in competitor_brand_names(positioning)}
    for key, label in _brand_names(files).items():
        if key in names and label:
            names[key] = label
    allowed = set(names)
    catalog = _catalog(src, positioning)
    left = _snap(files, date_from) if date_from else {}
    right = _snap(files, date_to) if date_to and date_to != date_from else {}
    buckets: dict[tuple, list[float]] = defaultdict(list)
    common = 0
    for key, old in left.items():
        new = right.get(key)
        if not new or not old:
            continue
        _brand = key[1]
        if _brand not in allowed:
            continue
        common += 1
        change = new / old - 1
        buckets[("", "", "", _brand)].append(change)
        seen: set[tuple[str, str, str]] = set()
        for place in catalog.get((_brand, key[2]), []):
            if place in seen:
                continue
            seen.add(place)
            category, group, subgroup = place
            if category:
                buckets[(category, "", "", _brand)].append(change)
            if group:
                buckets[(category, group, "", _brand)].append(change)
            if subgroup:
                buckets[(category, group, subgroup, _brand)].append(change)
    rows = []
    for (category, group, subgroup, brand), values in buckets.items():
        rows.append({
            "category": category, "group": group, "subgroup": subgroup,
            "brand": names.get(brand, brand), "n": len(values), "change": _median(values),
        })
    rows.sort(key=lambda r: (r["category"], r["group"], r["subgroup"], abs(r["change"] or 0)), reverse=False)
    rates = []
    for code, series in currency["rates"].items():
        old = series.get(date_from) if date_from != "без даты" else None
        new = series.get(date_to) if date_to != "без даты" else None
        rates.append({
            "code": code,
            "old": old,
            "new": new,
            "change": (new / old - 1) if old and new else None,
        })
    shots = []
    for pf in files:
        if not pf.enabled:
            continue
        shots.append({
            "date": pf.shot_date or "без даты",
            "supplier": pf.supplier,
            "file": pf.sheet.file.name,
            "offers": len(pf.offers),
        })
    payload = {
        "cached": False,
        "dates": priced,
        "date_from": date_from,
        "date_to": date_to,
        "compared": common,
        "currency_path": currency["path"],
        "currency_dates": currency["dates"],
        "rates": rates,
        "rows": rows,
        "shots": shots,
        "has_groups": any(r["group"] or r["category"] for r in rows),
        "conclusions": _conclusions(rows, rates),
    }
    _save("result", token, payload)
    return payload


EXPORT_NAME = "Динамика цен.xlsx"


def _article_changes(src: Sources, files: list[PriceFile], positioning, date_from: str, date_to: str) -> list[dict]:
    """Подробные строки для выгрузки: один артикул конкурента в каждой нашей подгруппе, где он есть в позиционировании."""
    names = {brand_key(name): name for name in competitor_brand_names(positioning)}
    for key, label in _brand_names(files).items():
        if key in names and label:
            names[key] = label
    suppliers = {}
    for pf in files:
        suppliers.setdefault(supplier_key(pf.supplier), pf.supplier)
    catalog = _catalog(src, positioning)
    left = _snap(files, date_from)
    right = _snap(files, date_to)
    out = []
    for (supplier, brand, article), old in left.items():
        new = right.get((supplier, brand, article))
        if not new or brand not in names:
            continue
        places = catalog.get((brand, article)) or []
        if not places:
            places = [("", "", "")]
        change = new / old - 1
        for category, group, subgroup in places:
            out.append({
                "supplier": suppliers.get(supplier, supplier),
                "brand": names.get(brand, brand),
                "article": article,
                "category": category,
                "group": group,
                "subgroup": subgroup,
                "old": old,
                "new": new,
                "change": change,
            })
    return out


def write_dynamics_export(path: Path, src: Sources, files: list[PriceFile], positioning, folder: Path,
                          date_from: str | None, date_to: str | None) -> dict:
    """Книга для руководства и для последующего разбора: выводы, срезы и каждый артикул."""
    from openpyxl import Workbook

    payload = build(src, files, positioning, folder, date_from, date_to)
    articles = _article_changes(src, files, positioning, payload["date_from"], payload["date_to"]) if payload["date_from"] else []
    book = Workbook(write_only=True)
    summary = book.create_sheet("Выводы")
    brands = book.create_sheet("Бренды")
    slices = book.create_sheet("Срезы")
    detail = book.create_sheet("Артикулы")
    summary.append(["Выводы"])
    summary.append([f"С {payload['date_from']} по {payload['date_to']}"])
    summary.append([f"Общих артикулов конкурентов: {payload['compared']}"])
    for line in payload.get("conclusions") or []:
        summary.append([line])
    summary.append([])
    summary.append(["Курс"])
    for rate in payload["rates"]:
        summary.append([rate["code"], rate.get("old"), rate.get("new"), rate.get("change")])
    brands.append(["Бренд", "Артикулов", "Изменение цены"] + [rate["code"] for rate in payload["rates"]])
    for row in payload["rows"]:
        if row["category"] or row["group"] or row["subgroup"] or not row["brand"]:
            continue
        brands.append([row["brand"], row["n"], row["change"]] + [
            None if row["change"] is None or rate.get("change") is None else row["change"] - rate["change"]
            for rate in payload["rates"]
        ])
    slices.append(["Категория", "Группа", "Подгруппа", "Бренд", "Артикулов", "Изменение цены"])
    for row in payload["rows"]:
        if not (row["category"] or row["group"] or row["subgroup"]):
            continue
        slices.append([row["category"], row["group"], row["subgroup"], row["brand"], row["n"], row["change"]])
    detail.append([
        "Дата с", "Дата по", "Категория", "Группа", "Подгруппа", "Поставщик",
        "Бренд конкурента", "Артикул", "Цена была", "Цена стала", "Изменение",
    ])
    for row in articles:
        detail.append([
            payload["date_from"], payload["date_to"], row["category"], row["group"], row["subgroup"],
            row["supplier"], row["brand"], row["article"], row["old"], row["new"], row["change"],
        ])
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return {"rows": len(articles), "brands": payload["compared"], "name": path.name}
