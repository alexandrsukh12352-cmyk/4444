"""Листы «Сводка» и «Цены конкурентов» итогового отчёта.

«Сводка» — первый лист на один экран: главное, мы против конкурентов, таблица по подгруппам и один
график. Главное и таблица — формулы по «Товарам»: поправили там цену — сводка пересчиталась.

«Цены конкурентов» — короткая таблица по товару: группа, подгруппа, разновидность, ABC, потенциал,
продажи, ВР, наша цена, цены конкурентов и отклонения. Числа — ссылки на «Товары».
"""

from __future__ import annotations

import datetime as dt

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import style
from .style import argb

NUM = ">-1E+99"  # условие SUMIFS / COUNTIFS «ячейка — число»
MONEY0 = '#,##0;-#,##0;""'
MONEY = '#,##0.00'
SIGNED_RUB = '+#,##0;-#,##0;0'
PCT1 = '0.0%;-0.0%;0.0%'
PCT1S = '+0.0%;-0.0%;0.0%'
SHARE = '0%;-0%;""'
INT = '#,##0'

INK = Font(name=style.FONT, color=style.INK)
MUTED = Font(name=style.FONT, color=style.MUTED, size=9)
BOLD = Font(name=style.FONT, color=style.INK, bold=True)
HEAD_FILL = PatternFill("solid", fgColor=argb("1F3A4D"))
HEAD_FONT = Font(name=style.FONT, color=argb("FFFFFF"), bold=True, size=10)
CARD_FILL = PatternFill("solid", fgColor=argb("F3F7F8"))
TOTAL_FILL = PatternFill("solid", fgColor=argb("E8EEF1"))
LINE = Side(style="thin", color=argb("D5E0E3"))
ROW_BORDER = Border(bottom=Side(style="thin", color=argb("E3E8EC")))
UP_FONT = Font(name=style.FONT, color=argb("1D5FA8"), bold=True)
DOWN_FONT = Font(name=style.FONT, color=argb("B45309"), bold=True)
DEAR_FILL, CHEAP_FILL = PatternFill(bgColor=argb("FDE2D4")), PatternFill(bgColor=argb("DCEBFA"))


def _q(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _crit(value: str) -> str:
    """Текст для условия SUMIFS: без подстановочных знаков и с удвоенными кавычками."""
    text = (value or "").replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""')
    return f'"{text}"'


class Across:
    """Формулы по всем листам отчёта, с условием на подгруппу или без."""

    def __init__(self, infos):
        self.infos = infos

    def parts(self, subgroup: str | None):
        for info in self.infos:
            if subgroup is None:
                yield info, None
                continue
            here = [r.subgroup for r in info.rows]
            if not here or subgroup not in here:
                continue
            if all(x == subgroup for x in here) or "subgroup" not in info.letters:
                yield info, None
            else:
                yield info, f'{info.rng("subgroup")},{_crit(subgroup)}'

    def sum(self, key: str, subgroup: str | None, *conds: tuple[str, str]) -> str:
        out = []
        for info, sg in self.parts(subgroup):
            if key not in info.letters or any(k not in info.letters for k, _ in conds):
                continue
            extra = [f"{info.rng(k)},{c}" for k, c in conds] + ([sg] if sg else [])
            out.append(f"SUMIFS({info.rng(key)},{','.join(extra)})" if extra else f"SUM({info.rng(key)})")
        return "(" + ("+".join(out) or "0") + ")"

    def count(self, key: str, subgroup: str | None, cond: str = f'"{NUM}"') -> str:
        out = []
        for info, sg in self.parts(subgroup):
            if key not in info.letters:
                continue
            out.append(f"COUNTIFS({info.rng(key)},{cond}{',' + sg if sg else ''})")
        return "(" + ("+".join(out) or "0") + ")"

    def rows(self, subgroup: str | None) -> str:
        out = []
        for info, sg in self.parts(subgroup):
            out.append(f"COUNTIFS({sg})" if sg else f"ROWS({info.rng('code')})")
        return "(" + ("+".join(out) or "0") + ")"

    def has(self, key: str) -> bool:
        return any(key in i.letters for i in self.infos)


def _cell(ws, ref: str, value, fmt: str | None = None, font: Font | None = None, fill=None, align=None):
    c = ws[ref]
    c.value = value
    if fmt:
        c.number_format = fmt
    c.font = font or INK
    if fill is not None:
        c.fill = fill
    if align is not None:
        c.alignment = align
    return c


def _header(ws, row: int, titles: list[str], first_col: int = 2, height: float = 30) -> None:
    for i, t in enumerate(titles):
        c = ws.cell(row, first_col + i, t)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def _section(ws, row: int, text: str, hint: str = "") -> None:
    _cell(ws, f"B{row}", text, font=Font(name=style.FONT, size=13, bold=True, color=style.INK))
    ws.row_dimensions[row].height = 22
    if hint:
        _cell(ws, f"B{row + 1}", hint, font=MUTED)


def _locate(infos) -> dict[str, tuple[str, int]]:
    where = {}
    for info in infos:
        for i, r in enumerate(info.rows):
            where.setdefault(r.code, (info.title, info.first + i))
    return where


def _link(where, letters_of, code: str, key: str) -> str | None:
    found = where.get(code)
    if not found:
        return None
    title, n = found
    col = letters_of[title].get(key)
    return f"{_q(title)}!${col}${n}" if col else None


def no_price_reason(r) -> str:
    picks = [p for lst in (r.picks or {}).values() for p in (lst or [])]
    if r.any_price and not r.target:
        names = ", ".join(c.name for c in r.competitors if r.prices.get(c.key))
        return f"цена есть только у {names}; цель — по первому в позиционировании"
    if any(p.article for p in picks):
        return "кросс есть, но его нет в прайсах поставщиков"
    return "нет кросса на конкурентов из позиционирования"


GREY, BLUE, RED_LINE = "A6A6A6", "4472C4", "C00000"
MLN = '#,##0.0,,;-#,##0.0,,;""'
MLN_SIGNED = '+#,##0.0,,;-#,##0.0,,;0.0'
HEAD_FILL_XL = PatternFill("solid", fgColor=style.XL_HEAD)
TOTAL_FILL_XL = PatternFill("solid", fgColor=style.XL_TOTAL)
TOTAL_LINE = Border(top=Side(style="thin", color=argb("595959")), bottom=Side(style="thin", color=argb("595959")))
KV_LINE = Border(bottom=Side(style="thin", color=style.XL_GRID))


def _xl_header(ws, row: int, titles: list[str], first_col: int = 1, height: float = 30) -> None:
    """Шапка таблицы, как её делают руками: серая заливка, жирный текст, тонкие границы."""
    for i, t in enumerate(titles):
        c = ws.cell(row, first_col + i, t)
        c.font = BOLD
        c.fill = HEAD_FILL_XL
        c.border = style.XL_HEAD_BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def _vr_chart(ws, data_ws, summary: dict, rec: bool, min_vr: float, anchor: str) -> None:
    """Один график: ВР по подгруппам (или по ABC, если подгруппа одна) — сейчас и после. Числа на скрытом листе."""
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.drawing.line import LineProperties

    groups = [x for x in summary.get("subgroups") or [] if x.get("items") and x.get("vr") is not None]
    what = "подгруппам"
    if len(groups) < 2:
        groups = [x for x in summary.get("abc") or [] if x.get("items") and x.get("vr") is not None]
        what = "ABC"
    if not groups:
        return
    data_ws.sheet_state = "hidden"
    data_ws.cell(1, 1, f"ВР по {what}").font = BOLD
    head = ["", "Сейчас", "После", "Мин. ВР"] if rec else ["", "Сейчас", "Мин. ВР"]
    for j, h in enumerate(head, 1):
        data_ws.cell(2, j, h)
    for i, x in enumerate(groups, 3):
        values = [x["name"] or "пусто", x["vr"]] + ([x["vr_rec"]] if rec else []) + [min_vr / 100]
        for j, v in enumerate(values, 1):
            c = data_ws.cell(i, j, v)
            if j > 1:
                c.number_format = "0.0%"
    last = 2 + len(groups)
    data_ws.column_dimensions["A"].width = 30
    n_bars = 3 if rec else 2
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    _title(bar, f"ВР по {what}: сейчас и после" if rec else f"ВР по {what}")
    bar.add_data(Reference(data_ws, min_col=2, max_col=n_bars, min_row=2, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(data_ws, min_col=1, min_row=3, max_row=last))
    for series, color in zip(bar.series, [GREY, BLUE]):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line = LineProperties(solidFill=color)
        series.dLbls = DataLabelList()
        series.dLbls.showVal = True
        series.dLbls.showSerName = series.dLbls.showCatName = series.dLbls.showLegendKey = False
    line = LineChart()
    line.add_data(Reference(data_ws, min_col=n_bars + 1, max_col=n_bars + 1, min_row=2, max_row=last), titles_from_data=True)
    line.series[0].graphicalProperties.line = LineProperties(solidFill=RED_LINE, w=15875, prstDash="dash")
    line.series[0].marker.symbol = "none"
    bar.gapWidth = 80
    bar.y_axis.majorGridlines = None
    bar.y_axis.numFmt = "0%"
    bar.y_axis.delete = False
    bar.x_axis.delete = False
    bar.legend.position = "b"
    bar += line
    bar.width, bar.height = min(32, 10 + 2.2 * len(groups)), 7.5
    ws.add_chart(bar, anchor)


def _title(chart, text: str) -> None:
    """Заголовок графика 11 пт, а не 14 пт, как ставит Excel по умолчанию."""
    from openpyxl.chart.text import RichText, Text
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun

    props = CharacterProperties(sz=1100, b=True, solidFill="404040")
    run = RegularTextRun(rPr=props, t=text)
    chart.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=props), r=[run])])), overlay=False)


def _comp_name(s: dict) -> str:
    return str(s["title"]).replace(" (без учёта % позиц.)", "").replace("Против ", "")


def write_summary_sheet(ws, rows, infos, opt, src, title: str, summary: dict, data_ws=None, signals=None) -> None:
    """«Сводка» на один экран: главное, мы против конкурентов, таблица по подгруппам и один график.

    Главное и таблица — формулы по листу «Товары»: поправили там цену — сводка пересчиталась.
    Против конкурентов и график — на момент формирования.
    """
    ro = opt.recs
    a = Across(infos)
    rec = ro.enabled and a.has("rec")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    widths = {"A": 40, "B": 11, "C": 11, "D": 11, "E": 11, "F": 12, "G": 11, "H": 3, "I": 20, "J": 11, "K": 11, "L": 11, "M": 11}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    _cell(ws, "A1", title, font=Font(name=style.FONT, size=14, bold=True, color=style.INK))
    ws.row_dimensions[1].height = 22
    bits = [f"Ценовой анализ на {dt.datetime.now():%d.%m.%Y}"]
    if getattr(src, "sales_period", ""):
        bits.append(f"продажи {src.sales_period}")
    if getattr(src, "sales3_period", ""):
        bits.append(f"последние продажи {src.sales3_period}")
    from .report import PRICE_PICKS, PRICE_SOURCES

    bits.append(f"цены конкурентов: {PRICE_SOURCES.get(opt.price_source, opt.price_source)}, "
                f"{PRICE_PICKS.get(getattr(opt, 'price_pick', 'median'), '').lower()}")
    _cell(ws, "A2", " · ".join(bits), font=Font(name=style.FONT, size=9, color=style.XL_MUTED))

    sub_font = Font(name=style.FONT, size=11, bold=True, color=style.INK)
    big = Font(name=style.FONT, bold=True, color=style.INK)
    right = Alignment(horizontal="right")

    # ---- Главное: формулы по «Товарам»
    sale = a.sum("sum_sale", None)
    profit = a.sum("sum_profit", None)
    sale_vr = a.sum("sum_sale", None, ("vr", f'"{NUM}"'))
    vr_now = f'IFERROR({profit}/{sale_vr},"")'
    items = a.rows(None)
    priced_sale = a.sum("sum_sale", None, ("target", f'"{NUM}"')) if a.has("target") else "0"
    priced_items = a.count("target", None) if a.has("target") else "0"
    _cell(ws, "A4", "Главное", font=sub_font)
    lines = []
    if rec:
        profit_r = a.sum("sum_profit_rec", None)
        sale_vr_r = a.sum("sum_sale_rec", None, ("vr_rec", f'"{NUM}"'))
        ups = a.count("change", None, '">0.0005"')
        downs = a.count("change", None, '"<-0.0005"')
        lines += [
            ("ВР сейчас / после", [(f"={vr_now}", "0.0%"), (f'=IFERROR({profit_r}/{sale_vr_r},"")', "0.0%")]),
            ("Прибыль, млн ₽ / изменение", [(f"={profit}", MLN), (f"={profit_r}-{profit}", MLN_SIGNED)]),
        ]
    else:
        lines += [("ВР", [(f"={vr_now}", "0.0%")]), ("Прибыль, млн ₽", [(f"={profit}", MLN)])]
    lines.append(("Выручка с ценой конкурента / товаров", [(f"=IFERROR({priced_sale}/{sale},0)", "0%"), (f"={priced_items}", INT)]))
    if rec:
        lines.append(("Цены: поднять / снизить / не трогать", [(f"={ups}", INT), (f"={downs}", INT), (f"={items}-{ups}-{downs}", INT)]))
    lines.append(("Товаров в отчёте", [(f"={items}", INT)]))
    from .salestrend import group_trend

    gt = group_trend(src, rows)
    if gt:
        lines.append((f"Продажи за 4 нед. до {gt['end']} к 4 нед. до них", [(gt["change"], "+0%;-0%;0%")]))
    for i, (label, values) in enumerate(lines):
        n = 5 + i
        _cell(ws, f"A{n}", label)
        for col in "ABCD":
            ws[f"{col}{n}"].border = KV_LINE
        for j, (v, fmt) in enumerate(values):
            _cell(ws, f"{'BCD'[j]}{n}", v, fmt, big, align=right)
    kv_last = 4 + len(lines)
    if rec:
        delta = "C6"
        ws.conditional_formatting.add(delta, FormulaRule(formula=["AND(ISNUMBER(C6),C6<0)"], font=Font(color=style.XL_RED)))
        ws.conditional_formatting.add(delta, FormulaRule(formula=["AND(ISNUMBER(C6),C6>0)"], font=Font(color=style.XL_GREEN)))

    # ---- Против конкурентов: снимок на момент формирования
    comps = [s for s in summary.get("slices") or [] if s.get("kind") == "comp" and s.get("n")]
    comp_last = 4
    if comps:
        _cell(ws, "I4", "Против конкурентов", font=sub_font)
        titles = ["Конкурент", "Товаров с ценой", "Разница сейчас"] + (["После"] if rec else []) + ["Цель"]
        _xl_header(ws, 5, titles, first_col=9)
        green, red = Font(color=style.XL_GREEN), Font(color=style.XL_RED)
        cheaper, dearer = (green, red) if opt.green_when != "above" else (red, green)
        for i, s_ in enumerate(comps):
            n = 6 + i
            now = s_.get("weighted") if s_.get("weighted") is not None else s_.get("median")
            after = s_.get("weighted_rec") if s_.get("weighted_rec") is not None else s_.get("median_rec")
            values = [(_comp_name(s_), None), (s_["n"], INT), (now, "+0%;-0%;0%")] + ([(after, "+0%;-0%;0%")] if rec else []) \
                + [(s_.get("target_dev"), "+0%;-0%;0%")]
            for j, (v, fmt) in enumerate(values):
                c = _cell(ws, f"{get_column_letter(9 + j)}{n}", v, fmt, align=None if j == 0 else right)
                c.border = style.XL_CELL
        comp_last = 5 + len(comps)
        dev_cols = f"K6:{'L' if rec else 'K'}{comp_last}"
        ws.conditional_formatting.add(dev_cols, FormulaRule(formula=["AND(ISNUMBER(K6),K6>0)"], font=cheaper))
        ws.conditional_formatting.add(dev_cols, FormulaRule(formula=["AND(ISNUMBER(K6),K6<0)"], font=dearer))
        _cell(ws, f"I{comp_last + 1}", "плюс — мы дешевле конкурента, минус — дороже; среднее по выручке",
              font=Font(name=style.FONT, size=9, color=style.XL_MUTED))
        comp_last += 1

    # ---- что изменилось у конкурентов с прошлого прогона
    if signals:
        n = kv_last + 2
        _cell(ws, f"A{n}", "Что изменилось на рынке", font=sub_font)
        for i, sg in enumerate(signals[:4], 1):
            _cell(ws, f"A{n + i}", f"{sg['text']}: {sg['hint']}")
        kv_last = n + min(len(signals), 4)

    # ---- По подгруппам: формулы по «Товарам»
    row = max(kv_last, comp_last) + 2
    subgroups = list(dict.fromkeys(r.subgroup for r in rows))
    _cell(ws, f"A{row}", "По подгруппам", font=sub_font)
    titles = ["Подгруппа", "Товаров", "Выручка, млн ₽", "ВР"]
    titles += ["ВР после", "Δ прибыли, млн ₽"] if rec else ["Прибыль, млн ₽"]
    titles += ["С ценой конк."]
    head = row + 1
    _xl_header(ws, head, titles)
    lines = ([(sg or "без подгруппы", sg) for sg in subgroups] if len(subgroups) > 1 else []) + [("Итого", None)]
    for i, (label, sg) in enumerate(lines):
        n = head + 1 + i
        total = sg is None
        font = BOLD if total else INK
        s_sale = a.sum("sum_sale", sg)
        s_profit = a.sum("sum_profit", sg)
        s_vr = a.sum("sum_sale", sg, ("vr", f'"{NUM}"'))
        vals = [(f"={a.rows(sg)}", INT), (f"={s_sale}", MLN), (f'=IFERROR({s_profit}/{s_vr},"")', "0.0%")]
        if rec:
            p_r = a.sum("sum_profit_rec", sg)
            s_vr_r = a.sum("sum_sale_rec", sg, ("vr_rec", f'"{NUM}"'))
            vals += [(f'=IFERROR({p_r}/{s_vr_r},"")', "0.0%"), (f"={p_r}-{s_profit}", MLN_SIGNED)]
        else:
            vals.append((f"={s_profit}", MLN))
        priced = a.sum("sum_sale", sg, ("target", f'"{NUM}"')) if a.has("target") else "0"
        vals.append((f'=IFERROR({priced}/{s_sale},"")', "0%"))
        c = ws.cell(n, 1, label)
        c.font = font
        c.border = TOTAL_LINE if total else style.XL_CELL
        for j, (v, fmt) in enumerate(vals):
            c = ws.cell(n, 2 + j, v)
            c.number_format = fmt
            c.font = font
            c.alignment = right
            c.border = TOTAL_LINE if total else style.XL_CELL
        if total:
            for j in range(1, len(vals) + 2):
                ws.cell(n, j).fill = TOTAL_FILL_XL
    last = head + len(lines)
    ws.conditional_formatting.add(f"D{head + 1}:{'E' if rec else 'D'}{last}", FormulaRule(
        formula=[f"AND(ISNUMBER(D{head + 1}),D{head + 1}<{ro.min_vr / 100})"], font=Font(color=style.XL_RED, bold=True)))
    if rec:
        ws.conditional_formatting.add(f"F{head + 1}:F{last}", FormulaRule(
            formula=[f"AND(ISNUMBER(F{head + 1}),F{head + 1}<0)"], font=Font(color=style.XL_RED)))
        ws.conditional_formatting.add(f"F{head + 1}:F{last}", FormulaRule(
            formula=[f"AND(ISNUMBER(F{head + 1}),F{head + 1}>0)"], font=Font(color=style.XL_GREEN)))

    # ---- один график
    if data_ws is not None:
        _vr_chart(ws, data_ws, summary, rec, ro.min_vr, f"A{last + 2}")
    style.print_wide(ws)


def write_prices_overview(ws, rows, infos, opt) -> None:
    """Лист «Цены конкурентов»: короткая таблица по товару, числа — ссылки на лист «Товары»."""
    from .report import competitor_order

    ro = opt.recs
    tol = ro.tolerance / 100
    where = _locate(infos)
    letters_of = {i.title: i.letters for i in infos}
    comps = competitor_order(rows)
    rec = ro.enabled and any("rec" in i.letters for i in infos)
    has_qty = any((r.sales_qty or 0) > 0 for r in rows)
    has_target = any(r.target for r in rows)
    has_potential = any("potential" in i.letters for i in infos)

    def link(r, key):
        ref = _link(where, letters_of, r.code, key)
        return f'=IF(ISNUMBER({ref}),{ref},"")' if ref else None

    # (заголовок, вид, ширина, значение(r, n, L), формат, ключ)
    cols: list[tuple] = [
        ("Артикул", "id", 14, lambda r, n, L: r.article, None, "article"),
        ("Наименование", "id", 30, lambda r, n, L: r.name, None, "name"),
        ("Группа", "info", 16, lambda r, n, L: r.group, None, "group"),
        ("Подгруппа", "info", 22, lambda r, n, L: r.subgroup, None, "subgroup"),
        ("Разновидность", "info", 36, lambda r, n, L: r.variety, None, "variety"),
        ("ABC", "calc", 6, lambda r, n, L: r.abc, None, "abc"),
        ("Потенциал, шт" if has_potential else "Остаток, шт", "calc", 10,
         lambda r, n, L: link(r, "potential" if has_potential else "stock"), '#,##0', "potential"),
        ("Продажи, ₽", "calc", 13, lambda r, n, L: link(r, "sales") or r.sales, '#,##0', "sales"),
    ]
    if has_qty:
        cols.append(("Продажи, шт", "calc", 9, lambda r, n, L: link(r, "sales_qty"), '#,##0', "sales_qty"))
    if any(getattr(r, "sales3", None) is not None for r in rows):
        qty3 = any(getattr(r, "sales3_qty", None) for r in rows)
        cols.append(("Продажи 3 мес, шт" if qty3 else "Продажи 3 мес, ₽", "calc", 10,
                     lambda r, n, L: link(r, "sales3"), '#,##0', "sales3"))
    cols += [
        ("ВР", "calc", 7, lambda r, n, L: link(r, "vr"), PCT1, "vr"),
        ("Кат. А, ₽", "calc", 11, lambda r, n, L: link(r, "cat_a"), INT, "cat_a"),
    ]
    for i, c in enumerate(comps):
        cols.append((f"{c.name}, ₽", "comp", 11, lambda r, n, L, k=c.key: link(r, f"price:{k}"), INT, f"price:{c.key}"))
        cols.append(("Откл.", "comp", 7, lambda r, n, L, k=c.key:
                     f'=IF(AND(ISNUMBER({L["price:" + k]}{n}),ISNUMBER({L["cat_a"]}{n})),1-{L["cat_a"]}{n}/{L["price:" + k]}{n},"")',
                     PCT1S, f"dev:{c.key}"))
    if has_target:
        cols.append(("Цель, ₽", "pos", 11, lambda r, n, L: link(r, "target"), INT, "target"))
        cols.append(("Откл. от цели", "pos", 8, lambda r, n, L:
                     f'=IF(AND(ISNUMBER({L["target"]}{n}),ISNUMBER({L["cat_a"]}{n})),1-{L["cat_a"]}{n}/{L["target"]}{n},"")',
                     PCT1S, "gap"))
    if rec:
        cols.append(("Новая цена, ₽", "rec", 11, lambda r, n, L: link(r, "rec"), INT, "rec"))
        cols.append(("Изм.", "rec", 7, lambda r, n, L:
                     f'=IF(AND(ISNUMBER({L["rec"]}{n}),ISNUMBER({L["cat_a"]}{n})),{L["rec"]}{n}/{L["cat_a"]}{n}-1,"")',
                     '+0.0%;-0.0%;""', "change"))

    L = {key: get_column_letter(i) for i, (*_, key) in enumerate(cols, 1)}
    ordered = sorted(rows, key=lambda r: (r.group or "", r.subgroup or "", -(r.sales or 0), r.article))
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.cell(1, 1, "Цены: наша, конкурентов и отклонения").font = Font(name=style.FONT, size=14, bold=True, color=style.INK)
    ws.cell(2, 1, "Откл. = 1 − Кат. А / цена конкурента: плюс — мы дешевле, минус — дороже. Числа — ссылки на лист «Товары».").font = \
        Font(name=style.FONT, size=9, color=style.XL_MUTED)
    ws.row_dimensions[1].height = 22
    head = 3
    for i, (t, kind, width, *_rest) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    _xl_header(ws, head, [c[0] for c in cols], height=32)
    first = head + 1
    tones = [PatternFill("solid", fgColor=x) for x in style.XL_COMP]
    comp_index = {f"price:{c.key}": i % 2 for i, c in enumerate(comps)} | {f"dev:{c.key}": i % 2 for i, c in enumerate(comps)}
    fills = {"cat_a": PatternFill("solid", fgColor=style.XL_CAT_A), "rec": PatternFill("solid", fgColor=style.XL_NEW)}
    for key, col in L.items():
        if key in comp_index:
            ws[f"{col}{head}"].fill = tones[comp_index[key]]
        elif key in fills:
            ws[f"{col}{head}"].fill = fills[key]
    for n, r in enumerate(ordered, first):
        for i, (_t, _kind, _w, value, fmt, key) in enumerate(cols, 1):
            v = value(r, n, L)
            c = ws.cell(n, i, v if v != "" else None)
            c.font = BOLD if key == "article" else INK
            c.border = style.XL_CELL
            if fmt:
                c.number_format = fmt
            if key in comp_index:
                c.fill = tones[comp_index[key]]
            elif key in fills:
                c.fill = fills[key]
    last = max(first, first + len(ordered) - 1)
    green, red = Font(color=style.XL_GREEN), Font(color=style.XL_RED)
    above, below = (green, red) if opt.green_when == "above" else (red, green)
    for key, col in L.items():
        rng = f"{col}{first}:{col}{last}"
        x = f"{col}{first}"
        if key.startswith("dev:"):
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}<0)"], font=above))
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}>0)"], font=below))
        elif key == "gap":
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}<-{tol})"], font=above))
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}>{tol})"], font=below))
        elif key == "change":
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}>0.0005)"], font=Font(bold=True)))
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}<-0.0005)"], font=Font(color=style.XL_RED, bold=True)))
        elif key == "vr":
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({x}),{x}<{ro.min_vr / 100})"], font=Font(color=style.XL_RED, bold=True)))
    ws.freeze_panes = f"C{first}"
    for sel in ws.sheet_view.selection:
        if sel.pane == "bottomRight":
            sel.activeCell = sel.sqref = f"C{first}"
        elif sel.pane == "topRight":
            sel.activeCell = sel.sqref = "C1"
        elif sel.pane == "bottomLeft":
            sel.activeCell = sel.sqref = f"A{first}"
    ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{last}"
    style.print_wide(ws, f"{head}:{head}")
