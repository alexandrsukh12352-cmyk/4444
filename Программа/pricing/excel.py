"""Чтение листов Excel/CSV в «сетку» значений и поиск многострочной шапки."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

SUPPORTED_EXT = {".xlsx", ".xlsm", ".xls", ".csv", ".txt"}
HEADER_SCAN_ROWS = 40
MAX_HEADER_DEPTH = 4
LEVEL_SEP = " / "


@dataclass
class Sheet:
    file: Path
    name: str  # имя листа; для csv — пусто
    grid: list[list]

    @property
    def label(self) -> str:
        return f"{self.file.name} / {self.name}" if self.name else self.file.name


@dataclass
class Header:
    sheet: Sheet
    row: int  # индекс первой строки шапки
    depth: int  # сколько строк занимает шапка
    raw: list[str]  # «Закупочная цена / Вал.» — как в файле
    keys: list[str]  # то же в нижнем регистре для поиска

    @property
    def data(self) -> list[list]:
        return self.sheet.grid[self.row + self.depth :]

    def find(self, patterns: list[str]) -> int | None:
        for pattern in patterns:
            rx = re.compile(pattern)
            for idx, key in enumerate(self.keys):
                if key and rx.search(key):
                    return idx
        return None

    def letter(self, idx: int) -> str:
        return openpyxl.utils.get_column_letter(idx + 1)


# ---------------------------------------------------------------- значения


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def norm_key(value) -> str:
    return clean(value).replace("ё", "е").replace("Ё", "Е").lower()


_NUM_RE = re.compile(r"^[-+]?\d+(?:[.,]\d+)?$")


def num(value) -> float | None:
    """Число из ячейки 1С: 4571.29, «4 571,29», «1 855 007». Пусто/текст -> None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\xa0", "").replace(" ", "").strip()
    if not _NUM_RE.match(text):
        return None
    return float(text.replace(",", "."))


def looks_numeric(value) -> bool:
    return num(value) is not None


def code_key(value) -> str:
    """Код товара: 130909, «130 909», 130909.0 -> «130909»."""
    n = num(value)
    if n is not None and float(n).is_integer():
        return str(int(n))
    return re.sub(r"\s+", "", clean(value))


# ---------------------------------------------------------------- чтение файлов


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_sheets(path: Path) -> list[Sheet]:
    ext = path.suffix.lower()
    if ext in (".csv", ".txt"):
        text = _decode(path.read_bytes())
        first = text.split("\n", 1)[0]
        delimiter = max((";", ",", "\t"), key=first.count)
        return [Sheet(path, "", [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)])]
    if ext == ".xls":
        from .xls import read_xls

        sheets = []
        for name, grid, merges in read_xls(path):
            for r1, r2, c1, c2 in merges:
                if r1 >= len(grid):
                    continue
                value = grid[r1][c1] if c1 < len(grid[r1]) else None
                for r in range(r1, min(r2 + 1, len(grid))):
                    row = grid[r]
                    if len(row) <= c2:
                        row.extend([None] * (c2 + 1 - len(row)))
                    for c in range(c1, c2 + 1):
                        row[c] = value
            sheets.append(Sheet(path, name, grid))
        return sheets
    if ext not in (".xlsx", ".xlsm"):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        grid = [list(r) for r in ws.iter_rows(values_only=True)]
        # Объединённые ячейки: значение левой верхней раскладываем на весь диапазон —
        # так шапка из нескольких строк и товар из нескольких строк читаются одинаково.
        for rng in ws.merged_cells.ranges:
            value = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else None
            for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
                row = grid[r]
                if len(row) < rng.max_col:
                    row.extend([None] * (rng.max_col - len(row)))
                for c in range(rng.min_col - 1, rng.max_col):
                    row[c] = value
        sheets.append(Sheet(path, ws.title, grid))
    wb.close()
    return sheets


# ---------------------------------------------------------------- шапка


def _compose(levels: list) -> tuple[str, str]:
    parts: list[str] = []
    for v in levels:
        text = clean(v)
        if text and (not parts or norm_key(parts[-1]) != norm_key(text)):
            parts.append(text)
    raw = LEVEL_SEP.join(parts)
    return raw, norm_key(raw)


def find_header(sheet: Sheet, required: list[list[str]]) -> Header | None:
    """Ищет строку шапки, где нашлись все обязательные столбцы (каждый — список regex-вариантов).

    Шапка может занимать несколько строк (объединённые ячейки 1С): строки ниже первой
    считаются шапкой, пока в них нет чисел.
    """
    grid = sheet.grid
    for start in range(min(HEADER_SCAN_ROWS, len(grid))):
        if not any(clean(v) for v in grid[start]):
            continue
        depth = 1
        while (
            depth < MAX_HEADER_DEPTH
            and start + depth < len(grid)
            and any(clean(v) for v in grid[start + depth])
            and not any(looks_numeric(v) for v in grid[start + depth])
        ):
            depth += 1
        for d in range(depth, 0, -1):
            width = max(len(grid[start + i]) for i in range(d))
            composed = [
                _compose([grid[start + i][c] if c < len(grid[start + i]) else None for i in range(d)])
                for c in range(width)
            ]
            header = Header(sheet, start, d, [c[0] for c in composed], [c[1] for c in composed])
            if all(header.find(p) is not None for p in required):
                return header
    return None


def cell(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]
