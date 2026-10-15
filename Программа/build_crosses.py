#!/usr/bin/env python3
"""Подбор кросс-артикулов по позиционированию.

Берёт «Кроссы свод», ваш прайс (артикул + бренд + подгруппа) и позиционирование,
и для каждого вашего артикула отбирает кроссы на бренды-конкуренты из позиционирования:
сначала кроссы с «Использовать для сравнения цен» = «Да», а если их мало — добирает «Нет».

Запуск: Запуск.bat (Windows) или `python app/build_crosses.py [файлы...]`.
"""

from __future__ import annotations

import argparse
import configparser
import csv
import datetime as dt
import io
import os
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Не установлен модуль openpyxl. Выполните: python -m pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from positioning_data import MONITORING_CSV, POSITIONING_CSV  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT_DIR = ROOT / "Входящие"
DEFAULT_OUTPUT_DIR = ROOT / "Результат"
DEFAULT_SETTINGS = ROOT / "настройки.ini"

SUPPORTED_EXT = {".xlsx", ".xlsm", ".csv", ".txt"}
HEADER_SCAN_ROWS = 30

DEFAULT_SETTINGS_DATA = {
    "Отбор": {"мин_да": "1", "макс_нет": "0", "открыть_результат": "да"},
    "Столбцы_кроссы": {
        "артикул": "Артикул",
        "бренд": "Бренд",
        "артикул_кросс": "Артикул кросс",
        "бренд_кросс": "Бренд кросс",
        "сравнение_цен": "Использовать для сравнения цен",
        "артикул_поиск": "Артикул поиск",
        "артикул_кросс_поиск": "Артикул кросс поиск",
        "код": "Код",
    },
    "Столбцы_прайс": {
        "артикул": "Артикул | Артикул товара | Каталожный номер",
        "бренд": "Бренд | Торговая марка | ТМ | Производитель",
        "категория": "Категория NEW | Категория",
        "группа": "Группа NEW | Группа",
        "подгруппа": "Подгруппа NEW | Подгруппа",
        "доп_столбцы": "Код | Код товара | Наименование | Номенклатура | Цена",
    },
    "Столбцы_позиционирование": {
        "категория": "Категория NEW | Категория",
        "группа": "Группа NEW | Группа",
        "подгруппа": "Подгруппа NEW | Подгруппа",
        "бренд": "Торговая марка | Бренд",
    },
}

CROSS_REQUIRED = ("артикул", "бренд", "артикул_кросс", "бренд_кросс", "сравнение_цен")
CROSS_OPTIONAL = ("артикул_поиск", "артикул_кросс_поиск", "код")
PRICE_REQUIRED = ("артикул", "бренд")
PRICE_OPTIONAL = ("категория", "группа", "подгруппа")
POS_REQUIRED = ("бренд", "подгруппа")
POS_OPTIONAL = ("категория", "группа")
COMPETITOR_RE = re.compile(r"^бренд\s*конкурент\w*\s*(\d+)$")

# Кириллица, похожая на латиницу, в артикулах встречается по ошибке ввода.
LOOKALIKE = str.maketrans("АВЕКМНОРСТХУ", "ABEKMHOPCTXY")
YES_VALUES = {"ДА", "YES", "Y", "1", "TRUE", "ИСТИНА", "+"}

STATUS_YES = "Да"
STATUS_NO_FILL = "Нет (добор)"
RESULT_SHEETS = {"Сводка", "Для заноса", "По артикулам"}


class UserError(Exception):
    """Ошибка, понятная пользователю: печатаем текст без трассировки."""


# ---------------------------------------------------------------- нормализация


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def norm_text(value) -> str:
    return clean(value).replace("ё", "е").replace("Ё", "Е").upper()


def norm_header(value) -> str:
    return norm_text(value).lower()


def norm_article(value) -> str:
    return re.sub(r"[^0-9A-ZА-Я]", "", norm_text(value).translate(LOOKALIKE))


def is_yes(value) -> bool:
    return norm_text(value) in YES_VALUES


def to_number_if_possible(value):
    text = clean(value)
    if re.fullmatch(r"[-+]?\d+", text):
        return int(text)
    if re.fullmatch(r"[-+]?\d+[.,]\d+", text):
        return float(text.replace(",", "."))
    return text


# ---------------------------------------------------------------- настройки


@dataclass
class Settings:
    min_yes: int
    max_no: int
    open_result: bool
    cross_cols: dict[str, list[str]]
    price_cols: dict[str, list[str]]
    price_extra: list[str]
    pos_cols: dict[str, list[str]]


def split_aliases(raw: str) -> list[str]:
    return [a.strip() for a in raw.split("|") if a.strip()]


def load_settings(path: Path | None) -> Settings:
    cp = configparser.ConfigParser(inline_comment_prefixes=(";",), comment_prefixes=(";", "#"))
    cp.read_dict(DEFAULT_SETTINGS_DATA)
    if path and path.exists():
        cp.read(path, encoding="utf-8-sig")

    def as_int(key: str) -> int:
        raw = cp.get("Отбор", key)
        try:
            return max(0, int(raw))
        except ValueError:
            raise UserError(f"настройки.ini: «{key} = {raw}» — нужно целое число") from None

    def section(name: str, skip=()) -> dict[str, list[str]]:
        return {k: split_aliases(v) for k, v in cp.items(name) if k not in skip}

    return Settings(
        min_yes=as_int("мин_да"),
        max_no=as_int("макс_нет"),
        open_result=is_yes(cp.get("Отбор", "открыть_результат")),
        cross_cols=section("Столбцы_кроссы"),
        price_cols=section("Столбцы_прайс", skip=("доп_столбцы",)),
        price_extra=split_aliases(cp.get("Столбцы_прайс", "доп_столбцы", fallback="")),
        pos_cols=section("Столбцы_позиционирование"),
    )


# ---------------------------------------------------------------- чтение файлов


@dataclass
class Table:
    source: str  # «файл.xlsx / Лист1»
    rows: list[list]
    sheet: str = ""


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_csv_text(text: str) -> list[list]:
    first = text.split("\n", 1)[0]
    delimiter = max((";", ",", "\t"), key=first.count)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_tables(path: Path) -> list[Table]:
    ext = path.suffix.lower()
    if ext in (".csv", ".txt"):
        return [Table(path.name, read_csv_text(decode_text(path.read_bytes())))]
    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                Table(f"{path.name} / {ws.title}", [list(r) for r in ws.iter_rows(values_only=True)], ws.title)
                for ws in wb.worksheets
            ]
        finally:
            wb.close()
    if ext == ".xls":
        raise UserError(f"{path.name}: старый формат .xls не поддерживается — пересохраните файл как .xlsx")
    return []


def match_column(headers: list[str], aliases: list[str]) -> int | None:
    wanted = [norm_header(a) for a in aliases]
    for alias in wanted:
        for idx, h in enumerate(headers):
            if h == alias:
                return idx
    return None


@dataclass
class Located:
    table: Table
    header_idx: int
    cols: dict[str, int]
    headers_raw: list[str]

    @property
    def data_rows(self):
        return self.table.rows[self.header_idx + 1 :]

    def describe_columns(self) -> list[tuple[str, str]]:
        return [
            (key, f"«{self.headers_raw[i]}» (столбец {get_column_letter(i + 1)})")
            for key, i in self.cols.items()
        ]


def locate(table: Table, columns: dict[str, list[str]], required, optional) -> Located | None:
    for idx, row in enumerate(table.rows[:HEADER_SCAN_ROWS]):
        headers = [norm_header(c) for c in row]
        cols = {}
        for key in required:
            pos = match_column(headers, columns.get(key, []))
            if pos is None:
                break
            cols[key] = pos
        else:
            for key in optional:
                pos = match_column(headers, columns.get(key, []))
                if pos is not None:
                    cols[key] = pos
            return Located(table, idx, cols, [clean(c) for c in row])
    return None


def locate_positioning(table: Table, settings: Settings) -> Located | None:
    found = locate(table, settings.pos_cols, POS_REQUIRED, POS_OPTIONAL)
    if not found:
        return None
    headers = [norm_header(c) for c in table.rows[found.header_idx]]
    if not any(COMPETITOR_RE.match(h) for h in headers):
        return None
    return found


def cell(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


# ---------------------------------------------------------------- позиционирование


@dataclass
class Competitor:
    order: int
    name: str
    key: str
    percents: list[str] = field(default_factory=list)

    @property
    def percent(self):
        if not self.percents:
            return ""
        if len(self.percents) == 1:
            return to_number_if_possible(self.percents[0])
        return " / ".join(self.percents) + " (зависит от подгруппы)"


@dataclass
class PosRow:
    category: str
    group: str
    subgroup: str
    brand: str
    competitors: list[Competitor]


class Positioning:
    def __init__(self, rows: list[PosRow], monitoring: dict[str, list[str]], source: str):
        self.rows = rows
        self.monitoring = monitoring
        self.source = source
        self.by_sub: dict[tuple[str, str], list[PosRow]] = defaultdict(list)
        self.by_group: dict[tuple[str, str], list[PosRow]] = defaultdict(list)
        self.by_brand: dict[str, list[PosRow]] = defaultdict(list)
        names: dict[str, str] = {}
        for r in rows:
            for c in r.competitors:
                c.name = names.setdefault(c.key, c.name)
            b = norm_text(r.brand)
            self.by_sub[(b, norm_text(r.subgroup))].append(r)
            self.by_group[(b, norm_text(r.group))].append(r)
            self.by_brand[b].append(r)

    def competitor_keys(self, brand_key: str) -> set[str]:
        return {c.key for r in self.by_brand.get(brand_key, []) for c in r.competitors}

    @staticmethod
    def merge(rows: list[PosRow]) -> list[Competitor]:
        merged: dict[str, Competitor] = {}
        for r in rows:
            for c in r.competitors:
                m = merged.setdefault(c.key, Competitor(c.order, c.name, c.key))
                m.order = min(m.order, c.order)
                for p in c.percents:
                    if p not in m.percents:
                        m.percents.append(p)
        return sorted(merged.values(), key=lambda c: (c.order, c.name))

    def resolve(self, brand: str, group: str, subgroup: str):
        """Возвращает (уровень, строка позиционирования|None, конкуренты, причина-отказа)."""
        b = norm_text(brand)
        if b not in self.by_brand:
            return None, None, [], "бренда нет в позиционировании"
        if subgroup:
            rows = self.by_sub.get((b, norm_text(subgroup)), [])
            if len(rows) > 1 and group:
                rows = [r for r in rows if norm_text(r.group) == norm_text(group)] or rows
            if not rows:
                return None, None, [], "подгруппа не найдена в позиционировании бренда"
            comps = self.merge(rows)
            if not comps:
                return "подгруппа", rows[0], [], "в позиционировании не указаны конкуренты"
            return "подгруппа", rows[0], comps, ""
        if group:
            rows = self.by_group.get((b, norm_text(group)), [])
            if not rows:
                return None, None, [], "группа не найдена в позиционировании бренда"
            comps = self.merge(rows)
            if not comps:
                return "группа", None, [], "в позиционировании не указаны конкуренты"
            return "группа", None, comps, ""
        comps = self.merge(self.by_brand[b])
        if not comps:
            return "бренд", None, [], "в позиционировании не указаны конкуренты"
        return "бренд", None, comps, ""


def parse_positioning(located: Located, source: str, monitoring: dict[str, list[str]]) -> Positioning:
    header = [norm_header(c) for c in located.table.rows[located.header_idx]]
    comp_cols: list[tuple[int, int, int | None]] = []
    for idx, h in enumerate(header):
        m = COMPETITOR_RE.match(h)
        if m:
            pct = idx + 1 if idx + 1 < len(header) and header[idx + 1].startswith("%") else None
            comp_cols.append((int(m.group(1)), idx, pct))
    comp_cols.sort()

    rows = []
    for raw in located.data_rows:
        brand = clean(cell(raw, located.cols["бренд"]))
        if not brand:
            continue
        comps = []
        for order, idx, pct_idx in comp_cols:
            name = clean(cell(raw, idx))
            if name:
                pct = clean(cell(raw, pct_idx))
                comps.append(Competitor(order, name, norm_text(name), [pct] if pct else []))
        rows.append(
            PosRow(
                category=clean(cell(raw, located.cols.get("категория"))),
                group=clean(cell(raw, located.cols.get("группа"))),
                subgroup=clean(cell(raw, located.cols["подгруппа"])),
                brand=brand,
                competitors=comps,
            )
        )
    if not rows:
        raise UserError(f"{source}: в позиционировании нет ни одной строки с брендом")
    return Positioning(rows, monitoring, source)


def builtin_monitoring() -> dict[str, list[str]]:
    result = {}
    for row in read_csv_text(MONITORING_CSV):
        if row and clean(row[0]):
            result[norm_text(row[0])] = [clean(x) for x in row[1:] if clean(x)]
    return result


def builtin_positioning(settings: Settings) -> Positioning:
    table = Table("встроенное позиционирование", read_csv_text(POSITIONING_CSV))
    located = locate_positioning(table, settings)
    assert located, "встроенное позиционирование не распознано"
    return parse_positioning(located, "встроенное (зашито в скрипт)", builtin_monitoring())


# ---------------------------------------------------------------- кроссы


@dataclass
class Cross:
    article: str
    article_key: str
    brand: str
    yes: bool
    code: str


@dataclass
class CrossBase:
    # (бренд, артикул-поиск) -> бренд-кросс -> [Cross]
    index: dict[tuple[str, str], dict[str, list[Cross]]]
    display: dict[tuple[str, str], tuple[str, str]]  # ключ -> (артикул, бренд) как в своде
    rows_total: int
    rows_skipped: int


def parse_crosses(located: Located) -> CrossBase:
    c = located.cols
    index: dict[tuple[str, str], dict[str, list[Cross]]] = defaultdict(dict)
    display: dict[tuple[str, str], tuple[str, str]] = {}
    total = skipped = 0
    for raw in located.data_rows:
        article = clean(cell(raw, c["артикул"]))
        brand = clean(cell(raw, c["бренд"]))
        cross_article = clean(cell(raw, c["артикул_кросс"]))
        cross_brand = clean(cell(raw, c["бренд_кросс"]))
        if not (article or brand or cross_article):
            continue
        total += 1
        art_key = norm_article(cell(raw, c.get("артикул_поиск")) or article)
        cross_key = norm_article(cell(raw, c.get("артикул_кросс_поиск")) or cross_article)
        if not (art_key and brand and cross_key and cross_brand):
            skipped += 1
            continue
        key = (norm_text(brand), art_key)
        display.setdefault(key, (article, brand))
        bucket = index[key].setdefault(norm_text(cross_brand), [])
        yes = is_yes(cell(raw, c["сравнение_цен"]))
        for existing in bucket:
            if existing.article_key == cross_key:
                existing.yes = existing.yes or yes
                break
        else:
            bucket.append(Cross(cross_article, cross_key, cross_brand, yes, clean(cell(raw, c.get("код")))))
    return CrossBase(dict(index), display, total, skipped)


# ---------------------------------------------------------------- прайс


@dataclass
class Item:
    article: str
    brand: str
    category: str
    group: str
    subgroup: str
    extras: dict[str, object]

    @property
    def key(self) -> tuple[str, str]:
        return norm_text(self.brand), norm_article(self.article)


def parse_price(located: Located, settings: Settings) -> tuple[list[Item], list[str], int]:
    c = located.cols
    headers = [norm_header(h) for h in located.table.rows[located.header_idx]]
    used = set(c.values())
    extra_cols: list[tuple[str, int]] = []
    for alias in settings.price_extra:
        idx = match_column(headers, [alias])
        if idx is not None and idx not in used:
            extra_cols.append((located.headers_raw[idx], idx))
            used.add(idx)

    items, seen, dups = [], set(), 0
    for raw in located.data_rows:
        article = clean(cell(raw, c["артикул"]))
        brand = clean(cell(raw, c["бренд"]))
        if not article or not brand:
            continue
        item = Item(
            article=article,
            brand=brand,
            category=clean(cell(raw, c.get("категория"))),
            group=clean(cell(raw, c.get("группа"))),
            subgroup=clean(cell(raw, c.get("подгруппа"))),
            extras={name: cell(raw, idx) for name, idx in extra_cols},
        )
        if item.key in seen:
            dups += 1
            continue
        seen.add(item.key)
        items.append(item)
    return items, [name for name, _ in extra_cols], dups


def items_from_crosses(base: CrossBase) -> list[Item]:
    items = [Item(art, brand, "", "", "", {}) for art, brand in base.display.values()]
    return sorted(items, key=lambda i: (norm_text(i.brand), norm_article(i.article)))


# ---------------------------------------------------------------- отбор


@dataclass
class Pick:
    competitor: Competitor
    crosses: list[tuple[Cross, str]]  # (кросс, статус)

    @property
    def kind(self) -> str:
        statuses = {s for _, s in self.crosses}
        if not statuses:
            return "не найден"
        if statuses == {STATUS_YES}:
            return "Да"
        if STATUS_YES in statuses:
            return "Да + Нет"
        return "Нет"


@dataclass
class ItemResult:
    item: Item
    level: str | None
    pos_row: PosRow | None
    picks: list[Pick]
    reason: str


def select_crosses(available: list[Cross], min_yes: int, max_no: int) -> list[tuple[Cross, str]]:
    yes = [c for c in available if c.yes]
    picked = [(c, STATUS_YES) for c in yes]
    if len(yes) < min_yes:
        no = [c for c in available if not c.yes]
        if max_no:
            no = no[:max_no]
        picked += [(c, STATUS_NO_FILL) for c in no]
    return picked


def build(items: list[Item], base: CrossBase, pos: Positioning, settings: Settings) -> list[ItemResult]:
    results = []
    for item in items:
        level, pos_row, comps, reason = pos.resolve(item.brand, item.group, item.subgroup)
        picks = []
        by_brand = base.index.get(item.key, {})
        for comp in comps:
            # Без подгруппы неизвестно, какие конкуренты относятся к артикулу,
            # поэтому «не найден» по ним не показываем — только тех, на кого есть кроссы.
            if level == "бренд" and comp.key not in by_brand:
                continue
            picks.append(Pick(comp, select_crosses(by_brand.get(comp.key, []), settings.min_yes, settings.max_no)))
        if comps and not picks:
            reason = "в своде нет кроссов ни на одного конкурента бренда"
        results.append(ItemResult(item, level, pos_row, picks, reason))
    return results


# ---------------------------------------------------------------- Excel

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
FILL_YES = PatternFill("solid", fgColor="E2EFDA")
FILL_NO = PatternFill("solid", fgColor="FCE4D6")
FILL_NONE = PatternFill("solid", fgColor="EDEDED")
FILL_MIX = PatternFill("solid", fgColor="FFF2CC")
KIND_FILL = {"Да": FILL_YES, "Нет": FILL_NO, "Да + Нет": FILL_MIX, "не найден": FILL_NONE}
TITLE_FONT = Font(bold=True, size=13)


def write_table(ws, headers: list[str], rows: list[list], widths: dict[int, int] | None = None):
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for col, title in enumerate(headers, 1):
        sample = [len(str(r[col - 1])) for r in rows[:500] if col - 1 < len(r) and r[col - 1] is not None]
        width = (widths or {}).get(col) or min(45, max(len(title) * 0.9, *(sample or [0])) + 2)
        ws.column_dimensions[get_column_letter(col)].width = max(8, width)
    ws.row_dimensions[1].height = 32


def item_head(r: ItemResult, extra_names: list[str]) -> list:
    it = r.item
    category = it.category or (r.pos_row.category if r.pos_row else "")
    group = it.group or (r.pos_row.group if r.pos_row else "")
    return [it.article, it.brand, *[it.extras.get(n) for n in extra_names], category, group, it.subgroup]


def write_workbook(path: Path, results: list[ItemResult], ctx: dict) -> dict:
    wb = openpyxl.Workbook()
    extra = ctx["extra_names"]
    head = ["Артикул", "Бренд", *extra, "Категория", "Группа", "Подгруппа"]
    monitoring = ctx["positioning"].monitoring
    stats = Counter()

    # 1. Для заноса: одна строка = один кросс.
    ws = wb.active
    ws.title = "Для заноса"
    long_rows, long_status = [], []
    for r in results:
        for pick in r.picks:
            comp = pick.competitor
            for cross, status in pick.crosses:
                long_rows.append(
                    [
                        *item_head(r, extra),
                        comp.order,
                        comp.name,
                        comp.percent,
                        cross.article,
                        cross.brand,
                        "Да" if cross.yes else "Нет",
                        ", ".join(monitoring.get(comp.key, [])),
                    ]
                )
                long_status.append(status)
                stats["cross_yes" if status == STATUS_YES else "cross_no"] += 1
    long_headers = [
        *head,
        "№ конкурента",
        "Бренд конкурент",
        "% отклонения",
        "Артикул кросс",
        "Бренд кросс",
        "Использовать для сравнения цен",
        "Где мониторить цену",
    ]
    write_table(ws, long_headers, long_rows)
    status_col = long_headers.index("Использовать для сравнения цен") + 1
    for i, status in enumerate(long_status, start=2):
        if status != STATUS_YES:
            for col in range(1, len(long_headers) + 1):
                ws.cell(row=i, column=col).fill = FILL_NO
        else:
            ws.cell(row=i, column=status_col).fill = FILL_YES

    # 2. По артикулам: одна строка = один ваш артикул, конкуренты в столбцах.
    ws = wb.create_sheet("По артикулам")
    slots = max((len(r.picks) for r in results), default=0)
    wide_headers = list(head) + ["Сопоставлено по"]
    for n in range(1, slots + 1):
        wide_headers += [f"Конкурент {n}", f"% {n}", f"Кроссы {n}", f"Тип {n}"]
    wide_headers.append("Итог")
    wide_rows, kinds = [], []
    for r in results:
        if not r.picks:
            continue
        row = [*item_head(r, extra), r.level]
        row_kinds = []
        for pick in r.picks:
            row += [
                pick.competitor.name,
                pick.competitor.percent,
                "; ".join(c.article for c, _ in pick.crosses),
                pick.kind,
            ]
            row_kinds.append(pick.kind)
            stats[f"pair_{pick.kind}"] += 1
        row += [""] * (4 * (slots - len(r.picks)))
        found = sum(k != "не найден" for k in row_kinds)
        row.append(f"кроссы есть по {found} из {len(row_kinds)}" if found else "кроссов нет")
        wide_rows.append(row)
        kinds.append(row_kinds)
    write_table(ws, wide_headers, wide_rows)
    first_slot = len(head) + 2
    for i, row_kinds in enumerate(kinds, start=2):
        for n, kind in enumerate(row_kinds):
            ws.cell(row=i, column=first_slot + 4 * n + 3).fill = KIND_FILL[kind]

    # 3. Не найдено: пары «артикул — конкурент», по которым в своде нет кроссов.
    ws = wb.create_sheet("Не найдено")
    nf_rows = [
        [*item_head(r, extra), p.competitor.order, p.competitor.name, p.competitor.percent]
        for r in results
        for p in r.picks
        if not p.crosses
    ]
    write_table(ws, [*head, "№ конкурента", "Бренд конкурент", "% отклонения"], nf_rows)

    # 4. Без позиционирования.
    ws = wb.create_sheet("Без позиционирования")
    np_rows = [[*item_head(r, extra), r.reason] for r in results if not r.picks]
    write_table(ws, [*head, "Причина"], np_rows)
    stats["no_pos"] = len(np_rows)

    write_summary(wb.create_sheet("Сводка", 0), results, ctx, stats)
    wb.active = 0
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return stats


def write_summary(ws, results: list[ItemResult], ctx: dict, stats: Counter):
    s: Settings = ctx["settings"]
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 18
    ws.column_dimensions["B"].width = 45

    def title(text):
        ws.append([])
        ws.append([text])
        ws.cell(row=ws.max_row, column=1).font = TITLE_FONT

    ws.append(["Подбор кросс-артикулов по позиционированию"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=15)
    ws.append(["Сформировано", dt.datetime.now().strftime("%d.%m.%Y %H:%M")])
    ws.append(["Режим", ctx["mode"]])
    ws.append(["Кроссы свод", ctx["cross_source"]])
    ws.append(["Прайс", ctx["price_source"] or "не указан"])
    ws.append(["Позиционирование", ctx["positioning"].source])
    ws.append(["Правило отбора", rule_text(s)])

    title("Итоги")
    with_pos = sum(1 for r in results if r.picks)
    ws.append(["Артикулов обработано", len(results)])
    ws.append(["  из них с конкурентами по позиционированию", with_pos])
    ws.append(["  без позиционирования (лист «Без позиционирования»)", stats["no_pos"]])
    ws.append(["Кроссов отобрано «Да»", stats["cross_yes"]])
    ws.append(["Кроссов добрано «Нет»", stats["cross_no"]])
    ws.append(["Пар «артикул — конкурент» без кроссов (лист «Не найдено»)", stats["pair_не найден"]])
    if ctx.get("price_dups"):
        ws.append(["Повторы артикулов в прайсе (пропущены)", ctx["price_dups"]])

    title("Покрытие по брендам-конкурентам")
    header = ["Бренд → конкурент", "Артикулов", "Есть «Да»", "Только «Нет» (добор)", "Да + Нет", "Не найдено", "Покрытие «Да»", "Покрытие всего"]
    ws.append(header)
    hdr_row = ws.max_row
    for col in range(1, len(header) + 1):
        c = ws.cell(row=hdr_row, column=col)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    cover: dict[tuple[str, str], Counter] = defaultdict(Counter)
    for r in results:
        for p in r.picks:
            cover[(r.item.brand.upper(), p.competitor.name)][p.kind] += 1
    for (brand, comp), cnt in sorted(cover.items()):
        total = sum(cnt.values())
        yes_any = cnt["Да"] + cnt["Да + Нет"]
        ws.append(
            [
                f"{brand} → {comp}",
                total,
                cnt["Да"],
                cnt["Нет"],
                cnt["Да + Нет"],
                cnt["не найден"],
                yes_any / total,
                (total - cnt["не найден"]) / total,
            ]
        )
        ws.cell(row=ws.max_row, column=7).number_format = "0%"
        ws.cell(row=ws.max_row, column=8).number_format = "0%"

    if ctx["unused"]:
        title("Кроссы в своде на бренды, которых нет в позиционировании этого бренда (не используются)")
        for (brand, cross_brand), n in sorted(ctx["unused"].items()):
            ws.append([f"{brand} → {cross_brand}", n])

    title("Как распознаны столбцы")
    for label, pairs in ctx["columns"]:
        for key, desc in pairs:
            ws.append([f"{label}: {key}", desc])

    if ctx["warnings"]:
        title("Предупреждения")
        for w in ctx["warnings"]:
            ws.append([w])


def rule_text(s: Settings) -> str:
    limit = "все найденные" if not s.max_no else f"не больше {s.max_no}"
    if s.min_yes <= 1:
        cond = "если кроссов «Да» нет"
    else:
        cond = f"если кроссов «Да» меньше {s.min_yes}"
    return f"Берём все кроссы «Да»; {cond} — добираем кроссы «Нет» ({limit}) по каждому бренду-конкуренту."


# ---------------------------------------------------------------- сборка


def collect_files(paths: list[Path]) -> list[Path]:
    files = []
    for p in paths:
        if p.is_dir():
            files += sorted(x for x in p.iterdir() if x.is_file())
        elif p.exists():
            files.append(p)
        else:
            raise UserError(f"Файл не найден: {p}")
    return [f for f in files if not f.name.startswith(("~$", ".")) and f.suffix.lower() in SUPPORTED_EXT | {".xls"}]


def classify(files: list[Path], settings: Settings, warnings: list[str]):
    found = {"crosses": [], "price": [], "positioning": []}
    for f in files:
        try:
            tables = read_tables(f)
        except UserError:
            raise
        except Exception as e:  # noqa: BLE001
            warnings.append(f"{f.name}: не удалось прочитать ({e})")
            continue
        if RESULT_SHEETS <= {t.sheet for t in tables}:
            warnings.append(f"{f.name}: это файл результата — пропущен")
            continue
        for t in tables:
            if (loc := locate(t, settings.cross_cols, CROSS_REQUIRED, CROSS_OPTIONAL)):
                found["crosses"].append((f, loc))
            elif (loc := locate_positioning(t, settings)):
                found["positioning"].append((f, loc))
            elif (loc := locate(t, settings.price_cols, PRICE_REQUIRED, PRICE_OPTIONAL)):
                found["price"].append((f, loc))
            else:
                continue
            break  # из одного файла берём первый распознанный лист
        else:
            warnings.append(f"{f.name}: не похож ни на кроссы, ни на прайс, ни на позиционирование — пропущен")
    for kind, label in (("crosses", "кроссов"), ("price", "прайсов"), ("positioning", "файлов позиционирования")):
        if len(found[kind]) > 1:
            found[kind].sort(key=lambda x: x[0].stat().st_mtime, reverse=True)
            names = ", ".join(f.name for f, _ in found[kind])
            warnings.append(f"Найдено несколько {label} ({names}) — взят самый свежий: {found[kind][0][0].name}")
    return {k: (v[0][1] if v else None) for k, v in found.items()}


def run(paths: list[Path], output_dir: Path, settings: Settings) -> tuple[Path, dict]:
    warnings: list[str] = []
    files = collect_files(paths)
    if not files:
        raise UserError(
            "Не найдено ни одного файла. Положите «Кроссы свод» и прайс (xlsx или csv) в папку «Входящие» "
            "или перетащите файлы на Запуск.bat."
        )
    found = classify(files, settings, warnings)
    cross_loc, price_loc, pos_loc = found["crosses"], found["price"], found["positioning"]
    if not cross_loc:
        need = ", ".join(f"«{settings.cross_cols[k][0]}»" for k in CROSS_REQUIRED)
        raise UserError(f"Не найден файл кроссов. В нём должны быть столбцы: {need}. Проверьте настройки.ini.")

    columns = [("Кроссы", cross_loc.describe_columns())]
    if pos_loc:
        positioning = parse_positioning(pos_loc, pos_loc.table.source, builtin_monitoring())
        columns.append(("Позиционирование", pos_loc.describe_columns()))
    else:
        positioning = builtin_positioning(settings)

    base = parse_crosses(cross_loc)
    if base.rows_skipped:
        warnings.append(f"В кроссах пропущено строк без артикула/бренда: {base.rows_skipped}")

    extra_names, price_dups = [], 0
    if price_loc:
        items, extra_names, price_dups = parse_price(price_loc, settings)
        columns.append(("Прайс", price_loc.describe_columns() + [("доп. столбцы", ", ".join(extra_names) or "—")]))
        if "подгруппа" in price_loc.cols:
            mode = "по прайсу: конкуренты по бренду + подгруппе"
        elif "группа" in price_loc.cols:
            mode = "по прайсу: конкуренты по бренду + группе (в прайсе нет столбца подгруппы)"
            warnings.append("В прайсе нет столбца подгруппы — конкуренты взяты по группе (объединение подгрупп).")
        else:
            mode = "по прайсу: конкуренты по бренду (в прайсе нет группы/подгруппы)"
            warnings.append("В прайсе нет столбцов группы/подгруппы — конкуренты взяты по всем подгруппам бренда.")
    else:
        items = items_from_crosses(base)
        mode = "без прайса: все артикулы из свода, конкуренты по всем подгруппам бренда"
        warnings.append(
            "Прайс не найден — подгруппы артикулов неизвестны, поэтому конкуренты взяты по всем подгруппам бренда. "
            "Положите прайс в «Входящие», чтобы отбор шёл по подгруппе."
        )

    unused: Counter = Counter()
    for (brand_key, _), by_brand in base.index.items():
        allowed = positioning.competitor_keys(brand_key)
        for cross_brand_key, crosses in by_brand.items():
            if cross_brand_key not in allowed:
                unused[(brand_key, crosses[0].brand)] += len(crosses)

    results = build(items, base, positioning, settings)
    ctx = {
        "settings": settings,
        "mode": mode,
        "cross_source": f"{cross_loc.table.source} — строк: {base.rows_total}",
        "price_source": f"{price_loc.table.source} — артикулов: {len(items)}" if price_loc else "",
        "positioning": positioning,
        "extra_names": extra_names,
        "price_dups": price_dups,
        "unused": unused,
        "columns": columns,
        "warnings": warnings,
    }
    out = output_dir / f"Кроссы_по_позиционированию_{dt.datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx"
    stats = write_workbook(out, results, ctx)
    stats["items"] = len(results)
    stats["warnings"] = warnings
    stats["mode"] = mode
    return out, stats


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass
    ap = argparse.ArgumentParser(description="Подбор кросс-артикулов по позиционированию")
    ap.add_argument("files", nargs="*", type=Path, help="файлы или папки (по умолчанию — папка «Входящие»)")
    ap.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT_DIR, help="папка для результата")
    ap.add_argument("-s", "--settings", type=Path, default=DEFAULT_SETTINGS, help="файл настроек")
    ap.add_argument("--no-open", action="store_true", help="не открывать результат")
    args = ap.parse_args(argv)

    try:
        settings = load_settings(args.settings)
        out, stats = run(args.files or [DEFAULT_INPUT_DIR], args.output, settings)
    except UserError as e:
        print(f"\nОШИБКА: {e}\n")
        return 1

    print(f"\nРежим: {stats['mode']}")
    print(f"Артикулов: {stats['items']}  |  без позиционирования: {stats['no_pos']}")
    print(f"Кроссов «Да»: {stats['cross_yes']}  |  добрано «Нет»: {stats['cross_no']}  |  пар без кроссов: {stats['pair_не найден']}")
    for w in stats["warnings"]:
        print(f"! {w}")
    print(f"\nГотово: {out}\n")
    if settings.open_result and not args.no_open and hasattr(os, "startfile"):
        os.startfile(out)  # type: ignore[attr-defined]
    return 0


if __name__ == "__main__":
    sys.exit(main())
